"""Generate ULT-10611 (Steel Origin | Add Steel Origin column to the Card History report, on-demand and scheduled) test cases DOCX."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = Path(r"c:\Prcatices-AI\docs\ULT-10611_test_cases.docx")

COLS = ["ID", "Title", "Preconditions", "Steps", "Expected Result", "Pri", "Type", "AC"]
WIDTHS = [0.55, 1.35, 1.55, 2.6, 2.3, 0.4, 0.7, 0.45]

PRE = "ULT-10607 delivered (steelorigin stored in cycle IGParams on Set Shipment Information). Steel-flagged supplier item with a kanban card."
PRE_AA = PRE + " Card C1 cycle shipped with Steel Origin 'AA'."
PRE_NONE = "Card C3 of a non-flagged item (or flagged but never shipped), so no steelorigin in IGParams."

SECTIONS = [
    ("1. Steel Origin Shown for a Shipped Flagged Card", [
        ("TC-01", "Shipped cycle with 'AA' - buyer export",
         PRE_AA + " Logged in as buyer.",
         "1. Reports > Card History (buyer side).\n2. Select a date range covering the shipment and export to Excel.\n3. Locate the row for C1's cycle.",
         "Excel downloads; the row for the cycle has Steel Origin = 'AA'.", "High", "Functional", "AC1"),
        ("TC-02", "Shipped cycle with 'AA' - supplier export",
         PRE_AA + " Logged in as supplier.",
         "1. Reports > Card History (supplier side).\n2. Export to Excel for the same date range.\n3. Locate the row.",
         "Row shows Steel Origin = 'AA'; same as buyer export.", "High", "Functional", "AC1"),
        ("TC-03", "Shipped cycle with 'US'",
         "Card C2 cycle shipped with Steel Origin 'US'.",
         "1. Export Card History covering the cycle.\n2. Locate the row.",
         "Steel Origin = 'US'.", "High", "Functional", "AC1"),
        ("TC-04", "Admin user export",
         PRE_AA + " Logged in as admin.",
         "1. Export Card History.\n2. Locate the row.",
         "Steel Origin = 'AA' visible to admin.", "Medium", "Functional", "AC1"),
        ("TC-05", "Shipped but not yet received (in-transit current cycle)",
         "Card C1 shipped with 'AA', not yet received (current-cycle branch of EboKanbanCardHistoryView2).",
         "1. Export Card History.",
         "Row for the current cycle shows 'AA' (value read from KCD / cycle-data IGParams).", "High", "Functional", "AC1"),
        ("TC-06", "One row per card cycle - no duplicates from new join",
         "Card with 3 completed cycles and a current cycle.",
         "1. Export Card History.\n2. Count rows for the card.\n3. Compare with row count on previous build.",
         "Exactly one row per cycle (4 rows); LEFT JOIN on EboKanbanCycleData does not duplicate rows.", "High", "Regression", "AC1/AC6"),
        ("TC-07", "Column position and header",
         PRE_AA,
         "1. Export Card History.\n2. Inspect header row.",
         "New column header is 'Steel Origin' and is placed directly after Packing Slip No.", "High", "Functional", "AC1"),
        ("TC-08", "Multiple cards in one export",
         "Cards shipped with US, AA and no value in the same period.",
         "1. Export Card History.",
         "Each row shows its own value (US / AA / blank); no value leaks between rows.", "High", "Functional", "AC1/AC4"),
        ("TC-09", "Filtered export (supplier / item / date filters)",
         PRE_AA + " Other cards present with different suppliers/items.",
         "1. Apply supplier, item and date filters.\n2. Export.",
         "Only filtered rows are returned; Steel Origin populated correctly for each.", "Medium", "Functional", "AC1"),
    ]),
    ("2. Completed Cycles Keep Their Own Value", [
        ("TC-10", "US in cycle 1, AA in cycle 2, both received",
         "Card C1: cycle 1 shipped with 'US' and received; cycle 2 shipped with 'AA' and received.",
         "1. Export Card History covering both cycles.\n2. Compare the two rows for C1.",
         "Cycle 1 row shows 'US'; cycle 2 row shows 'AA'.", "High", "Functional", "AC2"),
        ("TC-11", "Completed cycle (cross-tab history branch)",
         "Card C1 cycle shipped with 'AA' and received, so it is served by EboKbnCrossTabHistoryView.",
         "1. Export Card History.",
         "Row for the completed cycle shows 'AA' (value taken via CYD.IGParams by CycleGID = CT.CurrentCycleID).", "High", "Functional", "AC2"),
        ("TC-12", "Current cycle value differs from completed cycles",
         "C1: cycle 1 'US' received; current cycle 2 shipped 'AA' not received.",
         "1. Export Card History.",
         "Completed cycle = 'US', current cycle = 'AA'; values from both UNION branches are correct.", "High", "Functional", "AC2"),
        ("TC-13", "Later cycle without value does not overwrite earlier cycle",
         "C1: cycle 1 'AA' received; cycle 2 shipped after the item was unflagged (no value).",
         "1. Export Card History.",
         "Cycle 1 = 'AA', cycle 2 = blank.", "High", "Functional", "AC2/AC4"),
        ("TC-14", "Several completed cycles (5+)",
         "Card with 5 received cycles alternating US/AA.",
         "1. Export Card History.",
         "Every cycle shows its own value in the correct order.", "Medium", "Functional", "AC2"),
        ("TC-15", "Completed cycle with no matching cycle-data row",
         "Old completed cycle where EboKanbanCycleData has no row for CycleGID (archived / purged).",
         "1. Export Card History.",
         "Row is still returned (LEFT JOIN) with blank Steel Origin; no row dropped; no error.", "High", "Negative", "AC2/AC6"),
        ("TC-16", "Completed cycles of the same card in different CycleGIDs",
         "Card with two completed cycles.",
         "1. Check CycleGID per row.\n2. Compare Steel Origin to the value stored per cycle in EboKanbanCycleData.IGParams.",
         "Each row joins to its own CycleGID; value equals EboGetFlexValue(IGParams, ',', 'steelorigin') of that cycle.", "Medium", "DB", "AC2"),
    ]),
    ("3. Corrected Value", [
        ("TC-17", "Value corrected before receipt",
         "Card shipped with 'US', then corrected to 'AA' on Set Shipment Information before receipt.",
         "1. Export Card History.",
         "Row shows corrected value 'AA'; no 'US'.", "High", "Functional", "AC3"),
        ("TC-18", "Corrected from AA to US",
         "Card shipped with 'AA', corrected to 'US' before receipt.",
         "1. Export Card History.",
         "Row shows 'US'.", "High", "Functional", "AC3"),
        ("TC-19", "Corrected value retained after receipt",
         "As TC-17, then the card is received.",
         "1. Export Card History.",
         "Completed cycle still shows corrected 'AA'.", "High", "Functional", "AC2/AC3"),
        ("TC-20", "Multiple corrections",
         "Value changed US -> AA -> US before receipt.",
         "1. Export Card History.",
         "Row shows the last saved value ('US').", "Medium", "Functional", "AC3"),
        ("TC-21", "Corrected cycle does not affect other cycles",
         "C1 cycle 1 'US' received; cycle 2 shipped 'US' and corrected to 'AA'.",
         "1. Export Card History.",
         "Cycle 1 remains 'US'; cycle 2 = 'AA'.", "Medium", "Functional", "AC2/AC3"),
    ]),
    ("4. Unflagged Cards - Blank Cell", [
        ("TC-22", "Never-flagged item",
         PRE_NONE,
         "1. Export Card History.\n2. Locate the row.",
         "Steel Origin cell is blank (empty, not 'NULL' / '0' / 'N/A').", "High", "Functional", "AC4"),
        ("TC-23", "Flagged card, not yet shipped with a value",
         "Flagged item; card in the current cycle with no steelorigin stored.",
         "1. Export Card History.",
         "Steel Origin cell blank.", "High", "Functional", "AC4"),
        ("TC-24", "Flag changed after shipment",
         "Card shipped with 'AA'; item later unflagged.",
         "1. Export Card History.",
         "Row still shows 'AA' (stored value, not current flag).", "Medium", "Functional", "AC1/AC4"),
        ("TC-25", "steelorigin key present but empty",
         "IGParams contains 'steelorigin=' with no value.",
         "1. Export Card History.",
         "Cell blank; no error.", "Medium", "Negative", "AC4"),
        ("TC-26", "IGParams NULL / empty",
         "Cycle with IGParams NULL or ''.",
         "1. Export Card History.",
         "Cell blank; row still present; no SQL error from EboGetFlexValue.", "High", "Negative", "AC4"),
        ("TC-27", "Other IGParams keys do not interfere",
         "IGParams contains other keys (e.g. 'steelorigin2', 'xsteelorigin') but no 'steelorigin'.",
         "1. Export Card History.",
         "Cell blank; only the exact key 'steelorigin' is read.", "Low", "Negative", "AC4"),
        ("TC-28", "Consumption cards (EboConsumptionCardHistoryView)",
         "Consumption card history present.",
         "1. Export Card History for consumption cards.",
         "Report works; Steel Origin cell is blank ('' AS SteelOrigin) unless added to the view; column count equal across views.", "Medium", "Functional", "AC4/AC6"),
    ]),
    ("5. Scheduled Report", [
        ("TC-29", "Scheduled Card History contains Steel Origin",
         "Scheduled Card History report configured (UEKDataExtractionSvc JobExpReportScheduler).",
         "1. Trigger / wait for the scheduled run.\n2. Open the generated Excel / output file.",
         "File contains Steel Origin column with same header, same position and the same values as the on-demand export.", "High", "Functional", "AC5"),
        ("TC-30", "Scheduled value equals on-demand value",
         PRE_AA + " Same date range and filters.",
         "1. Run on-demand export.\n2. Run scheduled export.\n3. Compare rows.",
         "Both outputs identical including Steel Origin for each row.", "High", "Functional", "AC5"),
        ("TC-31", "Scheduled - completed cycles with US / AA",
         "As TC-10.",
         "1. Run the scheduled report.",
         "Cycle 1 = 'US', cycle 2 = 'AA'.", "High", "Functional", "AC2/AC5"),
        ("TC-32", "Scheduled - unflagged cards blank",
         PRE_NONE,
         "1. Run the scheduled report.",
         "Steel Origin blank.", "Medium", "Functional", "AC4/AC5"),
        ("TC-33", "Scheduled - both code paths (JobExpReportScheduler lines 2805 and 3749)",
         "Schedules that exercise both call sites (e.g. different views / filters).",
         "1. Run both schedules.\n2. Inspect outputs.",
         "Both call sites return the new column; no exception in the service log.", "High", "Functional", "AC5"),
        ("TC-34", "Scheduled - delivery (email / file share)",
         "Schedule configured to email or save the file.",
         "1. Run the schedule.\n2. Open the delivered file.",
         "File delivered and opens correctly with the new column.", "Medium", "E2E", "AC5"),
        ("TC-35", "Scheduled - EboHelperLib copy updated in every strColList branch",
         "Source access.",
         "1. Review CoreAppNetDev\\EboHelperLib\\EboReportHelper.cs GetCardHistory.",
         "SteelOrigin (or '' AS SteelOrigin) is present in all branches, matching LeanSuite copy.", "High", "Code review", "AC5"),
        ("TC-36", "Scheduled - customer-specific layout",
         "Schedule for a customer with a custom Card History grid override.",
         "1. Run the schedule.",
         "Layout unchanged vs previous build apart from the new column (where applicable); no error.", "Medium", "Regression", "AC5/AC6"),
    ]),
    ("6. Views and Database (10_CEK108 / 12_CEK1010)", [
        ("TC-37", "EboKanbanCardHistoryView2 exposes SteelOrigin in both UNION branches",
         "DB updated with 12_CEK1010.sql.",
         "1. SELECT TOP 100 * FROM EboKanbanCardHistoryView2.\n2. Check column SteelOrigin.\n3. Query rows from completed and current cycle.",
         "SteelOrigin exists; both completed and current branches return values; UNION column count and types are compatible.", "High", "DB", "AC1/AC2"),
        ("TC-38", "EboKanbanCardHistoryView exposes SteelOrigin",
         "As above.",
         "1. SELECT TOP 100 * FROM EboKanbanCardHistoryView.",
         "SteelOrigin column present and populated correctly.", "High", "DB", "AC1"),
        ("TC-39", "Value derived via EboGetFlexValue",
         "Cycle with IGParams containing steelorigin=AA among other keys.",
         "1. Compare SteelOrigin in the view with dbo.EboGetFlexValue(IGParams, ',', 'steelorigin').",
         "Values are identical.", "Medium", "DB", "AC1"),
        ("TC-40", "Script is re-runnable (DROP/CREATE)",
         "12_CEK1010.sql run once.",
         "1. Re-run 12_CEK1010.sql.\n2. Check for errors and that views exist.",
         "No error; views recreated with SteelOrigin.", "High", "DB", "AC6"),
        ("TC-41", "Script conventions",
         "Code review.",
         "1. Review 12_CEK1010.sql.",
         "Views in views section with DROP/CREATE; @CURRENT_REL_VERSION bumped; 'last edit' line added.", "Medium", "Code review", "AC6"),
        ("TC-42", "Existing view columns unchanged",
         "Previous build view definition available.",
         "1. Compare column names, order and types of the views before and after (excluding SteelOrigin).",
         "Only SteelOrigin is added; no existing column changed.", "High", "Regression", "AC6"),
        ("TC-43", "Dependent objects still compile",
         "Objects that reference these views (reports, other views, SPs).",
         "1. Run sp_refreshview / check dependencies after the script.\n2. Run dependent reports.",
         "No invalid object errors; dependent reports work.", "High", "Regression", "AC6"),
        ("TC-44", "Performance on large (Copeland-sized) DB",
         "Large DB with many cycles.",
         "1. Capture execution plan and runtime of Card History export (before/after).\n2. Check the LEFT JOIN on EboKanbanCycleData CYD by CycleGID.",
         "Index seek on CycleGID; runtime does not materially regress; export completes without timeout.", "High", "Performance", "AC6"),
        ("TC-45", "Zurn X* view variants",
         "Zurn views script (MiscScripts\\ZurnViews\\Zurn new Views.sql) applied.",
         "1. Run Card History export on a Zurn configuration.",
         "Report works without errors; if the shared column list references SteelOrigin the Zurn views either expose it or '' AS SteelOrigin is used.", "Medium", "Regression", "AC6"),
        ("TC-46", "Dataset shape for views without the column",
         "A view in a strColList branch that does not expose SteelOrigin.",
         "1. Run the export for that branch.",
         "Query succeeds with '' AS SteelOrigin; column present and blank.", "High", "Functional", "AC4/AC6"),
    ]),
    ("7. Grid / UI Configuration and Customer Overrides", [
        ("TC-47", "CardHistoryExcel contains SteelOrigin after PackingSlipNo",
         "Baseline config.",
         "1. Check BaseLineUIConfig.xml CardHistoryExcel.\n2. Export.",
         "UIElement SteelOrigin (header 'Steel Origin') exists with an <include> after PackingSlipNo; column appears after Packing Slip No.", "High", "Functional", "AC1"),
        ("TC-48", "CardHistoryExcel2_1 inherits new column",
         "Environment that uses CardHistoryExcel2_1.",
         "1. Export using this grid.",
         "Steel Origin column present.", "High", "Functional", "AC1/AC5"),
        ("TC-49", "CardHistoryExcel2_2 inherits new column",
         "Environment that uses CardHistoryExcel2_2.",
         "1. Export using this grid.",
         "Steel Origin column present.", "High", "Functional", "AC1/AC5"),
        ("TC-50", "CustomBaseLineUIConfig.xml override",
         "Customer using CustomBaseLineUIConfig.xml.",
         "1. Export Card History.\n2. Compare to previous build.",
         "Existing columns and order unchanged; Steel Origin added only where intended (or consciously left out) and consistent with the data.", "High", "Regression", "AC6"),
        ("TC-51", "CustomQAUIConfig.xml override",
         "QA custom config active.",
         "1. Export Card History.",
         "Existing layout unchanged; no error; new column handled per design.", "Medium", "Regression", "AC6"),
        ("TC-52", "CustomTFUSAUIConfig.xml (both copies)",
         "TFUSA config active (incl. CoreAppDevCustom\\TFUSA copy).",
         "1. Export Card History.",
         "Layout unchanged apart from the new column.", "Medium", "Regression", "AC6"),
        ("TC-53", "CustomTraneUIConfig.xml override",
         "Trane config active.",
         "1. Export Card History.",
         "Layout unchanged apart from the new column.", "Medium", "Regression", "AC6"),
        ("TC-54", "CustomLS2BaseLineUIConfig.xml override",
         "LS2 baseline custom config active.",
         "1. Export Card History.",
         "Layout unchanged apart from the new column.", "Medium", "Regression", "AC6"),
        ("TC-55", "Customer overrides that redefine the column list",
         "Customers with overrides at the listed line locations.",
         "1. Export per customer.\n2. Check column list against override definition.",
         "Overridden column list is respected; no missing/misaligned columns or shifted headers.", "High", "Regression", "AC6"),
        ("TC-56", "Header text and column width",
         PRE_AA,
         "1. Open the Excel export.\n2. Check header text, width and formatting.",
         "Header 'Steel Origin'; text is readable; format consistent with neighbouring text columns.", "Low", "UI", "AC1"),
    ]),
    ("8. No Regression", [
        ("TC-57", "All existing columns unchanged",
         "Previous build export of the same data available.",
         "1. Export Card History on new and old builds with identical filters.\n2. Compare all columns except Steel Origin.",
         "Names, order, values and formats identical; only Steel Origin is new.", "High", "Regression", "AC6"),
        ("TC-58", "Row count unchanged",
         "As above.",
         "1. Compare row count old vs new.",
         "Same number of rows.", "High", "Regression", "AC6"),
        ("TC-59", "Business unit columns (BSU.Code, BSU.Name)",
         "Multi-business-unit user.",
         "1. Export Card History.",
         "BSU Code / Name columns remain present and correct.", "Medium", "Regression", "AC6"),
        ("TC-60", "All four strColList branches",
         "Scenarios that select each view (via AppendCardHistoryActionFilterAndReturnViewName), including consumption.",
         "1. Run the export for each branch (buyer and supplier).",
         "Every branch runs without SQL error and includes the Steel Origin column.", "High", "Functional", "AC1/AC6"),
        ("TC-61", "Action filters (AppendCardHistoryActionFilter)",
         "Card History with action filters (e.g. shipped, received, cancelled).",
         "1. Apply each action filter and export.",
         "Filters behave as before; Steel Origin populated per row.", "Medium", "Regression", "AC6"),
        ("TC-62", "Buyer vs supplier controllers parity",
         PRE_AA,
         "1. Export from BuyerExcelReportController.CardHistory and SupplierExelReportController.ExcelReports_CardHistory.\n2. Compare.",
         "Both exports contain identical Steel Origin column and values.", "Medium", "Functional", "AC1/AC6"),
        ("TC-63", "Large export",
         "Large date range (thousands of cycles).",
         "1. Export Card History.",
         "Completes; Steel Origin populated; no timeout or memory error.", "Medium", "Performance", "AC6"),
        ("TC-64", "Localization / special characters",
         "User with non-English language.",
         "1. Export Card History.",
         "Report works; header displays 'Steel Origin' (or translated if resource exists); data not corrupted.", "Low", "Regression", "AC6"),
    ]),
]

AC_MATRIX = [
    ("AC1 - Steel Origin shown for a shipped flagged card", "TC-01 to TC-09, TC-24, TC-37 to TC-39, TC-47 to TC-49, TC-56, TC-60, TC-62"),
    ("AC2 - Completed cycles keep their own value", "TC-10 to TC-16, TC-19, TC-21, TC-31, TC-37"),
    ("AC3 - Corrected value", "TC-17 to TC-21"),
    ("AC4 - Unflagged cards blank", "TC-08, TC-13, TC-22 to TC-28, TC-32, TC-46"),
    ("AC5 - Scheduled report", "TC-29 to TC-36, TC-48, TC-49"),
    ("AC6 - No regression / customer layouts", "TC-06, TC-15, TC-28, TC-36, TC-40 to TC-46, TC-50 to TC-55, TC-57 to TC-64"),
]


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def write(cell, text, size=8, bold=False, color=None, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lines = text.split("\n")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if i < len(lines) - 1:
            run.add_break()


def heading(doc, text, size=13):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def set_widths(table, widths):
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Inches(0.5))

t = doc.add_paragraph()
r = t.add_run("ULT-10611 - Test Cases")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
doc.add_paragraph("Steel Origin | Add Steel Origin column to the Card History report (on-demand and scheduled)").runs[0].font.size = Pt(11)

info = doc.add_table(rows=0, cols=2)
info.style = "Table Grid"
for k, v in [
    ("Ticket", "ULT-10611 (Epic ULT-10604) - Story - Status: In QA Review - Fix version 10.10.6"),
    ("Labels", "Copeland, SteelOrigin"),
    ("Depends on", "ULT-10607 (Steel Origin stored in cycle IGParams)"),
    ("Scope", "Card History Excel export (buyer: BuyerExcelReportController.CardHistory; supplier: SupplierExelReportController.ExcelReports_CardHistory). "
              "Views EboKanbanCardHistoryView2 (both UNION branches) and EboKanbanCardHistoryView get SteelOrigin = dbo.EboGetFlexValue(IGParams, ',', 'steelorigin'); "
              "completed cycles via LEFT JOIN EboKanbanCycleData CYD ON CYD.CycleGID = CT.CurrentCycleID. "
              "SteelOrigin added to every strColList branch in LeanSuite and EboHelperLib GetCardHistory ('' AS SteelOrigin where the view lacks it). "
              "New UIElement 'Steel Origin' in BaseLineUIConfig CardHistoryExcel after PackingSlipNo (inherited by 2_1 and 2_2). Views redefined in 12_CEK1010.sql."),
    ("Scheduled copy", "CoreAppNetDev\\EboHelperLib\\EboReportHelper.cs GetCardHistory, used by UEKDataExtractionSvc JobExpReportScheduler (2 call sites)."),
    ("Customer overrides", "CustomBaseLineUIConfig, CustomQAUIConfig, CustomTFUSAUIConfig (both copies), CustomTraneUIConfig, CustomLS2BaseLineUIConfig."),
    ("Out of scope", "Consumption card history (not applicable unless trivial); Zurn X* views changed only if the shared column list would break."),
    ("Open points", "Whether customer overrides get the new column or stay unchanged (TC-50 to TC-55); Zurn view handling (TC-45); consumption view behaviour (TC-28)."),
    ("Total test cases", str(sum(len(c) for _, c in SECTIONS))),
]:
    row = info.add_row()
    write(row.cells[0], k, size=9, bold=True)
    shade(row.cells[0], "DCE6F1")
    write(row.cells[1], v, size=9)
set_widths(info, [1.5, 8.5])

for title, cases in SECTIONS:
    heading(doc, f"{title} ({len(cases)})")
    tbl = doc.add_table(rows=1, cols=len(COLS))
    tbl.style = "Table Grid"
    for i, h in enumerate(COLS):
        write(tbl.rows[0].cells[i], h, size=9, bold=True, color="FFFFFF", center=True)
        shade(tbl.rows[0].cells[i], "1F4E79")
    for n, case in enumerate(cases):
        row = tbl.add_row()
        for i, val in enumerate(case):
            write(row.cells[i], val, bold=(i == 0), center=(i in (0, 5, 6, 7)))
            if n % 2:
                shade(row.cells[i], "F2F2F2")
    set_widths(tbl, WIDTHS)

heading(doc, "Acceptance Criteria Traceability")
m = doc.add_table(rows=1, cols=2)
m.style = "Table Grid"
for i, h in enumerate(["Acceptance Criteria", "Test Cases"]):
    write(m.rows[0].cells[i], h, size=9, bold=True, color="FFFFFF")
    shade(m.rows[0].cells[i], "1F4E79")
for a, b in AC_MATRIX:
    row = m.add_row()
    write(row.cells[0], a, size=9)
    write(row.cells[1], b, size=9)
set_widths(m, [4.0, 6.0])

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"Created: {OUT}")
