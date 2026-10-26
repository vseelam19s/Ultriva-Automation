"""One-time script to generate test_data/test_data.xlsx."""
import os
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

thin   = Side(style="thin")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

HEADER_BG  = "1F4E79"
PASS_BG    = "C6EFCE"
FAIL_BG    = "FFC7CE"
SECTION_BG = "D9E1F2"
ALT_BG     = "F2F2F2"

def hdr_font():  return Font(bold=True, color="FFFFFF", size=11, name="Calibri")
def cell_font(): return Font(size=10, name="Calibri")
def bold_font(): return Font(bold=True, size=10, name="Calibri")
def fill(h):     return PatternFill(start_color=h, end_color=h, fill_type="solid")
def center():    return Alignment(horizontal="center", vertical="center", wrap_text=True)
def left():      return Alignment(horizontal="left",   vertical="center", wrap_text=True)


def write_header(ws, headers, col_widths):
    for ci, (hdr, w) in enumerate(zip(headers, col_widths), 1):
        c = ws.cell(row=1, column=ci, value=hdr)
        c.font = hdr_font(); c.fill = fill(HEADER_BG)
        c.alignment = center(); c.border = border
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def write_row(ws, ri, vals, bg=None):
    for ci, val in enumerate(vals, 1):
        c = ws.cell(row=ri, column=ci, value=val)
        c.font = cell_font()
        c.alignment = center() if ci in (1, 2, 5, 6, 7) else left()
        c.border = border
        if bg:
            c.fill = fill(bg)
        elif ri % 2 == 0:
            c.fill = fill(ALT_BG)
    ws.row_dimensions[ri].height = 18


wb = openpyxl.Workbook()

# ── Sheet 1 : Login Test Data ─────────────────────────────────────────────────
ws1 = wb.active
ws1.title = "Login_TestData"
write_header(ws1,
    ["TC_ID", "Description", "Username", "Password", "Expected_Result", "Expected_Status"],
    [10, 45, 22, 22, 30, 16])

login_rows = [
    ("TC-001", "Login page loads correctly",                  "",               "",            "Fields visible",      "PASS"),
    ("TC-002", "Valid login - buyer user",                    "buyeruser",       "Ultriva!23",  "Login success",       "PASS"),
    ("TC-003", "Invalid username only",                       "invaliduser",     "Ultriva!23",  "Login rejected",      "PASS"),
    ("TC-004", "Invalid password only",                       "buyeruser",       "WrongPass!",  "Login rejected",      "PASS"),
    ("TC-005", "Both username and password empty",            "",               "",            "Login rejected",      "PASS"),
    ("TC-006", "Valid login then logout",                     "buyeruser",       "Ultriva!23",  "Redirect to login",   "PASS"),
    ("TC-007", "Login page browser title not empty",          "",               "",            "Non-empty title",     "PASS"),
    ("TC-008", "Username with embedded space",                "buyer user",      "Ultriva!23",  "Login rejected",      "PASS"),
    ("TC-009", "SQL injection in username",                   "' OR 1=1 --",     "Ultriva!23",  "Login rejected",      "PASS"),
    ("TC-010", "All-uppercase username",                      "BUYERUSER",       "Ultriva!23",  "Login rejected",      "PASS"),
    ("TC-011", "Empty username with valid password",          "",               "Ultriva!23",  "Login rejected",      "PASS"),
    ("TC-012", "Valid username with empty password",          "buyeruser",       "",            "Login rejected",      "PASS"),
]
for ri, row in enumerate(login_rows, 2):
    write_row(ws1, ri, row, PASS_BG if row[5] == "PASS" else FAIL_BG)

# ── Sheet 2 : Dashboard Test Data ─────────────────────────────────────────────
ws2 = wb.create_sheet("Dashboard_TestData")
write_header(ws2,
    ["TC_ID", "Description", "Login_User", "Login_Pass", "Element_To_Check", "Expected_Result", "Expected_Status"],
    [10, 45, 15, 15, 35, 28, 16])

dash_rows = [
    ("TC-101", "Dashboard loads after valid login",          "buyeruser", "Ultriva!23", "vertical-nav",             "Visible",    "PASS"),
    ("TC-102", "Vertical navigation menu is present",        "buyeruser", "Ultriva!23", "class=vertical-nav",       "Present",    "PASS"),
    ("TC-103", "URL changes to dashboard after login",       "buyeruser", "Ultriva!23", "Current URL",              "Changed",    "PASS"),
    ("TC-104", "Main content area visible",                  "buyeruser", "Ultriva!23", "class=main-content",       "Visible",    "PASS"),
    ("TC-105", "Nav menu contains icon links",               "buyeruser", "Ultriva!23", "vertical-nav-items //a",   "> 0 links",  "PASS"),
    ("TC-106", "Logout button in avatar dropdown toolbar",   "buyeruser", "Ultriva!23", "href=Account/Logout",      "Present",    "PASS"),
    ("TC-107", "Browser tab title is non-empty",             "buyeruser", "Ultriva!23", "document.title",           "Non-empty",  "PASS"),
    ("TC-108", "Page ready state within page load timeout",  "buyeruser", "Ultriva!23", "readyState == complete",   "True",       "PASS"),
    ("TC-109", "Avatar dropdown trigger button visible",     "buyeruser", "Ultriva!23", "class=avatar-btn",         "Present",    "PASS"),
    ("TC-110", "My Account link inside avatar dropdown",     "buyeruser", "Ultriva!23", "href=MyAccount/Home",      "Present",    "PASS"),
]
for ri, row in enumerate(dash_rows, 2):
    write_row(ws2, ri, row, PASS_BG if row[6] == "PASS" else FAIL_BG)

# ── Sheet 3 : Navigation Test Data ────────────────────────────────────────────
ws3 = wb.create_sheet("Navigation_TestData")
write_header(ws3,
    ["TC_ID", "Description", "URL_Path", "Login_Required", "Expected_Behavior", "Expected_Status"],
    [10, 45, 42, 16, 36, 16])

nav_rows = [
    ("TC-201", "Browser back button after login",           "/QALS1011/KanbanItem/List",    "Yes", "No crash / white screen",    "PASS"),
    ("TC-202", "Base URL without login redirects",          "/qals1011",                    "No",  "Redirect to login page",     "PASS"),
    ("TC-203", "Page refresh keeps authenticated session",  "/QALS1011/KanbanItem/List",    "Yes", "Still authenticated",        "PASS"),
    ("TC-204", "All vertical-nav links have href",          "/QALS1011/KanbanItem/List",    "Yes", "No links without href",      "PASS"),
    ("TC-205", "Invalid route shows 404/error page",        "/qals1011/invalid_xyz9999",    "Yes", "Error indicator shown",      "PASS"),
    ("TC-206", "Home landing page loads",                   "/QALS1011/Home/LandingHome",   "Yes", "Page loads successfully",    "PASS"),
    ("TC-207", "Material Status page loads",                "/QALS1011/MaterialStatus/Index","Yes","Page loads successfully",    "PASS"),
    ("TC-208", "Report home page loads",                    "/QALS1011/Report/Home",        "Yes", "Page loads successfully",    "PASS"),
    ("TC-209", "Admin home page loads",                     "/QALS1011/Admin/Home",         "Yes", "Page loads successfully",    "PASS"),
    ("TC-210", "Forgot Password link on login page",        "/qals1011",                    "No",  "Link present and clickable", "PASS"),
]
for ri, row in enumerate(nav_rows, 2):
    write_row(ws3, ri, row, PASS_BG if row[5] == "PASS" else FAIL_BG)

# ── Sheet 4 : Edge Case / Negative Test Data ──────────────────────────────────
ws4 = wb.create_sheet("EdgeCase_TestData")
write_header(ws4,
    ["TC_ID", "Description", "Input_Field", "Input_Value", "Expected_Result", "Expected_Status"],
    [10, 50, 18, 40, 35, 16])

edge_rows = [
    ("TC-301", "XSS attempt in username field",          "Username", "<script>alert(1)</script>", "Login rejected / sanitized", "PASS"),
    ("TC-302", "Very long username (200 chars)",         "Username", "a" * 60 + "...(200 total)", "Login rejected",             "PASS"),
    ("TC-303", "Very long password (200 chars)",         "Password", "a" * 60 + "...(200 total)", "Login rejected",             "PASS"),
    ("TC-304", "Special characters in username",         "Username", "!@#$%^&*()",                "Login rejected",             "PASS"),
    ("TC-305", "Whitespace-only username",               "Username", "(spaces only)",             "Login rejected",             "PASS"),
    ("TC-306", "Whitespace-only password",               "Password", "(spaces only)",             "Login rejected",             "PASS"),
    ("TC-307", "Tab character in username",              "Username", "buyer[TAB]user",            "Login rejected",             "PASS"),
    ("TC-308", "Null-byte in username",                  "Username", "buyeruser[NULL]",           "Login rejected / sanitized", "PASS"),
    ("TC-309", "Password with embedded space",           "Password", "Ultriva !23",               "Login rejected",             "PASS"),
    ("TC-310", "Case-sensitive password (all lowercase)","Password", "ultriva!23",                "Login rejected",             "PASS"),
]
for ri, row in enumerate(edge_rows, 2):
    write_row(ws4, ri, row, PASS_BG if row[5] == "PASS" else FAIL_BG)

# ── Sheet 5 : Test Summary ────────────────────────────────────────────────────
ws5 = wb.create_sheet("Test_Summary")
ws5.column_dimensions["A"].width = 28
ws5.column_dimensions["B"].width = 12
ws5.column_dimensions["C"].width = 42

t = ws5.cell(row=1, column=1, value="QA Test Data Summary — Ultriva QALS1011")
t.font = Font(bold=True, size=14, color=HEADER_BG, name="Calibri")
t.alignment = center()
ws5.merge_cells("A1:C1")
ws5.row_dimensions[1].height = 30

for ci, hdr in enumerate(["Module", "# Test Cases", "Coverage Area"], 1):
    c = ws5.cell(row=2, column=ci, value=hdr)
    c.font = bold_font(); c.fill = fill(SECTION_BG)
    c.alignment = center(); c.border = border

summary_data = [
    ("Authentication (Login)",  len(login_rows),  "Login/logout, valid/invalid credentials, edge cases"),
    ("Dashboard",               len(dash_rows),   "Post-login navigation, content, session, toolbar"),
    ("Navigation",              len(nav_rows),     "URL routing, back/forward, 404, direct page access"),
    ("Edge Cases",              len(edge_rows),    "Injection, long inputs, special chars, case-sensitivity"),
]
for ri, (mod, cnt, cov) in enumerate(summary_data, 3):
    ws5.cell(row=ri, column=1, value=mod).font = cell_font()
    ws5.cell(row=ri, column=2, value=cnt).font = cell_font()
    ws5.cell(row=ri, column=3, value=cov).font = cell_font()
    for ci in range(1, 4):
        ws5.cell(row=ri, column=ci).fill = fill(ALT_BG if ri % 2 == 0 else "FFFFFF")
        ws5.cell(row=ri, column=ci).alignment = center() if ci == 2 else left()
        ws5.cell(row=ri, column=ci).border = border

total_ri = 3 + len(summary_data)
total_cnt = sum(x[1] for x in summary_data)
for ci, val in enumerate([("TOTAL", left()), (total_cnt, center()), ("All functional test data", left())], 1):
    c = ws5.cell(row=total_ri, column=ci, value=val[0])
    c.font = bold_font(); c.fill = fill(SECTION_BG)
    c.alignment = val[1]; c.border = border

for ri, (k, v) in enumerate([
    ("Application URL",  "https://engineering.ultrivalms.com/qals1011"),
    ("Test User",        "buyeruser"),
    ("Created On",       time.strftime("%Y-%m-%d %H:%M:%S")),
], total_ri + 2):
    ws5.cell(row=ri, column=1, value=k).font = bold_font()
    ws5.cell(row=ri, column=2, value=v).font = cell_font()
    ws5.merge_cells(f"B{ri}:C{ri}")
    for ci in range(1, 4):
        ws5.cell(row=ri, column=ci).border = border

os.makedirs("test_data", exist_ok=True)
wb.save("test_data/test_data.xlsx")
print("Saved: test_data/test_data.xlsx")
print()
for ws in wb.worksheets:
    nrows = sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if r[0])
    print(f"  {ws.title:<28} {nrows:>3} rows")
print(f"\n  TOTAL test data rows: {total_cnt}")
