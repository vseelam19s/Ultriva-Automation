"""
ExcelDataReader — reads rows from test_data/test_data.xlsx and returns
them as named-tuple lists so test files can import them cleanly.
"""
import os
from collections import namedtuple
from openpyxl import load_workbook

TEST_DATA_PATH = os.path.join(os.path.dirname(__file__), "..", "test_data", "test_data.xlsx")
ADMIN_DATA_PATH = os.path.join(os.path.dirname(__file__), "..", "test_data", "admin_test_data.xlsx")


def _sheet_rows(sheet_name: str, filepath: str = None):
    """Return a list of namedtuples (one per data row) for *sheet_name*."""
    path = filepath if filepath else TEST_DATA_PATH
    wb  = load_workbook(path, read_only=True, data_only=True)
    ws  = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return []

    # Build a namedtuple type from header row, replacing spaces with underscores
    raw_headers = [str(h).strip().replace(" ", "_").replace("(", "").replace(")", "").replace("/", "_")
                   for h in rows[0]]
    RowType = namedtuple(sheet_name.replace(" ", "_"), raw_headers)
    return [RowType(*r) for r in rows[1:] if r[0]]  # skip blank rows


def login_test_data():
    """Rows from Login_TestData sheet."""
    return _sheet_rows("Login_TestData")


def dashboard_test_data():
    """Rows from Dashboard_TestData sheet."""
    return _sheet_rows("Dashboard_TestData")


def navigation_test_data():
    """Rows from Navigation_TestData sheet."""
    return _sheet_rows("Navigation_TestData")


def edge_case_test_data():
    """Rows from EdgeCase_TestData sheet."""
    return _sheet_rows("EdgeCase_TestData")


if __name__ == "__main__":
    for fn, label in [
        (login_test_data,      "Login"),
        (dashboard_test_data,  "Dashboard"),
        (navigation_test_data, "Navigation"),
        (edge_case_test_data,  "Edge Cases"),
    ]:
        rows = fn()
        print(f"\n{label} ({len(rows)} rows):")
        for r in rows:
            print(f"  {r.TC_ID:<10} {r.Description[:55]}")
