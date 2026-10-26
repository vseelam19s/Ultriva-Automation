"""
QALS11 — Shared browser fixture + result reporting
URL  : https://engineering.ultrivalms.com/QALS1011
Scope: module  →  one browser per test file, kept open across tests
"""

import time
import sys
import pytest
from selenium import webdriver

# ── In-memory result store used by hooks ────────────────────────────────
_results: list[dict] = []   # {"name", "status", "duration", "error"}

# ═══════════════════════════════════════════════════════════════════════════════
# Browser fixture
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.fixture(scope="module")
def driver():
    """Launch Chrome (visible), shared across all tests in the module."""
    options = webdriver.ChromeOptions()
    options.add_argument("--start-maximized")
    options.add_experimental_option("excludeSwitches", ["enable-logging"])

    browser = webdriver.Chrome(options=options)
    browser.set_page_load_timeout(30)

    yield browser   # ← tests run here

    try:
        browser.quit()
    except Exception:
        pass

# ═══════════════════════════════════════════════════════════════════════════════
# Hooks — capture every test outcome
# ═══════════════════════════════════════════════════════════════════════════════


def pytest_runtest_logreport(report):
    """Called after each test phase (setup / call / teardown)."""
    if report.when != "call":
        return                          # only record the actual test call

    status = "PASS" if report.passed else ("FAIL" if report.failed else "SKIP")
    error = ""
    if report.failed and report.longrepr:
        error = str(report.longrepr).splitlines()[-1][:200]

    _results.append({
        "name": report.nodeid.split("::")[-1],
        "module": report.nodeid.split("::")[-2] if "::" in report.nodeid else "-",
        "status": status,
        "duration": round(getattr(report, "duration", 0.0), 2),
        "error": error,
    })

# ═══════════════════════════════════════════════════════════════════════════════
# Terminal summary — printed at the end of the run
# ═══════════════════════════════════════════════════════════════════════════════


def pytest_terminal_summary(terminalreporter, exitstatus, config):
    """Print a formatted summary table and save an Excel report."""
    if not _results:
        return

    passed = sum(1 for r in _results if r["status"] == "PASS")
    failed = sum(1 for r in _results if r["status"] == "FAIL")
    skipped = sum(1 for r in _results if r["status"] == "SKIP")
    total = len(_results)

    # ── ANSI colours (work in most terminals) ─────────────────────────────
    GREEN = "\033[92m"
    RED = "\033[91m"
    YELLOW = "\033[93m"
    CYAN = "\033[96m"
    BOLD = "\033[1m"
    RESET = "\033[0m"

    def colour(text, status):
        c = {"PASS": GREEN, "FAIL": RED, "SKIP": YELLOW}.get(status, "")
        return f"{c}{text}{RESET}"

    w_name = max(len(r["name"]) for r in _results) + 2
    sep = "─" * (w_name + 42)

    print(f"\n{BOLD}{CYAN}{'═' * 60}{RESET}")
    print(f"{BOLD}{CYAN}  QALS11 – Test Summary Report{RESET}")
    print(f"{BOLD}{CYAN}{'═' * 60}{RESET}")
    print(
        "  " + f"{'#':<4} {'Test Case':<{w_name}} {'Status':<8} {'Duration':>10}  Error")
    print(f"  {sep}")

    for i, r in enumerate(_results, 1):
        status_col = colour(r["status"], r["status"])
        err_col = (RED + r["error"][:60] + RESET) if r["error"] else ""
        print(
            "  " + f"{i:<4} {r['name']:<{w_name}} {status_col:<8} {r['duration']:>8.2f}s  {err_col}")

    print(f"  {sep}")
    print(
        f"  {BOLD}Total: {total}  |  "
        f"{GREEN}Passed: {passed}{RESET}{BOLD}  |  "
        f"{RED}Failed: {failed}{RESET}{BOLD}  |  "
        f"{YELLOW}Skipped: {skipped}{RESET}"
    )
    print(f"{BOLD}{CYAN}{'═' * 60}{RESET}\n")

    # ── Save Excel report ─────────────────────────────────────────────────
    _save_excel_report(passed, failed, skipped)

# ── Excel report helper ─────────────────────────────────────────────────


def _save_excel_report(passed, failed, skipped):
    try:
        import os
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        def fill(hex_color):
            return PatternFill(
                start_color=hex_color,
                end_color=hex_color,
                fill_type="solid")

        wb = Workbook()

        # ── Summary sheet ──────────────────────────────────────────────────
        ws_sum = wb.active
        ws_sum.title = "Summary"
        ws_sum.column_dimensions["A"].width = 22
        ws_sum.column_dimensions["B"].width = 18

        hdr = ws_sum.cell(1, 1, "QALS11 — Edit User Test Report")
        hdr.font = Font(bold=True, size=13, color="1F4E79")
        hdr.alignment = Alignment(horizontal="center")
        ws_sum.merge_cells("A1:B1")

        meta = [
            ("Run Date", time.strftime("%Y-%m-%d %H:%M:%S")),
            ("Total", passed + failed + skipped),
            ("Passed", passed),
            ("Failed", failed),
            ("Skipped", skipped),
        ]
        STATUS_COLORS = {
            "Passed": "92D050",
            "Failed": "FF0000",
            "Skipped": "FFC000"}
        for row_i, (k, v) in enumerate(meta, start=2):
            c1 = ws_sum.cell(row_i, 1, k)
            c2 = ws_sum.cell(row_i, 2, v)
            c1.font = Font(bold=True)
            if k in STATUS_COLORS:
                c2.fill = fill(STATUS_COLORS[k])
            for c in (c1, c2):
                c.border = border
                c.alignment = Alignment(horizontal="center")

        # ── Results sheet ──────────────────────────────────────────────────
        ws_res = wb.create_sheet("Test Results")
        headers = [
            "#",
            "Test Case",
            "Module",
            "Status",
            "Duration (s)",
            "Error"]
        widths = [5, 40, 30, 10, 14, 60]
        for col_i, (h, w) in enumerate(zip(headers, widths), 1):
            c = ws_res.cell(1, col_i, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = fill("1F4E79")
            c.alignment = Alignment(horizontal="center")
            c.border = border
            ws_res.column_dimensions[
                __import__("openpyxl.utils", fromlist=["get_column_letter"]).get_column_letter(col_i)
            ].width = w

        STATUS_FILL = {"PASS": "92D050", "FAIL": "FF0000", "SKIP": "FFC000"}
        for row_i, r in enumerate(_results, start=2):
            vals = [
                row_i - 1,
                r["name"],
                r["module"],
                r["status"],
                r["duration"],
                r["error"]]
            for col_i, v in enumerate(vals, 1):
                c = ws_res.cell(row_i, col_i, v)
                c.border = border
                c.alignment = Alignment(
                    horizontal="center" if col_i in (
                        1, 4, 5) else "left", wrap_text=True)
                if col_i == 4:
                    c.fill = fill(STATUS_FILL.get(r["status"], "FFFFFF"))

        report_dir = os.path.join(os.path.dirname(
            __file__), "..", "..", "..", "reports")
        os.makedirs(report_dir, exist_ok=True)
        path = os.path.join(report_dir, "qals11_edit_user_report.xlsx")
        wb.save(path)
        print(f"  [Excel] Report saved → {os.path.abspath(path)}\n")
    except Exception as exc:
        print(f"  [Excel] Could not save report: {exc}")
