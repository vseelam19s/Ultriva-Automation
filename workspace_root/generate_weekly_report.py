"""Generate a weekly QA test case DOCX for This Week vs Last Week (no dates shown)."""

import ast
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - optional fallback when workbook parsing is unavailable
    load_workbook = None


ROOT = Path(r"c:\Prcatices-AI")
TEST_ROOT = ROOT / "qa_automation" / "QASite" / "QALS11"
REPORTS_ROOT = ROOT / "qa_automation" / "reports"
SCREENSHOTS_ROOT = ROOT / "qa_automation" / "tests" / "reports" / "screenshots"
EXECUTION_REPORT_PATH = REPORTS_ROOT / "test_results.xlsx"
OUTPUT_BASENAME = "QA_TestCases_Weekly_Report"

HEADER_BLUE = "1F497D"
THIS_WEEK_BG = "D9EAD3"
LAST_WEEK_BG = "FCE5CD"
UNKNOWN_BG = "E7E6E6"


@dataclass
class TestRecord:
    week: str
    scope: str
    module: str
    suite: str
    testcase: str
    status: str
    source_file: str


def set_cell_bg(cell, hex_color: str):
    """Set background colour of a table cell."""
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def write_cell(
    cell,
    text: str,
    *,
    bold: bool = False,
    size: float = 9.5,
    color: str | None = None,
    align=WD_ALIGN_PARAGRAPH.LEFT,
):
    """Write a formatted value into a cell."""
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def status_fill(status: str) -> str:
    status = status.upper()
    if status == "PASS":
        return "C6E0B4"
    if status == "FAIL":
        return "F4CCCC"
    return UNKNOWN_BG


def section_heading(document: Document, text: str):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(HEADER_BLUE)


def prettify_test_name(name: str) -> str:
    """Convert function name to a readable test case label."""
    parts = name.split("_")
    start = 2 if len(parts) > 2 and parts[1].isdigit() else 1
    return " ".join(p.capitalize() for p in parts[start:])


def normalize_text(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def split_tokens(value: str) -> list[str]:
    return [token for token in re.split(r"[^a-z0-9]+", value.lower()) if token]


def file_slug_from_path(relative_path: Path) -> str:
    stem = relative_path.stem
    stem = stem.removeprefix("test_")
    stem = stem.removesuffix("_report")
    return stem


def get_week_windows(now: datetime):
    """Return datetime windows for this week and last week."""
    this_week_start = datetime(now.year, now.month, now.day) - timedelta(days=now.weekday())
    last_week_start = this_week_start - timedelta(days=7)
    next_week_start = this_week_start + timedelta(days=7)
    return {
        "This Week": (this_week_start, next_week_start),
        "Last Week": (last_week_start, this_week_start),
    }


def week_from_modified(modified: datetime, windows: dict[str, tuple[datetime, datetime]]):
    """Map file modified time to week label, or None when outside range."""
    for label, (start, end) in windows.items():
        if start <= modified < end:
            return label
    return None


def scope_and_module_from_path(relative_path: Path):
    """Derive scope and module labels from test file path."""
    parts = relative_path.parts
    scope = parts[0] if parts else "Other"
    parent = relative_path.parent.name
    stem = relative_path.stem.replace("test_", "")
    module = " ".join(segment.capitalize() for segment in stem.split("_"))
    if parent not in ("", ".") and parent != scope:
        module = f"{parent} - {module}"
    return scope, module


def extract_suite_cases(file_path: Path):
    """Return list of (suite_name, function_name, case_name) tuples from a test file."""
    try:
        source = file_path.read_text(encoding="utf-8")
        tree = ast.parse(source)
    except (OSError, UnicodeDecodeError, SyntaxError):
        return []

    records = []
    for node in tree.body:
        if isinstance(node, ast.ClassDef):
            suite = node.name
            funcs = [
                n.name
                for n in node.body
                if isinstance(n, ast.FunctionDef) and n.name.startswith("test_")
            ]
            for func_name in funcs:
                records.append((suite, func_name, prettify_test_name(func_name)))

    # Fallback for function-based tests (no class)
    if not records:
        for node in tree.body:
            if isinstance(node, ast.FunctionDef) and node.name.startswith("test_"):
                records.append((file_path.stem, node.name, prettify_test_name(node.name)))

    return records


def load_execution_results():
    """Load PASS/FAIL results from the latest execution workbook when available."""
    if load_workbook is None or not EXECUTION_REPORT_PATH.exists():
        return {}, {}

    try:
        workbook = load_workbook(EXECUTION_REPORT_PATH, read_only=True, data_only=True)
    except Exception:
        return {}, {}

    if "Test Results" not in workbook.sheetnames:
        return {}, {}

    worksheet = workbook["Test Results"]
    case_status = {}
    module_case_status = {}

    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 4:
            continue
        test_case = row[1]
        module = row[2]
        status = row[3]
        if not test_case or not status:
            continue

        normalized_case = normalize_text(str(test_case))
        normalized_module = normalize_text(str(module)) if module else ""
        normalized_status = str(status).upper()
        case_status[normalized_case] = normalized_status
        if normalized_module:
            module_case_status[(normalized_module, normalized_case)] = normalized_status

    return case_status, module_case_status


def resolve_status_from_screenshots(file_slug: str, function_name: str, testcase: str) -> str | None:
    """Infer PASS/FAIL from screenshot artifacts produced for the test file."""
    if not SCREENSHOTS_ROOT.exists():
        return None

    slug_tokens = split_tokens(file_slug)
    function_tokens = split_tokens(function_name.removeprefix("test_"))
    case_tokens = split_tokens(testcase)

    best_match = None
    best_score = 0

    for screenshot in SCREENSHOTS_ROOT.glob("*.png"):
        name = screenshot.name.lower()
        if slug_tokens and not all(token in name for token in slug_tokens):
            continue

        score = 0
        for token in slug_tokens:
            if token in name:
                score += 8
        for token in function_tokens:
            if token in name:
                score += 4
        for token in case_tokens:
            if token in name:
                score += 2
        if "fail" in name:
            score += 1

        if score > best_score or (
            score == best_score and best_match and screenshot.stat().st_mtime > best_match.stat().st_mtime
        ):
            if score > 0:
                best_match = screenshot
                best_score = score

    if not best_match:
        return None

    return "FAIL" if "fail" in best_match.name.lower() else "PASS"


def resolve_status(
    relative_path: Path,
    file_slug: str,
    suite: str,
    function_name: str,
    testcase: str,
    case_status: dict[str, str],
    module_case_status: dict[tuple[str, str], str],
) -> tuple[str, str]:
    """Resolve a case status from execution artifacts, with a visible fallback."""
    normalized_case = normalize_text(testcase)
    normalized_suite = normalize_text(suite)

    status = module_case_status.get((normalized_suite, normalized_case))
    if not status:
        status = case_status.get(normalized_case)
    if not status:
        status = resolve_status_from_screenshots(file_slug, function_name, testcase)
    if not status:
        status = "UNKNOWN"

    return status, relative_path.as_posix()


def collect_records():
    """Collect test case records for this week and last week."""
    now = datetime.now()
    windows = get_week_windows(now)
    case_status, module_case_status = load_execution_results()
    collected: list[TestRecord] = []

    for file_path in TEST_ROOT.rglob("test_*.py"):
        modified = datetime.fromtimestamp(file_path.stat().st_mtime)
        week = week_from_modified(modified, windows)
        if not week:
            continue

        relative = file_path.relative_to(TEST_ROOT)
        scope, module = scope_and_module_from_path(relative)
        file_slug = file_slug_from_path(relative)
        suite_cases = extract_suite_cases(file_path)
        for suite, function_name, case in suite_cases:
            status, source_file = resolve_status(
                relative,
                file_slug,
                suite,
                function_name,
                case,
                case_status,
                module_case_status,
            )
            collected.append(
                TestRecord(
                    week=week,
                    scope=scope,
                    module=module,
                    suite=suite,
                    testcase=case,
                    status=status,
                    source_file=source_file,
                )
            )

    # Keep output stable and readable
    collected.sort(key=lambda r: (r.week, r.scope, r.module, r.suite, r.testcase))
    return collected


def summarize(records: list[TestRecord]):
    """Return week-wise summary counts of suites and test cases."""
    summary = {
        "This Week": {"suites": set(), "cases": 0},
        "Last Week": {"suites": set(), "cases": 0},
    }
    for rec in records:
        summary[rec.week]["suites"].add((rec.scope, rec.module, rec.suite))
        summary[rec.week]["cases"] += 1
    return {
        "This Week": {
            "suites": len(summary["This Week"]["suites"]),
            "cases": summary["This Week"]["cases"],
        },
        "Last Week": {
            "suites": len(summary["Last Week"]["suites"]),
            "cases": summary["Last Week"]["cases"],
        },
    }


def build_document(records: list[TestRecord]):
    """Build the weekly report DOCX document."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    # Title block (no dates)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("QA Test Cases Weekly Report")
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor.from_string(HEADER_BLUE)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Weekly comparison of test cases created across This Week and Last Week")
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    document.add_paragraph()

    # Summary section
    section_heading(document, "Summary")
    totals = summarize(records)

    summary_table = document.add_table(rows=1, cols=3)
    summary_table.style = "Table Grid"
    summary_table.alignment = WD_TABLE_ALIGNMENT.LEFT

    headers = ["Scope", "Test Suites", "Total Test Cases"]
    for idx, text in enumerate(headers):
        write_cell(
            summary_table.cell(0, idx),
            text,
            bold=True,
            size=10,
            color="FFFFFF",
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_bg(summary_table.cell(0, idx), HEADER_BLUE)

    summary_rows = [
        ("This Week", totals["This Week"]["suites"], totals["This Week"]["cases"], THIS_WEEK_BG),
        ("Last Week", totals["Last Week"]["suites"], totals["Last Week"]["cases"], LAST_WEEK_BG),
    ]

    for label, suites, cases, bg in summary_rows:
        row = summary_table.add_row()
        write_cell(row.cells[0], label, bold=True)
        write_cell(row.cells[1], str(suites), align=WD_ALIGN_PARAGRAPH.CENTER)
        write_cell(row.cells[2], str(cases), align=WD_ALIGN_PARAGRAPH.CENTER)
        for cell in row.cells:
            set_cell_bg(cell, bg)

    document.add_paragraph()

    # Detail section
    section_heading(document, "Detailed Test Cases")
    detail_table = document.add_table(rows=1, cols=6)
    detail_table.style = "Table Grid"
    detail_table.alignment = WD_TABLE_ALIGNMENT.LEFT

    detail_headers = ["#", "Scope", "Module", "Test Suite", "Test Case", "Pass/Fail"]
    for idx, text in enumerate(detail_headers):
        write_cell(
            detail_table.cell(0, idx),
            text,
            bold=True,
            size=10,
            color="FFFFFF",
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_bg(detail_table.cell(0, idx), HEADER_BLUE)

    # This Week first, then Last Week
    ordered_records = [
        *[r for r in records if r.week == "This Week"],
        *[r for r in records if r.week == "Last Week"],
    ]

    for idx, rec in enumerate(ordered_records, start=1):
        row = detail_table.add_row()
        bg = THIS_WEEK_BG if rec.week == "This Week" else LAST_WEEK_BG
        write_cell(row.cells[0], str(idx), align=WD_ALIGN_PARAGRAPH.CENTER)
        write_cell(row.cells[1], rec.scope)
        write_cell(row.cells[2], rec.module)
        write_cell(row.cells[3], rec.suite)
        write_cell(row.cells[4], rec.testcase)
        write_cell(row.cells[5], rec.status, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        for cell in row.cells:
            set_cell_bg(cell, bg)
        set_cell_bg(row.cells[5], status_fill(rec.status))

    return document, totals, len(ordered_records)


def main():
    records = collect_records()
    REPORTS_ROOT.mkdir(parents=True, exist_ok=True)
    document, totals, total_cases = build_document(records)
    output_path = REPORTS_ROOT / f"{OUTPUT_BASENAME}_{datetime.now():%Y%m%d_%H%M%S}.docx"
    document.save(output_path)

    pass_count = sum(1 for record in records if record.status == "PASS")
    fail_count = sum(1 for record in records if record.status == "FAIL")
    unknown_count = sum(1 for record in records if record.status not in {"PASS", "FAIL"})

    print(f"Report saved to: {output_path}")
    print(
        "This Week -> suites: "
        f"{totals['This Week']['suites']}, cases: {totals['This Week']['cases']}"
    )
    print(
        "Last Week -> suites: "
        f"{totals['Last Week']['suites']}, cases: {totals['Last Week']['cases']}"
    )
    print(f"Combined cases listed: {total_cases}")
    print(f"Status summary -> PASS: {pass_count}, FAIL: {fail_count}, UNKNOWN: {unknown_count}")


if __name__ == "__main__":
    main()
