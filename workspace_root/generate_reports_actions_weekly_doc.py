import ast
from collections import OrderedDict
from datetime import datetime, timedelta
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORKSPACE_ROOT = Path(r"c:\Prcatices-AI")
TEST_ROOT = WORKSPACE_ROOT / "qa_automation" / "QASite" / "QALS11"
OUTPUT_PATH = Path(r"c:\Upland\Weekly_TestCases_Report.docx")
FALLBACK_OUTPUT_PATH = Path(r"c:\Upland\Weekly_TestCases_Report_Reports_Actions.docx")

HEADER_BLUE = "1F497D"
THIS_WEEK_BG = "D9EAD3"
LAST_WEEK_BG = "FCE5CD"


def set_cell_bg(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def write_cell(cell, text: str, *, bold: bool = False, font_size: float = 9.5, color: str | None = None, center: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def prettify_case(name: str) -> str:
    parts = name.split("_")
    start = 2 if len(parts) > 2 and parts[1].isdigit() else 1
    return " ".join(part.capitalize() for part in parts[start:])


def prettify_module(file_name: str) -> str:
    name = file_name.replace("test_", "").replace("_report", "")
    return " ".join(part.capitalize() for part in name.split("_"))


def docstring_title(tree: ast.AST) -> str | None:
    text = ast.get_docstring(tree)
    if not text:
        return None

    first_line = text.strip().splitlines()[0].strip()
    if ">" in first_line:
        first_line = first_line.split(">")[-1].strip()
    first_line = first_line.replace("-", " ")
    first_line = re.sub(r"\s+", " ", first_line).strip()
    return first_line or None


def extract_suites(file_path: Path) -> tuple[str, OrderedDict[str, list[str]]]:
    try:
        source = file_path.read_text(encoding="utf-8")
        tree = ast.parse(source)
    except (OSError, SyntaxError, UnicodeDecodeError):
        return prettify_module(file_path.stem), OrderedDict()

    module_name = docstring_title(tree) or prettify_module(file_path.stem)
    suites: OrderedDict[str, list[str]] = OrderedDict()
    for node in tree.body:
        if isinstance(node, ast.ClassDef):
            test_names = [
                child.name
                for child in node.body
                if isinstance(child, ast.FunctionDef) and child.name.startswith("test_")
            ]
            if test_names:
                suites[node.name] = test_names
    return module_name, suites


def week_windows() -> dict[str, tuple[datetime, datetime]]:
    today = datetime.now()
    this_week_start = datetime(today.year, today.month, today.day) - timedelta(days=today.weekday())
    last_week_start = this_week_start - timedelta(days=7)
    return {
        "This Week": (this_week_start, this_week_start + timedelta(days=7)),
        "Last Week": (last_week_start, this_week_start),
    }


def collect_entries() -> dict[str, list[dict[str, str]]]:
    windows = week_windows()
    selected = {
        "This Week": TEST_ROOT / "Reports",
        "Last Week": TEST_ROOT / "Actions",
    }
    entries: dict[str, list[dict[str, str]]] = {"This Week": [], "Last Week": []}

    for week_name, folder in selected.items():
        start, end = windows[week_name]
        for file_path in sorted(folder.glob("test_*.py")):
            modified = datetime.fromtimestamp(file_path.stat().st_mtime)
            if not (start <= modified < end):
                continue

            module_name, suites = extract_suites(file_path)
            for suite_name, test_names in suites.items():
                for test_name in test_names:
                    entries[week_name].append(
                        {
                            "module": module_name,
                            "suite": suite_name,
                            "case": prettify_case(test_name),
                        }
                    )
    return entries


def add_summary_table(document: Document, entries: dict[str, list[dict[str, str]]]) -> None:
    heading = document.add_paragraph()
    run = heading.add_run("Summary")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(HEADER_BLUE)

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    headers = ["Scope", "Test Suites", "Total Test Cases"]
    for index, header in enumerate(headers):
        write_cell(table.cell(0, index), header, bold=True, font_size=10, color="FFFFFF", center=True)
        set_cell_bg(table.cell(0, index), HEADER_BLUE)

    this_week_modules = {entry["module"] for entry in entries["This Week"]}
    last_week_modules = {entry["module"] for entry in entries["Last Week"]}
    summary_rows = [
        ("This Week - Reports", len(this_week_modules), len(entries["This Week"]), THIS_WEEK_BG),
        ("Last Week - Actions", len(last_week_modules), len(entries["Last Week"]), LAST_WEEK_BG),
    ]

    for scope, suites, cases, bg in summary_rows:
        row = table.add_row()
        write_cell(row.cells[0], scope, bold=True)
        write_cell(row.cells[1], str(suites), center=True)
        write_cell(row.cells[2], str(cases), center=True)
        for cell in row.cells:
            set_cell_bg(cell, bg)


def add_details_table(document: Document, entries: dict[str, list[dict[str, str]]]) -> None:
    document.add_paragraph()
    heading = document.add_paragraph()
    run = heading.add_run("Detailed Test Cases")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(HEADER_BLUE)

    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"

    headers = ["#", "Scope", "Module", "Test Suite", "Test Case"]
    widths = [0.4, 1.6, 1.9, 2.1, 3.3]
    for index, (header, width) in enumerate(zip(headers, widths)):
        write_cell(table.cell(0, index), header, bold=True, font_size=9.5, color="FFFFFF", center=True)
        table.columns[index].width = Inches(width)
        set_cell_bg(table.cell(0, index), HEADER_BLUE)

    row_number = 0
    for week_name in ("This Week", "Last Week"):
        bg = THIS_WEEK_BG if week_name == "This Week" else LAST_WEEK_BG
        scope = "Reports" if week_name == "This Week" else "Actions"
        for entry in entries[week_name]:
            row_number += 1
            row = table.add_row()
            write_cell(row.cells[0], str(row_number), center=True, font_size=9)
            write_cell(row.cells[1], scope, font_size=9)
            write_cell(row.cells[2], entry["module"], font_size=9)
            write_cell(row.cells[3], entry["suite"], font_size=9)
            write_cell(row.cells[4], entry["case"], font_size=9)
            for cell in row.cells:
                set_cell_bg(cell, bg)

    document.add_paragraph()
    footer = document.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(
        f"Total test cases reported: {row_number}  (This week - Reports: {len(entries['This Week'])}  |  Last week - Actions: {len(entries['Last Week'])})"
    )
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    run.bold = True


def build_report(entries: dict[str, list[dict[str, str]]]) -> Document:
    document = Document()
    for section in document.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("QA Test Cases Weekly Report")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(HEADER_BLUE)

    document.add_paragraph()
    add_summary_table(document, entries)
    document.add_paragraph()
    add_details_table(document, entries)
    return document


def main() -> None:
    entries = collect_entries()
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = build_report(entries)
    saved_path = OUTPUT_PATH
    try:
        document.save(OUTPUT_PATH)
    except PermissionError:
        saved_path = FALLBACK_OUTPUT_PATH
        document.save(saved_path)

    print(f"Report saved to: {saved_path}")
    print(f"This Week Reports: {len(entries['This Week'])}")
    print(f"Last Week Actions: {len(entries['Last Week'])}")


if __name__ == "__main__":
    main()