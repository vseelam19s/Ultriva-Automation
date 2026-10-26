from __future__ import annotations

import ast
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"c:\Prcatices-AI")
TEST_ROOT = ROOT / "qa_automation" / "QASite" / "QALS11"
OUTPUT_PATH = ROOT / "qa_automation" / "reports" / "QA_Automation_Test_Cases_Report.docx"

HEADER_BLUE = "1F4E79"
THIS_WEEK_BG = "D9EAD3"
LAST_WEEK_BG = "FCE5CD"


def set_cell_bg(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def write_cell(cell, text: str, *, bold: bool = False, size: float = 9.5, color: str | None = None, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(HEADER_BLUE)


def module_label(relative_path: Path) -> str:
    parts = relative_path.parts
    if not parts:
        return "General"
    top = parts[0]
    if top == "Admin" and len(parts) > 1:
        return f"Admin - {parts[1]}"
    if top in {"Actions", "Reports", "Supplier", "Buyer", "Invoices"}:
        return top
    return top


def prettify_name(name: str) -> str:
    cleaned = name.replace("test_", "").replace("_report", "")
    return " ".join(part.capitalize() for part in cleaned.split("_"))


def extract_test_count(file_path: Path) -> int:
    try:
        source = file_path.read_text(encoding="utf-8")
        tree = ast.parse(source)
    except (OSError, SyntaxError, UnicodeDecodeError):
        return 0
    return sum(1 for node in ast.walk(tree) if isinstance(node, ast.FunctionDef) and node.name.startswith("test_"))


def week_windows(today: datetime) -> dict[str, tuple[datetime, datetime]]:
    this_week_start = datetime(today.year, today.month, today.day) - timedelta(days=today.weekday())
    last_week_start = this_week_start - timedelta(days=7)
    next_week_start = this_week_start + timedelta(days=7)
    return {
        "This Week": (this_week_start, next_week_start),
        "Last Week": (last_week_start, this_week_start),
    }


def collect_recent_tests() -> dict[str, dict[str, dict[str, object]]]:
    windows = week_windows(datetime.now())
    grouped: dict[str, dict[str, dict[str, object]]] = {
        "This Week": defaultdict(lambda: {"files": [], "cases": 0}),
        "Last Week": defaultdict(lambda: {"files": [], "cases": 0}),
    }

    for file_path in TEST_ROOT.rglob("test_*.py"):
        modified = datetime.fromtimestamp(file_path.stat().st_mtime)
        relative = file_path.relative_to(TEST_ROOT)
        matched_week = None
        for week_name, (start, end) in windows.items():
            if start <= modified < end:
                matched_week = week_name
                break
        if not matched_week:
            continue

        label = module_label(relative)
        friendly_name = prettify_name(file_path.stem)
        test_count = extract_test_count(file_path)
        if test_count == 0:
            continue

        grouped[matched_week][label]["files"].append(friendly_name)
        grouped[matched_week][label]["cases"] += test_count

    return grouped


def summarize_week(week_data: dict[str, dict[str, object]]) -> tuple[int, int]:
    suite_count = sum(len(data["files"]) for data in week_data.values())
    case_count = sum(int(data["cases"]) for data in week_data.values())
    return suite_count, case_count


def coverage_text(names: list[str], limit: int = 8) -> str:
    ordered = sorted(names)
    if len(ordered) <= limit:
        return ", ".join(ordered)
    shown = ", ".join(ordered[:limit])
    return f"{shown}, +{len(ordered) - limit} more"


def add_summary_table(document: Document, grouped: dict[str, dict[str, dict[str, object]]]) -> None:
    section_heading(document, "Summary")
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    headers = ["Period", "Automation Files", "Test Cases"]
    for index, header in enumerate(headers):
        write_cell(table.cell(0, index), header, bold=True, size=10, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_bg(table.cell(0, index), HEADER_BLUE)

    for week_name in ("This Week", "Last Week"):
        suites, cases = summarize_week(grouped[week_name])
        row = table.add_row()
        write_cell(row.cells[0], week_name, bold=True)
        write_cell(row.cells[1], str(suites), align=WD_ALIGN_PARAGRAPH.CENTER)
        write_cell(row.cells[2], str(cases), align=WD_ALIGN_PARAGRAPH.CENTER)
        bg = THIS_WEEK_BG if week_name == "This Week" else LAST_WEEK_BG
        for cell in row.cells:
            set_cell_bg(cell, bg)

    document.add_paragraph()


def add_week_section(document: Document, week_name: str, week_data: dict[str, dict[str, object]]) -> None:
    section_heading(document, week_name)
    suites, cases = summarize_week(week_data)

    intro = document.add_paragraph()
    intro_run = intro.add_run(f"Completed automation across {suites} files covering {cases} validation scenarios.")
    intro_run.font.size = Pt(10)

    if not week_data:
        note = document.add_paragraph()
        note.add_run("No new automation items were identified in this period.").italic = True
        return

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    headers = ["Area", "Files", "Scenarios", "Highlights"]
    for index, header in enumerate(headers):
        write_cell(table.cell(0, index), header, bold=True, size=10, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_bg(table.cell(0, index), HEADER_BLUE)

    bg = THIS_WEEK_BG if week_name == "This Week" else LAST_WEEK_BG
    for module_name in sorted(week_data):
        data = week_data[module_name]
        files = sorted(data["files"])
        row = table.add_row()
        write_cell(row.cells[0], module_name, bold=True)
        write_cell(row.cells[1], str(len(files)), align=WD_ALIGN_PARAGRAPH.CENTER)
        write_cell(row.cells[2], str(int(data["cases"])), align=WD_ALIGN_PARAGRAPH.CENTER)
        write_cell(row.cells[3], coverage_text(files))
        for cell in row.cells:
            set_cell_bg(cell, bg)

    document.add_paragraph()


def build_document(grouped: dict[str, dict[str, dict[str, object]]]) -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("QA Automation – Test Cases Report")
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor.from_string(HEADER_BLUE)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Summary of test cases covered in this week and last week")
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    document.add_paragraph()
    add_summary_table(document, grouped)
    add_week_section(document, "This Week", grouped["This Week"])
    add_week_section(document, "Last Week", grouped["Last Week"])

    closing = document.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    total_suites = sum(summarize_week(grouped[week])[0] for week in grouped)
    total_cases = sum(summarize_week(grouped[week])[1] for week in grouped)
    closing_run = closing.add_run(f"Combined total: {total_suites} automation files covering {total_cases} scenarios")
    closing_run.bold = True
    closing_run.font.size = Pt(9)
    closing_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    return document


def main() -> None:
    grouped = collect_recent_tests()
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = build_document(grouped)
    document.save(OUTPUT_PATH)

    this_week_suites, this_week_cases = summarize_week(grouped["This Week"])
    last_week_suites, last_week_cases = summarize_week(grouped["Last Week"])
    print(f"Saved: {OUTPUT_PATH}")
    print(f"This Week -> suites: {this_week_suites}, cases: {this_week_cases}")
    print(f"Last Week -> suites: {last_week_suites}, cases: {last_week_cases}")


if __name__ == "__main__":
    main()
