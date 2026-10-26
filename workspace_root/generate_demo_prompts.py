"""
QA Demo Prompts Document Generator
Organises all prompts (TestRail, JIRA, QA Automation) into a clean Word doc.
Run: python generate_demo_prompts.py
Output: QA_Demo_Prompts.docx
"""

from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT_DOC = r"c:\Prcatices-AI\QA_Demo_Prompts.docx"

# ── Style helpers ─────────────────────────────────────────────────────────────

def heading(doc, text, level=1, rgb="1F4E79"):
    h = doc.add_heading(text, level=level)
    run = h.runs[0] if h.runs else h.add_run(text)
    run.font.color.rgb = RGBColor.from_string(rgb)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    h.paragraph_format.space_after  = Pt(4)

def subhead(doc, text, rgb="2E75B6"):
    heading(doc, text, level=2, rgb=rgb)

def body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        run.font.size = Pt(10)

def badge(doc, text, color="1F4E79"):
    """Inline coloured label at start of a paragraph."""
    p = doc.add_paragraph()
    run = p.add_run(f"  {text}  ")
    run.font.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "EEF4FB")
    pPr.append(shd)
    p.paragraph_format.space_after = Pt(2)
    return p

def prompt_block(doc, label, steps):
    """Numbered list of steps under a named block."""
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.font.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor.from_string("C55A11")
    p.paragraph_format.space_after = Pt(2)
    for i, step in enumerate(steps, 1):
        s = doc.add_paragraph(style="List Number")
        s.add_run(step).font.size = Pt(9)
        s.paragraph_format.left_indent  = Inches(0.3)
        s.paragraph_format.space_after  = Pt(1)
    doc.add_paragraph()  # spacer

def info_table(doc, rows, headers):
    tbl = doc.add_table(rows=0, cols=len(headers))
    tbl.style = "Table Grid"
    hr = tbl.add_row()
    for i, h in enumerate(headers):
        c = hr.cells[i]
        c.text = h
        c.paragraphs[0].runs[0].font.bold  = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        c.paragraphs[0].runs[0].font.size  = Pt(9)
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"),"clear"); sh.set(qn("w:color"),"auto"); sh.set(qn("w:fill"),"1F4E79")
        c._tc.get_or_add_tcPr().append(sh)
    for j, row in enumerate(rows):
        dr = tbl.add_row()
        for i, val in enumerate(row):
            c = dr.cells[i]
            c.text = str(val)
            c.paragraphs[0].runs[0].font.size = Pt(9)
            if j % 2 == 0:
                sh = OxmlElement("w:shd")
                sh.set(qn("w:val"),"clear"); sh.set(qn("w:color"),"auto"); sh.set(qn("w:fill"),"EEF4FB")
                c._tc.get_or_add_tcPr().append(sh)
    doc.add_paragraph()

# ── Document builder ──────────────────────────────────────────────────────────

def build():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(2.5); sec.right_margin = Cm(2)

    stl = doc.styles["Normal"]
    stl.font.name = "Calibri"; stl.font.size = Pt(10)

    # ── Cover ────────────────────────────────────────────────────────────────
    doc.add_paragraph()
    t = doc.add_heading("QA Automation Demo — Prompt Reference", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in t.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79); r.font.size = Pt(26)

    sub = doc.add_paragraph("TestRail · JIRA · Selenium Automation — All Prompts Organised")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(14); sub.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    meta = doc.add_paragraph(
        f"Prepared by: QA Engineer  |  Date: {datetime.now().strftime('%B %d, %Y')}\n"
        "App: https://engineering.ultrivalms.com/QALS1011  |  TestRail: upland.testrail.com"
    )
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in meta.runs:
        r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    doc.add_page_break()

    # ── TOC ──────────────────────────────────────────────────────────────────
    heading(doc, "Table of Contents")
    for num, item in [
        ("1", "TestRail Prompts"),
        ("2", "JIRA Ticket Prompts"),
        ("3", "QA Automation — Admin → User"),
        ("4", "QA Automation — Admin → Item"),
        ("5", "QA Automation — Admin → Supplier"),
        ("6", "Postman API Automation"),
        ("7", "Quick Reference Table"),
    ]:
        p = doc.add_paragraph(f"  {num}.  {item}")
        p.paragraph_format.space_after = Pt(3)
        p.runs[0].font.size = Pt(10)
    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════════
    # SECTION 1 — TESTRAIL
    # ══════════════════════════════════════════════════════════════════════════
    heading(doc, "1.  TestRail Prompts", rgb="538135")
    body(doc,
        "All TestRail tasks are executed through VS Code Copilot Chat using the TestRail MCP Server. "
        "Files are under: qa_automation/Practise on test rail/"
    )

    subhead(doc, "1.1  Create Milestone", rgb="538135")
    prompt_block(doc, "File: Practise on test rail/01_create_milestone.py", [
        "Open VS Code → Copilot Chat → ensure TestRail MCP server is running.",
        "Prompt: '@testrail Add a new milestone named \"AuomationPractiseQA\" in project 15'",
        "MCP creates the milestone and returns the milestone ID.",
        "Verify in TestRail: Milestones tab → 'AuomationPractiseQA' appears.",
    ])

    subhead(doc, "1.2  Create Test Plan", rgb="538135")
    prompt_block(doc, "File: Practise on test rail/02_create_plan.py", [
        "Prompt: '@testrail Create a test plan named \"AuomationPractiseQA\" in project 15'",
        "Assign milestone: '@testrail Assign milestone \"AuomationPractiseQA\" to plan \"AuomationPractiseQA\"'",
        "MCP returns the test plan ID.",
        "Verify in TestRail: Test Plans tab → plan visible with milestone linked.",
    ])

    subhead(doc, "1.3  Add Admin Module S2571 to Test Plan", rgb="538135")
    prompt_block(doc, "File: Practise on test rail/02_create_plan.py  (continued)", [
        "Prompt: '@testrail Add suite S2571 (Admin module) to test plan \"AuomationPractiseQA\"'",
        "MCP tool: add_plan_entry with suite_id=2571 and plan_id obtained in step above.",
        "Confirm run entry created inside the plan for Admin module.",
    ])

    subhead(doc, "1.4  Add Test Case C5157814 to Admin Module", rgb="538135")
    prompt_block(doc, "File: Practise on test rail/03_pass_testcase.py", [
        "Prompt: '@testrail Add test case C5157814 to Admin module S2571 in plan \"AuomationPractiseQA\"'",
        "MCP includes case_id=5157814 in the plan entry run.",
        "Verify case appears under Admin module inside the plan.",
    ])

    subhead(doc, "1.5  Add Test Case for Bug ULT-10069", rgb="538135")
    prompt_block(doc, "File: JIRA Tickets/Bugs/ULT-10069_add_testrail_case.py", [
        "Prompt: '@testrail Get steps and expected results from reference case C13569701'",
        "MCP returns all Steps and Expected Results from the reference case.",
        "Prompt: '@testrail Create a new test case in section 372888 (Kanban Loop Design)'",
        "  Title: 'Stop Manage Background Service — KLD HTTP exchange'",
        "  Preconditions: 'Stop Manage Background Service'",
        "  Copy steps + expected results from C13569701.",
        "Verify new case ID is returned and visible in TestRail under S2571.",
    ])

    subhead(doc, "1.6  Add Test Cases for Enhancement ULT-10425", rgb="538135")
    prompt_block(doc, "File: JIRA Tickets/Enhancements/ULT-10425_add_testrail_cases.py", [
        "Prompt: '@testrail Add test cases for ULT-10425 under section S7352 → Home page changes'",
        "MCP creates cases: label rename, BU filter persistence, default filter, only-this-BU option.",
        "Verify all new case IDs appear under Enhancements → Home page changes.",
    ])

    info_table(doc,
        rows=[
            ("01_create_milestone.py", "Practise on test rail/", "Create milestone AuomationPractiseQA"),
            ("02_create_plan.py",      "Practise on test rail/", "Create plan + assign milestone + add S2571"),
            ("03_pass_testcase.py",    "Practise on test rail/", "Add C5157814, pass test case"),
            ("ULT-10069_add_testrail_case.py", "JIRA Tickets/Bugs/", "Copy C13569701 steps → new bug case"),
            ("ULT-10425_add_testrail_cases.py","JIRA Tickets/Enhancements/","Add home-page change cases to S7352"),
        ],
        headers=["File", "Folder", "Purpose"],
    )
    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════════
    # SECTION 2 — JIRA
    # ══════════════════════════════════════════════════════════════════════════
    heading(doc, "2.  JIRA Ticket Prompts", rgb="C55A11")
    body(doc,
        "JIRA access: https://uplandsoftware.atlassian.net  |  User: veseelam@uplandsoftware.com\n"
        "All files under: qa_automation/JIRA Tickets/"
    )

    subhead(doc, "2.1  Enhancement — ULT-10511", rgb="C55A11")
    prompt_block(doc, "Files: Enhancements/ ULT-10511_test_plan.md, _test_scenarios.py, _test_cases.py, _generate_excel.py", [
        "Access JIRA → UltrivaApp project → open ticket ULT-10511.",
        "Note scope: read ticket summary, description, acceptance criteria.",
        "Create test scenarios file with major scenario headings (no verbose comments).",
        "Create test plan file: scope, objectives, in/out of scope, entry/exit criteria.",
        "Generate test cases (short form): TC-01 to TC-05 covering major flows only.",
        "Run _generate_excel.py to export cases to Excel.",
        "Prompt (JIRA MCP): '@jira Create subtask under ULT-10511: Write automated test cases'",
        "Prompt: '@jira Fill subtask fields: title, test case IDs, TestRail suite reference'",
    ])

    subhead(doc, "2.2  Enhancement — ULT-10425", rgb="C55A11")
    prompt_block(doc, "Files: Enhancements/ ULT-10425_test_plan.md, _test_scenarios.py, _test_cases.py, _add_testrail_cases.py", [
        "Access JIRA → open ULT-10425 (Plant Home Page filter / label rename).",
        "QA site: https://engineering.ultrivalms.com/QALS1011  login: buyeruser / Ultriva!23",
        "Create test scenarios: label renamed, filter persists, default filter, Only-this-BU present.",
        "Create test plan: scope based on parent ticket + acceptance criteria.",
        "Generate test cases (short, no lengthy steps): TC-01 to TC-05.",
        "Prompt (JIRA MCP): '@jira Create subtask under ULT-10425: Write test cases' → fill mandatory fields.",
        "Prompt: '@testrail Add test cases under Enhancements S7352 → Home page changes'",
    ])

    subhead(doc, "2.3  Enhancement — ULT-9764", rgb="C55A11")
    prompt_block(doc, "Files: Enhancements/ ULT-9764_test_plan.md, _test_scenarios.py, _test_cases.py, _generate_excel.py", [
        "Access JIRA → open ULT-9764. Review attached screenshots.",
        "QA site: https://engineering.ultrivalms.com/QALS1011",
        "Create short test plan (scope, objectives, test case IDs only).",
        "Generate test cases based on screenshots — keep cases concise (5–8 max).",
        "Run _generate_excel.py to export cases.",
    ])

    subhead(doc, "2.4  Bug — ULT-10069", rgb="C55A11")
    prompt_block(doc, "Files: Bugs/ ULT-10069_test_plan.md, _test_scenarios.py, _test_cases.py, _add_testrail_case.py", [
        "Access JIRA → UltrivaApp → open ULT-10069 (Kanban Loop Design bug).",
        "Create test scenarios: KLD processes on HTTP exchange, batch job transitions, regression.",
        "Create test plan: scope, pre-conditions (Stop Background Service).",
        "Generate test cases short form — major flows only.",
        "Add new TestRail case: Admin module S2571 → sub-module 'Kanban Loop Design'.",
        "  Preconditions: 'Stop Manage Background Service'  (no manual steps — use reference).",
        "  Copy steps + expected results from reference C13569701.",
    ])

    info_table(doc,
        rows=[
            ("ULT-10511", "Enhancement", "Enhancements/", "Test plan, scenarios, cases, Excel export, JIRA subtask"),
            ("ULT-10425", "Enhancement", "Enhancements/", "Home page filter/label — plan, cases, TestRail add"),
            ("ULT-9764",  "Enhancement", "Enhancements/", "Short cases based on screenshots, Excel export"),
            ("ULT-10069", "Bug",         "Bugs/",         "KLD bug — scenarios, plan, cases, TestRail case copy"),
        ],
        headers=["Ticket", "Type", "Folder", "Files Generated"],
    )
    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════════
    # SECTION 3 — ADMIN → USER
    # ══════════════════════════════════════════════════════════════════════════
    heading(doc, "3.  QA Automation — Admin → User", rgb="1F4E79")
    body(doc, "Folder: qa_automation/tests/QALS11/Admin/  |  URL: https://engineering.ultrivalms.com/QALS1011")

    subhead(doc, "3.1  Create 5 Users")
    prompt_block(doc, "File: QALS11/Admin/test_create_users.py", [
        "Login with buyeruser / Ultriva!23.",
        "Click Admin module → User menu.",
        "Click Add button → Create New User form.",
        "Create 5 users: different User IDs, all with email vseelam@uplandsoftware.com.",
        "Account Type: Plant Administrators for all 5.",
        "Passwords: 8+ chars, 1 number, 1 uppercase, 1 lowercase, 1 punctuation (90-day expiry).",
        "Save each user — capture User ID on success.",
        "Export all 5 User IDs + passwords to Excel: QALS11/Reports/users_created.xlsx",
    ])

    subhead(doc, "3.2  Edit / Update User ID")
    prompt_block(doc, "File: QALS11/Admin/test_edit_userid.py", [
        "Login → Admin → User menu.",
        "Search for user ID PlantAdm01 in Users list.",
        "Click Edit → update mandatory fields with new values.",
        "Save and verify success toast message.",
        "Show summary report of test case result in UI.",
        "Screenshot saved to QALS11/Reports/screenshots/.",
    ])

    info_table(doc,
        rows=[
            ("test_create_users.py", "QALS11/Admin/", "Create 5 Plant Admins, export to Excel"),
            ("test_edit_userid.py",  "QALS11/Admin/", "Edit PlantAdm01 — update fields, summary report"),
        ],
        headers=["File", "Folder", "Purpose"],
    )
    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════════
    # SECTION 4 — ADMIN → ITEM
    # ══════════════════════════════════════════════════════════════════════════
    heading(doc, "4.  QA Automation — Admin → Item", rgb="1F4E79")
    body(doc, "Folder: qa_automation/tests/QALS11/Admin/Item/  |  URL: https://engineering.ultrivalms.com/QALS1011")

    subhead(doc, "4.1  Create New Item")
    prompt_block(doc, "File: QALS11/Admin/Item/test_create_item.py", [
        "Login → Admin → Item menu.",
        "Click Add → enter Item No: IT010426, Item Coordinator: buyeruser, other random values.",
        "Click Supplier tab → select Supplier: supnew, Supplier User: supnewuser, other randoms.",
        "Click Save → verify item created and appears in Items list.",
    ])

    subhead(doc, "4.2  Adjust Initial Loop Size (IT020426)")
    prompt_block(doc, "File: QALS11/Admin/Item/ItemInitialLoopSize.py", [
        "Login → Admin → Item menu → select item IT020426.",
        "Click Cards button → Click Add Cards button.",
        "Adjust Loop Size wizard: enter random Units On Hand and Units On Order.",
        "Click Next through all wizard steps → Finish.",
        "Verify loop size applied successfully.",
    ])

    subhead(doc, "4.3  Add Replenishment Cards (IT020426)")
    prompt_block(doc, "File: QALS11/Admin/Item/ItemAdjustLoopSize.py", [
        "Login → Admin → Item → select IT020426 → Cards.",
        "Click Add/Remove Cards → Select 'Add Replenishment Cards' radio → Next.",
        "Enter number of replenishment cards > 10.",
        "Select 'Release to Supplier' option → Finish.",
        "Verify cards added and loop size updated.",
    ])

    subhead(doc, "4.4  Add Non-Replenishment Cards (IT020426)")
    prompt_block(doc, "File: QALS11/Admin/Item/ItemAdjustLoopSize-NonReplinishmentcards.py", [
        "Login → Admin → Item → select IT020426 → Cards.",
        "Click Add/Remove Cards → Select 'Add Non-replenishment Cards' radio → Next.",
        "Enter number of non-replenishment cards > 10.",
        "Enter number of cycles: 5.",
        "Select 'Release to Supplier' option → Finish.",
        "Verify non-replenishment cards added.",
    ])

    info_table(doc,
        rows=[
            ("test_create_item.py",       "QALS11/Admin/Item/", "Create item IT010426 with supplier supnew"),
            ("ItemInitialLoopSize.py",    "QALS11/Admin/Item/", "Set initial loop size for IT020426"),
            ("ItemAdjustLoopSize.py",     "QALS11/Admin/Item/", "Add replenishment cards to IT020426"),
            ("ItemAdjustLoopSize-NonReplinishmentcards.py", "QALS11/Admin/Item/", "Add non-replenishment cards, 5 cycles"),
        ],
        headers=["File", "Folder", "Purpose"],
    )
    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════════
    # SECTION 5 — ADMIN → SUPPLIER
    # ══════════════════════════════════════════════════════════════════════════
    heading(doc, "5.  QA Automation — Admin → Supplier", rgb="1F4E79")
    body(doc, "Folder: qa_automation/tests/QALS11/Admin/Supplier/")

    subhead(doc, "5.1  Create New Supplier")
    prompt_block(doc, "File: QALS11/Admin/Supplier/test_create_supplier.py", [
        "Login → Admin → Supplier menu.",
        "Click Add → enter random mandatory fields.",
        "Set Label Format: eBizKanban  |  Timezone: GMT+5:30.",
        "Click Save → verify supplier appears in External Supplier list.",
    ])

    subhead(doc, "5.2  Create Supplier User (TestSup197)")
    prompt_block(doc, "File: QALS11/Admin/Supplier/supplierusercreation.py", [
        "Login → Admin → Supplier menu → select TestSup197.",
        "Click Add User → Create User form.",
        "Account Type: Supplier Administrators.",
        "Enter unique User ID and password (8+ chars, 1 num, 1 upper, 1 lower, 1 punct).",
        "Click Save → verify user created under TestSup197.",
    ])

    info_table(doc,
        rows=[
            ("test_create_supplier.py",  "QALS11/Admin/Supplier/", "Create supplier with eBizKanban label, GMT+5:30"),
            ("supplierusercreation.py",  "QALS11/Admin/Supplier/", "Create Supplier Admin user under TestSup197"),
        ],
        headers=["File", "Folder", "Purpose"],
    )
    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════════
    # SECTION 6 — POSTMAN / POST API
    # ══════════════════════════════════════════════════════════════════════════
    heading(doc, "6.  Postman API Automation", rgb="7030A0")
    body(doc, "Folder: qa_automation/tests/QALS11/Actions/POSTAPI/  |  URL: https://engineering.ultrivalms.com/QALS1011")

    subhead(doc, "6.1  Setup & Login via API", rgb="7030A0")
    prompt_block(doc, "File: QALS11/Actions/POSTAPI/test_api_login.py", [
        "Navigate to https://engineering.ultrivalms.com/QALS1011.",
        "POST /Account/Login with username=buyeruser, password=Ultriva!23.",
        "Capture session cookie / auth token from response.",
        "Verify 200 OK response and redirect to dashboard.",
    ])

    subhead(doc, "6.2  Create Postman Collection — Add User", rgb="7030A0")
    prompt_block(doc, "Postman Collection: Ultriva_QALS1011.postman_collection.json", [
        "Create Postman collection: Ultriva QALS1011.",
        "Add request: POST /Admin/User/Create with Add User payload.",
        "Set environment variable: {{base_url}} = https://engineering.ultrivalms.com/QALS1011",
        "Set Auth: use session cookie captured from login request.",
        "Add Tests tab: pm.test status 200, pm.test user created assertion.",
        "Run collection → verify Add User page accessible and user created.",
        "Export collection as JSON and attach to JIRA ticket ULT-10425 (if required).",
    ])

    info_table(doc,
        rows=[
            ("test_api_login.py",      "QALS11/Actions/POSTAPI/", "Session login via POST API"),
            ("test_api_add_user.py",   "QALS11/Actions/POSTAPI/", "Add user via API, assert 200 OK"),
            ("Ultriva_QALS1011.postman_collection.json", "QALS11/Actions/POSTAPI/", "Postman collection export"),
        ],
        headers=["File", "Folder", "Purpose"],
    )
    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════════
    # SECTION 7 — QUICK REFERENCE
    # ══════════════════════════════════════════════════════════════════════════
    heading(doc, "7.  Quick Reference Table")
    body(doc, "Summary of all prompts, files, and folders for the demo.")

    info_table(doc,
        rows=[
            # TestRail
            ("TestRail","Milestone","01_create_milestone.py",       "Practise on test rail/"),
            ("TestRail","Test Plan","02_create_plan.py",            "Practise on test rail/"),
            ("TestRail","Add Suite S2571","02_create_plan.py",      "Practise on test rail/"),
            ("TestRail","Pass Case C5157814","03_pass_testcase.py", "Practise on test rail/"),
            ("TestRail","Bug Case ULT-10069","ULT-10069_add_testrail_case.py","JIRA Tickets/Bugs/"),
            ("TestRail","Enhancement ULT-10425","ULT-10425_add_testrail_cases.py","JIRA Tickets/Enhancements/"),
            # JIRA
            ("JIRA","ULT-10511 plan+cases","ULT-10511_test_cases.py",  "JIRA Tickets/Enhancements/"),
            ("JIRA","ULT-10425 subtask",   "ULT-10425_test_cases.py",  "JIRA Tickets/Enhancements/"),
            ("JIRA","ULT-9764 cases",      "ULT-9764_test_cases.py",   "JIRA Tickets/Enhancements/"),
            ("JIRA","ULT-10069 bug cases", "ULT-10069_test_cases.py",  "JIRA Tickets/Bugs/"),
            # QA Automation
            ("Selenium","Create Users","test_create_users.py",           "QALS11/Admin/"),
            ("Selenium","Edit User",   "test_edit_userid.py",            "QALS11/Admin/"),
            ("Selenium","Create Item", "test_create_item.py",            "QALS11/Admin/Item/"),
            ("Selenium","Initial Loop","ItemInitialLoopSize.py",         "QALS11/Admin/Item/"),
            ("Selenium","Add Rep Cards","ItemAdjustLoopSize.py",         "QALS11/Admin/Item/"),
            ("Selenium","Add Non-Rep Cards","ItemAdjustLoopSize-NonReplinishmentcards.py","QALS11/Admin/Item/"),
            ("Selenium","Create Supplier","test_create_supplier.py",     "QALS11/Admin/Supplier/"),
            ("Selenium","Supplier User","supplierusercreation.py",       "QALS11/Admin/Supplier/"),
            # API
            ("Postman/API","Login API",  "test_api_login.py",            "QALS11/Actions/POSTAPI/"),
            ("Postman/API","Add User API","test_api_add_user.py",        "QALS11/Actions/POSTAPI/"),
        ],
        headers=["Area", "Task", "File", "Folder"],
    )

    # ── Footer ────────────────────────────────────────────────────────────────
    p = doc.add_paragraph(
        f"\nGenerated: {datetime.now().strftime('%B %d, %Y at %H:%M')}  |  "
        "QA Demo Prompts — Ultriva LMS QALS1011"
    )
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(8)
    p.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.save(OUTPUT_DOC)
    print(f"✓ Saved → {OUTPUT_DOC}")


if __name__ == "__main__":
    build()
