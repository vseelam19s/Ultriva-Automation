"""Generate ULT-10604 Set Shipment Information test cases DOCX."""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(r"c:\Prcatices-AI\docs\ULT-10604_test_cases.docx")

COLS = ["ID", "Title", "Preconditions", "Steps", "Expected Result", "Pri", "Type", "AC"]
WIDTHS = [0.55, 1.35, 1.55, 2.6, 2.3, 0.4, 0.7, 0.45]

PRE = (
    "Supplier user SUP-A has access to Set Tracking/Packing Slip Numbers and "
    "Set Shipment Information. Card C1 belongs to SUP-A, is in SHIP state, "
    "has not been received, and has a current shipment cycle."
)
PRE_FLAGGED_US = PRE + " C1 is Steel Origin flagged and currently stores steelorigin=US."
PRE_FLAGGED_EMPTY = PRE + " C2 is Steel Origin flagged but has no steelorigin value."
PRE_UNFLAGGED = PRE + " C3 is not Steel Origin flagged."

SECTIONS = [
    ("1. Menu, Access and Navigation", [
        ("TC-01", "New menu item is visible to an authorized supplier",
         "Supplier user has access to ITMLS02190 (Set Tracking/Packing Slip Numbers).",
         "1. Sign in as the supplier user.\n2. Open Supplier Actions.",
         "'Set Shipment Information' appears under Supplier Actions.", "High", "Functional", "AC1"),
        ("TC-02", "Existing Set Tracking page remains present and unchanged",
         "Supplier user has access to both supplier actions.",
         "1. Open Supplier Actions.\n2. Open 'Set Tracking/Packing Slip Numbers'.\n3. Verify its title, route, fields, and normal save behavior.",
         "The old action remains available and behaves as before; the new page is a separate action.", "High", "Regression", "AC1"),
        ("TC-03", "Menu permission matches existing action",
         "Test supplier user with the access rights copied from ITMLS02190; also have a user without that access.",
         "1. Sign in as the authorized user and inspect Supplier Actions.\n2. Sign in as the unauthorized user and inspect Supplier Actions and attempt the new route directly.",
         "Authorized user can see and open the new page. Unauthorized user cannot use the new action; access is not broader than the existing action.", "High", "Security/Access", "AC1"),
        ("TC-04", "Page title, breadcrumb, help and icon are correct",
         "Supplier user can open Set Shipment Information.",
         "1. Open the new page.\n2. Inspect page title, breadcrumb, help message and menu icon.",
         "Title and breadcrumb read 'Set Shipment Information'; help explains that tracking, packing slip and Steel Origin can be updated for shipped cards; expected supplier-action icon is shown.", "Medium", "UI", "AC1"),
        ("TC-05", "Filter and result navigation use the new page",
         PRE,
         "1. Open Set Shipment Information.\n2. Submit the filter.\n3. Select a result and navigate through Items, Cards and Results.",
         "Each step stays on the new action and renders its own page/grid; it does not route to the old page or its results grid.", "High", "Functional", "AC1"),
    ]),
    ("2. Shipped Card List and Grid", [
        ("TC-06", "Only this supplier's shipped or quarantined cards are listed",
         "SUP-A has cards in SHIP, QUARANTINE, RECEIVED and other states; SUP-B also has shipped cards.",
         "1. Filter for SUP-A in Set Shipment Information.\n2. Inspect returned cards and their states/owners.",
         "Only SUP-A's eligible SHIP/QUARANTINE cards are listed. Received cards and other suppliers' cards are excluded.", "High", "Functional", "AC2"),
        ("TC-07", "Tracking and Packing Slip columns match the old page",
         PRE,
         "1. Open the new action for an eligible card.\n2. Compare tracking and packing slip columns with the existing page.",
         "Tracking # and Packing Slip # are available with the same editing behavior as the old page.", "High", "Regression", "AC2"),
        ("TC-08", "Flagged row shows Steel Origin prefilled with current value",
         PRE_FLAGGED_US,
         "1. Open the card list.\n2. Locate C1 and inspect its Steel Origin control.",
         "C1 shows a Steel Origin dropdown prefilled with its stored current value, US.", "High", "Functional", "AC2"),
        ("TC-09", "Unflagged row has no Steel Origin control",
         PRE_UNFLAGGED,
         "1. Open the card list.\n2. Locate C3 and inspect row controls.",
         "C3 has no Steel Origin dropdown or editable Steel Origin field; tracking and packing slip remain editable.", "High", "Functional", "AC2"),
        ("TC-10", "Steel Origin column hidden when no rows are flagged",
         "Filter results contain only unflagged eligible cards.",
         "1. Open Set Shipment Information with the filter returning no flagged cards.\n2. Inspect the grid headers and rows.",
         "The Steel Origin column is hidden when no returned row is flagged.", "Medium", "UI", "AC2"),
        ("TC-11", "Steel Origin column shown when at least one row is flagged",
         "Filter results contain both flagged and unflagged eligible cards.",
         "1. Open the card list.\n2. Inspect the Steel Origin column and each row.",
         "The column is visible; only flagged rows expose the dropdown, and unflagged rows do not.", "High", "Functional", "AC2"),
        ("TC-12", "Steel Origin apply-to-all updates checked flagged rows only",
         "List contains two flagged checked cards (US, AA), one unflagged checked card, and one flagged unchecked card.",
         "1. Select Steel Origin AA in the apply-to-all control.\n2. Apply to all.\n3. Inspect row values and selection state.",
         "Only checked flagged rows receive AA. Unflagged and unchecked rows are unchanged.", "High", "Functional", "AC2"),
        ("TC-13", "Apply-to-all does not update unselected rows",
         "List contains multiple flagged rows with different values.",
         "1. Check one flagged row.\n2. Set the apply-to-all value to AA.\n3. Apply.",
         "Only the checked flagged row changes; all other rows keep their existing values.", "High", "Functional", "AC2"),
        ("TC-14", "Result grid uses its own SetShipmentInfo results GID",
         "Code/configuration or runtime grid diagnostics are available.",
         "1. Complete a search and open the Results view.\n2. Inspect the grid identity/configuration used by the page.",
         "The Results view uses SetShipmentInfoCardsResults, not the existing incorrect SetWorkOrderNoCardsResults GID.", "Medium", "Configuration", "AC2"),
    ]),
    ("3. Save - Tracking, Packing Slip and Expected Receive Date", [
        ("TC-15", "Save corrected tracking and packing slip for shipped card",
         PRE,
         "1. Set a valid Tracking # and Packing Slip # for C1.\n2. Save.\n3. Reopen the action and inspect C1.",
         "Save succeeds and both values persist on the eligible shipped card.", "High", "Functional", "AC3"),
        ("TC-16", "Tracking number remains required",
         PRE,
         "1. Clear Tracking # for a selected card.\n2. Enter or retain a Packing Slip #.\n3. Save.",
         "Save is rejected with the established required-tracking validation; no card changes are committed.", "High", "Negative", "AC3"),
        ("TC-17", "Packing Slip number is massaged using existing rules",
         PRE,
         "1. Enter a Packing Slip # requiring the existing MassagePackingSlipNo normalization.\n2. Save.\n3. Inspect the persisted/displayed value.",
         "Packing Slip # is normalized exactly as on Set Tracking/Packing Slip Numbers; no regression in existing formatting rules.", "Medium", "Regression", "AC3"),
        ("TC-18", "Expected Receive Date remains optional",
         PRE,
         "1. Leave Expected Receive Date blank.\n2. Enter valid required shipment fields.\n3. Save.",
         "Save succeeds with no date supplied, matching the existing page's optional-date behavior.", "Medium", "Functional", "AC3"),
        ("TC-19", "Expected Receive Date saves when supplied",
         PRE,
         "1. Enter a valid Expected Receive Date and valid required shipment fields.\n2. Save.\n3. Reopen and inspect the date.",
         "Save succeeds and the supplied Expected Receive Date persists as expected.", "Medium", "Functional", "AC3"),
        ("TC-20", "Invalid date does not partially save card updates",
         PRE,
         "1. Enter an invalid Expected Receive Date and change tracking/packing slip.\n2. Save.",
         "Validation reports the invalid date according to existing behavior and no partial card update is committed.", "Medium", "Negative", "AC3"),
        ("TC-21", "Updating one card does not modify other selected or unselected cards",
         "At least two eligible cards have distinct tracking/packing slip values.",
         "1. Select only C1 and edit its values.\n2. Save.\n3. Reopen C2 and compare its values.",
         "Only the intended card is updated; other cards retain their original values.", "High", "Regression", "AC3"),
        ("TC-22", "Custom before/after action hooks still execute",
         "Test environment can observe OnBeforeCardAction and OnAfterCardAction hooks.",
         "1. Save a valid change through Set Shipment Information.\n2. Inspect hook invocation/logs.",
         "Both existing custom hooks run at the expected points around the save.", "Medium", "Integration", "AC3"),
    ]),
    ("4. Steel Origin Validation and Persistence", [
        ("TC-23", "Correct flagged card from US to AA",
         PRE_FLAGGED_US,
         "1. Change C1 Steel Origin from US to AA.\n2. Save.\n3. Reopen the page and inspect C1.",
         "Save succeeds; the current card shipment displays and stores AA.", "High", "E2E", "AC4"),
        ("TC-24", "Set missing Steel Origin for flagged card shipped by scan or EDI",
         PRE_FLAGGED_EMPTY + " C2 was shipped via scan or EDI.",
         "1. Set Steel Origin to US.\n2. Save.\n3. Reopen the card.",
         "Save succeeds and US is stored for the current shipment; the card is now eligible to print with that value.", "High", "E2E", "AC4"),
        ("TC-25", "Flagged row requires Steel Origin",
         PRE_FLAGGED_EMPTY,
         "1. Leave the flagged row's Steel Origin blank.\n2. Enter a valid Tracking #.\n3. Save.",
         "Save is rejected with 'Steel Origin is required.' and the card is not updated.", "High", "Negative", "AC4"),
        ("TC-26", "Only US and AA are accepted",
         PRE_FLAGGED_US,
         "1. Attempt to submit a flagged card with an unsupported Steel Origin value (for example CA).\n2. Save.",
         "Save is rejected with 'Invalid Steel Origin value.'; no card state or settings are changed.", "High", "Negative", "AC4"),
        ("TC-27", "Blank and unsupported values have distinct status messages",
         PRE_FLAGGED_EMPTY,
         "1. Submit a blank Steel Origin and record the message.\n2. Submit an unsupported value and record the message.",
         "Blank shows 'Steel Origin is required.'; unsupported value shows 'Invalid Steel Origin value.'", "Medium", "Validation", "AC4"),
        ("TC-28", "Unflagged card does not require or accept a Steel Origin edit",
         PRE_UNFLAGGED,
         "1. Update Tracking # and Packing Slip # while leaving Steel Origin unavailable.\n2. Save.",
         "Tracking and Packing Slip save normally; no Steel Origin validation is applied to the unflagged card.", "High", "Functional", "AC4"),
        ("TC-29", "Apply-to-all accepts only whitelisted values",
         "At least one flagged card is checked.",
         "1. Enter or select a value other than US or AA in Steel Origin apply-to-all.\n2. Apply and save.",
         "Unsupported value is rejected and no affected card's Steel Origin is changed.", "Medium", "Negative", "AC4"),
        ("TC-30", "Steel Origin merge preserves other card IGParams keys",
         "Flagged card has steelorigin=US and additional settings keys in current cycle IGParams.",
         "1. Correct Steel Origin to AA.\n2. Save.\n3. Compare the complete current-cycle IGParams string before and after.",
         "steelorigin is updated to AA and every unrelated key/value remains intact.", "High", "Data integrity", "AC4"),
        ("TC-31", "Steel Origin is written to card settings and current cycle",
         PRE_FLAGGED_US,
         "1. Change Steel Origin to AA and save.\n2. Inspect the card settings string and EboKanbanCycleData.IGParams for CurrentCycleID.",
         "Both storage locations contain steelorigin=AA for the current shipment cycle.", "High", "Integration", "AC4"),
        ("TC-32", "Past completed cycle is not changed",
         "Card has current shipment cycle and one or more completed historical cycles with steelorigin=US.",
         "1. Correct current cycle Steel Origin to AA.\n2. Inspect current and completed cycle IGParams.",
         "Current shipment is AA. Past completed cycles remain unchanged.", "High", "Data integrity", "AC4"),
        ("TC-33", "Invalid Steel Origin does not partially save tracking or packing slip",
         PRE_FLAGGED_US,
         "1. Change Tracking # and Packing Slip #.\n2. Set Steel Origin to an invalid value.\n3. Save.\n4. Reload and inspect all values.",
         "Validation fails and none of the edited values are partially committed.", "High", "Negative", "AC3/AC4"),
    ]),
    ("5. Card History, Audit and Label Reprint", [
        ("TC-34", "Correction creates a new card-history version with actor and time",
         PRE_FLAGGED_US + " Capture the current history version and user identity before editing.",
         "1. Correct C1 from US to AA and save.\n2. Inspect EboKanbanHistory version rows.",
         "A new history version contains AA, the logged-in user's ActionBy, and BPFL_ACTION_TIME for the correction.", "High", "Audit", "AC5"),
        ("TC-35", "Previous card-history version retains old Steel Origin",
         PRE_FLAGGED_US,
         "1. Correct C1 from US to AA and save.\n2. View current and prior history versions.",
         "Current version contains AA; prior version retains US. Existing version history is not overwritten.", "High", "Audit", "AC5"),
        ("TC-36", "History is committed through AddNotes and ActionUpdate",
         "Test environment can inspect the card history produced by AddNotes(false).",
         "1. Save a valid Steel Origin correction.\n2. Inspect the resulting history action/version.",
         "Save uses the AddNotes(false) -> EbnKanbanCard.AddNotes -> CopyCard2History(ActionUpdate) flow and produces the expected history version.", "Medium", "Integration", "AC5"),
        ("TC-37", "Reprinted label displays corrected Steel Origin",
         "Flagged card was corrected from US to AA; Emerson print plug-in and label template are configured.",
         "1. Reprint the corrected card.\n2. Inspect the human-readable label text.",
         "Label shows 'Steel Origin: AA', not US.", "High", "E2E", "AC6"),
        ("TC-38", "Both 2D barcode variants encode corrected Steel Origin",
         "Flagged card was corrected from US to AA; capture PX2dBarCode and PXIR2dBarCode.",
         "1. Reprint the corrected card.\n2. Inspect both 2D barcode payloads.",
         "Both variants encode the corrected AA value (5ZAA using the configured DI); neither contains the stale 5ZUS value.", "High", "E2E", "AC6"),
        ("TC-39", "Linear barcode encodes corrected Steel Origin",
         "Flagged card was corrected from US to AA; label includes PXSteelOrigin.",
         "1. Reprint the corrected card.\n2. Scan/read the linear barcode.",
         "Linear barcode encodes 5ZAA; it does not encode the stale 5ZUS value.", "High", "E2E", "AC6"),
        ("TC-40", "All reprint representations agree after correction",
         "Flagged card was corrected from US to AA and all label data is available.",
         "1. Reprint the card.\n2. Compare visible text, linear barcode, PX2dBarCode and PXIR2dBarCode.",
         "All label representations carry AA consistently; no output carries US.", "High", "E2E", "AC6"),
        ("TC-41", "Non-UI ship with missing Steel Origin is fixed before next print",
         PRE_FLAGGED_EMPTY + " Card was shipped by scan or EDI without Steel Origin.",
         "1. Set Steel Origin to US in Set Shipment Information.\n2. Save.\n3. Print/reprint the label.",
         "Label displays 'Steel Origin: US'; linear barcode and both 2D variants encode US.", "High", "E2E", "AC6"),
    ]),
    ("6. Regression and Edge Cases", [
        ("TC-42", "Received cards cannot be corrected through this page",
         "C4 belongs to SUP-A but is already RECEIVED.",
         "1. Search for C4 in Set Shipment Information.\n2. Attempt to open or update it.",
         "C4 is not listed and cannot be updated through this action.", "High", "Regression", "AC2"),
        ("TC-43", "Quarantined card remains eligible where query includes QUARANTINE",
         "C5 belongs to SUP-A and is in QUARANTINE state, not received.",
         "1. Search for C5 in Set Shipment Information.",
         "C5 appears in the eligible list, consistent with the reused shipped-or-quarantined query.", "Medium", "Functional", "AC2"),
        ("TC-44", "Concurrent changes do not overwrite unrelated card settings",
         "Card has multiple current-cycle IGParams keys; a second process updates a non-Steel-Origin key before save.",
         "1. Load the action page.\n2. Update a non-Steel-Origin key in the second process.\n3. Correct Steel Origin and save from the page.\n4. Inspect current-cycle IGParams.",
         "Steel Origin correction does not silently discard unrelated settings; conflicts follow the application's established persistence behavior.", "Medium", "Concurrency", "AC4"),
        ("TC-45", "No change to old action route, results grid or scan variant",
         "Existing ITMLS02190 page and SUPPLIERSCAN/SETTRACKINGNO variant are available.",
         "1. Exercise the old page and scan variant.\n2. Compare routes and page behavior before and after enabling the new page.",
         "The old page and scan/handheld variant remain unchanged; only the new supplier action is added.", "Medium", "Regression", "AC1"),
    ]),
]

AC_MATRIX = [
    ("AC1 - New menu item; old page remains", "TC-01 to TC-05, TC-45"),
    ("AC2 - Eligible cards and row/grid behavior", "TC-06 to TC-14, TC-42, TC-43"),
    ("AC3 - Tracking/Packing Slip behavior and existing hooks", "TC-07, TC-15 to TC-22, TC-33"),
    ("AC4 - Steel Origin required/whitelisted and persisted to current shipment", "TC-08 to TC-13, TC-23 to TC-33, TC-44"),
    ("AC5 - New history version records correction while retaining prior value", "TC-34 to TC-36"),
    ("AC6 - Corrected value appears in reprinted label and all barcode forms", "TC-37 to TC-41"),
]


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def write(cell, text, size=8, bold=False, color=None, center=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lines = text.split("\n")
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if index < len(lines) - 1:
            run.add_break()


def heading(doc, text, size=13):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)


def set_widths(table, widths):
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(section, margin, Inches(0.5))

title = doc.add_paragraph()
run = title.add_run("ULT-10604 - Test Cases")
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
doc.add_paragraph(
    "Set Shipment Information | Correct shipped-card tracking, packing slip and Steel Origin"
).runs[0].font.size = Pt(11)

info = doc.add_table(rows=0, cols=2)
info.style = "Table Grid"
for key, value in [
    ("Epic / feature", "ULT-10604 - Feature 6"),
    ("Related tickets", "Depends on ULT-10605 and ULT-10607; access ticket ULT-10610 (Atlassian API access unavailable during preparation)."),
    ("Decision", "D4: create a new cloned Set Shipment Information page; keep ITMLS02190 Set Tracking/Packing Slip Numbers unchanged."),
    ("Scope", "Update Tracking #, Packing Slip #, optional Expected Receive Date and flagged-row Steel Origin (US/AA) for this supplier's SHIP/QUARANTINE cards that are not yet received. Verify current-cycle settings merge, history versioning, and corrected label reprint values."),
    ("Out of scope", "Scan/handheld variant; received cards; updates to past completed cycles; implementation of the AG Grid forward-merge version for release/10.11."),
    ("Open points", "ULT-10610 ticket details could not be fetched because the Atlassian API connection denied access. Test scenarios are derived from the supplied epic description and acceptance criteria."),
    ("Total test cases", str(sum(len(cases) for _, cases in SECTIONS))),
]:
    row = info.add_row()
    write(row.cells[0], key, size=9, bold=True)
    shade(row.cells[0], "DCE6F1")
    write(row.cells[1], value, size=9)
set_widths(info, [1.5, 8.5])

for title_text, cases in SECTIONS:
    heading(doc, f"{title_text} ({len(cases)})")
    table = doc.add_table(rows=1, cols=len(COLS))
    table.style = "Table Grid"
    for index, column in enumerate(COLS):
        write(table.rows[0].cells[index], column, size=9, bold=True, color="FFFFFF", center=True)
        shade(table.rows[0].cells[index], "1F4E79")
    for case_index, case in enumerate(cases):
        row = table.add_row()
        for index, value in enumerate(case):
            write(row.cells[index], value, bold=(index == 0), center=(index in (0, 5, 6, 7)))
            if case_index % 2:
                shade(row.cells[index], "F2F2F2")
    set_widths(table, WIDTHS)

heading(doc, "Acceptance Criteria Traceability")
matrix = doc.add_table(rows=1, cols=2)
matrix.style = "Table Grid"
for index, column in enumerate(("Acceptance Criteria", "Test Cases")):
    write(matrix.rows[0].cells[index], column, size=9, bold=True, color="FFFFFF")
    shade(matrix.rows[0].cells[index], "1F4E79")
for acceptance, test_cases in AC_MATRIX:
    row = matrix.add_row()
    write(row.cells[0], acceptance, size=9)
    write(row.cells[1], test_cases, size=9)
set_widths(matrix, [4.0, 6.0])

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"Created: {OUT} ({sum(len(cases) for _, cases in SECTIONS)} test cases)")
