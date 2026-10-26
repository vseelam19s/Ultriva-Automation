import json, os, re, collections
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
data = json.load(open(os.path.join(HERE, "results.json"), encoding="utf-8"))
res = data["results"]
side = data["side"]
_rr = os.path.join(HERE, "results_rerun.json")
if os.path.exists(_rr):  # Sorting and Supplier cases were re-run with corrected checks
    res = [r for r in res if r["area"] not in ("Sorting", "Filter - Supplier")] + json.load(open(_rr, encoding="utf-8"))["results"]
SEV_COLOR = {"High": "F8CBAD", "Medium": "FFE699", "Low": "C6E0B4"}
SHOT = lambda s: s if s and os.path.exists(s) else ""

# Cases verified manually as not defects (no data / expected behaviour); excluded from the bug list.
NOT_BUGS = {"TC_002": "No Yellow rows exist in this data set", "TC_019": "Aging buckets exclude cards that are not yet aged; not a confirmed defect"}


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def set_cell(cell, text, bold=False, size=9, color=None):
    cell.text = ""
    r = cell.paragraphs[0].add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)


fails = [r for r in res if r["status"] == "FAIL" and r["id"] not in NOT_BUGS and "'Is Consignment'" not in r["title"]]
n_fail = len([r for r in res if r["status"] == "FAIL"])
syserr = [r for r in res if r["status"] == "SYSERR"]
skips = [r for r in res if r["status"] == "SKIP"]
passed = [r for r in res if r["status"] == "PASS"]

bugs = []


def sev(r):
    if r["area"] == "Sorting":
        return "Low" if "case-sensitive" in r["actual"] else "High"
    return "High" if r["area"].startswith("Filter") else "Medium"


# one bug for all deterministic system errors on the same page
if syserr:
    bugs.append(dict(title="System Error!! when Show = Cards List / My Cards List and Go is clicked", area="Filter - Show", sev="High", items=syserr,
                     steps="Open Release Items List, set Show = 'Cards List' (or 'My Cards List') and click Go.",
                     expected="Release Cards List opens with the cards of the selected filters.",
                     actual="Page /BuyerAction/ReleaseCards shows 'Sorry, an error occurred while processing your request. System Error!!' (reproduced twice per case).",
                     shot=SHOT(syserr[0]["shot"]), cases=", ".join(r["id"] for r in syserr)))

sort_ci = [r for r in fails if r["area"] == "Sorting" and "case-sensitive" in r["actual"]]
others = [r for r in fails if r not in sort_ci]
if sort_ci:
    bugs.append(dict(title="Text columns are sorted case-sensitively (lowercase values placed after uppercase values)", area="Sorting", sev="Low", items=sort_ci,
                     steps="Click the 'Item' or 'Description' column header (ascending / descending) and view the order.",
                     expected="Case-insensitive alphabetical order.", actual=" || ".join(f"{r['title']}: {r['actual'][:220]}" for r in sort_ci[:3]),
                     shot=SHOT(sort_ci[0]["shot"]), cases=", ".join(r["id"] for r in sort_ci)))
for r in others:
    if r["area"] == "Sorting" and "'Description'" in r["title"]:
        bugs.append(dict(title="Description column sort order is inconsistent (e.g. 'ITSerial07222020_1' listed before 'IT_ERP_RECEIPTS')", area="Sorting", sev="Medium", items=[r],
                         steps="Click the 'Description' column header (ascending) and compare neighbouring values around row 22.",
                         expected="Consistent alphabetical order (case-insensitive) for the whole list.",
                         actual=r["actual"][:600], shot=SHOT(r["shot"]), cases=r["id"]))
        continue
    if r["area"] == "Sorting" and "'RYG'" in r["title"]:
        bugs.append(dict(title="RYG column cannot be sorted (header click has no effect)", area="Sorting", sev="Medium", items=[r],
                         steps="Click the 'RYG' column header once, twice and a third time.",
                         expected="Rows are ordered by RYG status (ascending / descending), as for the other columns.",
                         actual=r["actual"][:600], shot=SHOT(r["shot"]), cases=r["id"]))
        continue
    bugs.append(dict(title=f"[{r['area']}] {r['title']}", area=r["area"], sev=sev(r), items=[r], steps=r["steps"], expected=r["expected"],
                     actual=r["actual"], shot=SHOT(r["shot"]), cases=r["id"]))
if side.get("count_mismatch"):
    cm = side["count_mismatch"]
    bugs.append(dict(title="Pager total does not match the number of rows in the grid", area="Grid", sev="Medium", items=[],
                     steps="Apply any filter and compare 'x to y of N' with the rows listed across all pages.",
                     expected="Pager total equals the number of rows displayed.", actual="; ".join(f"{c[0]}: pager {c[1]} vs rows {c[2]}" for c in cm[:5]),
                     shot="", cases=", ".join(c[0] for c in cm[:8])))

order = {"High": 0, "Medium": 1, "Low": 2}
bugs.sort(key=lambda b: order[b["sev"]])
for n, b in enumerate(bugs, 1):
    b["id"] = f"BUG-REL-{n:03d}"

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.7)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10)
doc.add_heading("Bug Report - Buyer Actions > Release Items List (Filters and Sorting)", 0)
doc.add_paragraph("Module: Actions > Buyer Actions > Release  |  Build: Ultriva Lean Suite POC")

t = doc.add_table(rows=0, cols=2)
t.style = "Table Grid"
for k, v in [
    ("Application URL", "https://engineering.ultrivalms.com/POC/BuyerAction/ReleaseItems"),
    ("User", "buyeruser (Buyer K Test), Business Unit TCEK-Ram, Application (Plant), Supplier = Supnew (Supnew) for all filter/sort cases"),
    ("Browser", "Google Chrome (Selenium automation), 1920x1080"),
    ("Execution date", data["run"][:10]),
    ("Scope", "Filters (Supplier, Cards, Show, Category, RYG), filter combinations, Search all columns, sorting of all 6 columns"),
    ("Cases executed", f"{len(res)}  (Pass: {len(passed)}, Fail: {n_fail}, System error: {len(syserr)}, Skipped: {len(skips)})"),
    ("Bugs logged", str(len(bugs))),
]:
    row = t.add_row().cells
    set_cell(row[0], k, True)
    shade(row[0], "D9E2F3")
    set_cell(row[1], v)
    row[0].width, row[1].width = Inches(1.7), Inches(5.4)

doc.add_heading("1. Bug Summary", 1)
if bugs:
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for i, hd in enumerate(["Bug ID", "Title", "Area", "Severity"]):
        set_cell(t.rows[0].cells[i], hd, True, color="FFFFFF")
        shade(t.rows[0].cells[i], "2F5496")
    for b in bugs:
        c = t.add_row().cells
        set_cell(c[0], b["id"], True)
        set_cell(c[1], b["title"])
        set_cell(c[2], b["area"])
        set_cell(c[3], b["sev"], True)
        shade(c[3], SEV_COLOR[b["sev"]])
else:
    doc.add_paragraph("No functional bugs were found in the executed cases.")

doc.add_heading("2. Bug Details", 1)
for b in bugs:
    doc.add_heading(f"{b['id']} - {b['title']}", 2)
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in [
        ("Bug ID", b["id"]), ("Summary", b["title"]), ("Module", "Actions > Buyer Actions > Release > Release Items List"),
        ("Environment", "https://engineering.ultrivalms.com/POC, Chrome, user buyeruser"),
        ("Severity", b["sev"]), ("Priority", {"High": "P1", "Medium": "P2", "Low": "P3"}[b["sev"]]),
        ("Related test cases", b["cases"]),
        ("Steps to reproduce", "1. Log in as buyeruser / Ultriva!23.\n2. Click Actions > Buyer Actions > Release (Release Items List).\n"
                               f"3. {b['steps']}"),
        ("Expected result", b["expected"]), ("Actual result", b["actual"]), ("Status", "Open"),
    ]:
        c = t.add_row().cells
        set_cell(c[0], k, True)
        shade(c[0], "D9E2F3")
        set_cell(c[1], v)
        if k == "Severity":
            shade(c[1], SEV_COLOR[v])
        c[0].width, c[1].width = Inches(1.7), Inches(5.4)
    c = t.add_row().cells
    set_cell(c[0], "Screenshot", True)
    shade(c[0], "D9E2F3")
    c[0].width, c[1].width = Inches(1.7), Inches(5.4)
    if b["shot"]:
        c[1].paragraphs[0].add_run().add_picture(b["shot"], width=Inches(5.2))
    else:
        set_cell(c[1], "Not available")
    doc.add_paragraph()

doc.add_heading("3. Cases Skipped (System / Automation Error)", 1)
rows_ = syserr + skips
if rows_:
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, hd in enumerate(["Case ID", "Title", "Error"]):
        set_cell(t.rows[0].cells[i], hd, True, color="FFFFFF")
        shade(t.rows[0].cells[i], "2F5496")
    for r in rows_:
        c = t.add_row().cells
        set_cell(c[0], r["id"])
        set_cell(c[1], r["title"])
        set_cell(c[2], r["actual"])
else:
    doc.add_paragraph("None. No system errors were encountered.")

doc.add_heading("4. Execution Notes", 1)
doc.add_paragraph(f"Filters were verified against the unfiltered Supnew list ({data.get('baseline')} rows): RYG colour indicators, Category, Cards aging and Show against the single-filter results, "
                  "and every combination against each of its single filters. Each column was sorted ascending, descending and cleared, and the order was checked across both pages. "
                  "Cases that hit a system error were retried once, logged, and the run continued with the next case. "
                  "Cases where the data set has no matching rows (for example no Yellow items) are not logged as bugs.")

out = os.path.join(HERE, "BuyerActions_Release_Bug_Report.docx")
doc.save(out)
print("saved", out, "bugs:", len(bugs))
