"""Convert Item_TestCases.md to a Word document."""
import re
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH


def parse_md_and_create_docx():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    md_path = os.path.join(script_dir, "Item_TestCases.md")
    docx_path = os.path.join(script_dir, "Item_TestCases.docx")

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Style defaults
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip("\n")

        # Heading 1
        if line.startswith("# ") and not line.startswith("## "):
            doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue

        # Heading 2
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue

        # Heading 3
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue

        # Bold line (like **Application**: ...)
        if line.startswith("**") and "**" in line[2:]:
            p = doc.add_paragraph()
            # Parse bold and normal segments
            parts = re.split(r"(\*\*.*?\*\*)", line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Blockquote
        if line.startswith("> "):
            p = doc.add_paragraph(line[2:].strip(), style="Intense Quote")
            i += 1
            continue

        # Code block
        if line.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].rstrip(
                    "\n").startswith("```"):
                code_lines.append(lines[i].rstrip("\n"))
                i += 1
            i += 1  # skip closing ```
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            pf = p.paragraph_format
            pf.space_before = Pt(4)
            pf.space_after = Pt(4)
            continue

        # Bullet list
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            text = line[2:].strip()
            # Handle inline formatting
            parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
            for part in parts:
                if part.startswith("`") and part.endswith("`"):
                    run = p.add_run(part[1:-1])
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
                elif part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Table
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].rstrip("\n").startswith("|"):
                table_lines.append(lines[i].rstrip("\n"))
                i += 1

            # Parse header
            header_cells = [c.strip() for c in table_lines[0].split("|")[1:-1]]
            # Skip separator line (index 1)
            data_rows = []
            for tl in table_lines[2:]:
                cells = [c.strip() for c in tl.split("|")[1:-1]]
                data_rows.append(cells)

            num_cols = len(header_cells)
            table = doc.add_table(rows=1 + len(data_rows), cols=num_cols)
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT

            # Header row
            for ci, val in enumerate(header_cells):
                cell = table.rows[0].cells[ci]
                cell.text = ""
                run = cell.paragraphs[0].add_run(val)
                run.bold = True
                run.font.size = Pt(9)

            # Data rows
            for ri, row_data in enumerate(data_rows):
                for ci in range(min(num_cols, len(row_data))):
                    cell = table.rows[ri + 1].cells[ci]
                    cell.text = ""
                    text = row_data[ci]
                    # Strip backticks for display
                    clean = re.sub(r"`([^`]+)`", r"\1", text)
                    run = cell.paragraphs[0].add_run(clean)
                    run.font.size = Pt(9)
                    if "`" in text:
                        run.font.name = "Consolas"

            doc.add_paragraph()  # spacing after table
            continue

        # Horizontal rule
        if line.startswith("---"):
            doc.add_paragraph("─" * 60)
            i += 1
            continue

        # Empty / other
        if line.strip() == "":
            i += 1
            continue

        # Default paragraph
        p = doc.add_paragraph()
        parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", line)
        for part in parts:
            if part.startswith("`") and part.endswith("`"):
                run = p.add_run(part[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            elif part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                p.add_run(part)
        i += 1

    doc.save(docx_path)
    print(f"Created: {docx_path}")


if __name__ == "__main__":
    parse_md_and_create_docx()
