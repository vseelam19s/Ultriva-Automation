"""Generate test_data/credentials.xlsx and test_data/admin_test_data.xlsx."""
import os, time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

thin   = Side(style="thin")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

HDR_BG   = "1F4E79"
PASS_BG  = "C6EFCE"
FAIL_BG  = "FFC7CE"
INFO_BG  = "BDD7EE"
WARN_BG  = "FFE699"
ALT_BG   = "F2F2F2"
ADMIN_BG = "E2EFDA"

def fill(h): return PatternFill(start_color=h, end_color=h, fill_type="solid")
def hf(c="FFFFFF"): return Font(bold=True, color=c, size=11, name="Calibri")
def cf(): return Font(size=10, name="Calibri")
def bf(): return Font(bold=True, size=10, name="Calibri")
def center(): return Alignment(horizontal="center", vertical="center", wrap_text=True)
def left():   return Alignment(horizontal="left", vertical="center", wrap_text=True)

def write_hdr(ws, headers, widths, row=1):
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=row, column=ci, value=h)
        c.font = hf(); c.fill = fill(HDR_BG)
        c.alignment = center(); c.border = border
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[row].height = 22
    ws.freeze_panes = f"A{row+1}"

def wr(ws, ri, vals, bg=None):
    for ci, v in enumerate(vals, 1):
        c = ws.cell(row=ri, column=ci, value=v)
        c.font = cf(); c.border = border
        c.alignment = center() if ci in (1,2,5,6,7) else left()
        c.fill = fill(bg if bg else (ALT_BG if ri%2==0 else "FFFFFF"))
    ws.row_dimensions[ri].height = 18

os.makedirs("test_data", exist_ok=True)

# ══════════════════════════════════════════════════════════════════════════════
# FILE 1 : credentials.xlsx
# ══════════════════════════════════════════════════════════════════════════════
wb1 = openpyxl.Workbook()

# Sheet 1 — Valid Credentials
ws_creds = wb1.active
ws_creds.title = "Valid_Credentials"
write_hdr(ws_creds,
    ["#","Role","Username","Password","Access_Level","Module_Access","Notes"],
    [5, 18, 18, 18, 18, 45, 50])

cre_rows = [
    (1, "Buyer User",   "buyeruser",  "Ultriva!23", "Admin",     "All modules incl. Admin, User Mgmt, Config", "Primary test account — verified working"),
]
for ri, row in enumerate(cre_rows, 2):
    wr(ws_creds, ri, row, PASS_BG)

# Sheet 2 — Invalid / Test Credentials
ws_inv = wb1.create_sheet("Invalid_Credentials")
write_hdr(ws_inv,
    ["#","Scenario","Username","Password","Expected_Outcome","Risk_Level"],
    [5, 40, 20, 25, 35, 14])

inv_rows = [
    (1,  "Correct username, wrong password",           "buyeruser",       "WrongPass!99",   "Login rejected",           "Low"),
    (2,  "Wrong username, correct password",           "wronguser",       "Ultriva!23",     "Login rejected",           "Low"),
    (3,  "Both wrong",                                 "wronguser",       "WrongPass!99",   "Login rejected",           "Low"),
    (4,  "Empty username",                             "",                "Ultriva!23",     "Login rejected",           "Low"),
    (5,  "Empty password",                             "buyeruser",       "",               "Login rejected",           "Low"),
    (6,  "Both empty",                                 "",                "",               "Login rejected",           "Low"),
    (7,  "SQL injection attempt",                      "' OR 1=1 --",     "anything",       "Login rejected/sanitized", "High"),
    (8,  "Uppercase username (case-insensitive bug)",  "BUYERUSER",       "Ultriva!23",     "ACCEPTED -- security gap",  "Medium"),
    (9,  "Whitespace-only password (validation bug)",  "buyeruser",       "     ",          "ACCEPTED -- security gap",  "High"),
    (10, "Null-byte in username (bypass attempt)",     "buyeruser [NULL]","Ultriva!23",     "ACCEPTED -- security gap",  "Medium"),
]
for ri, row in enumerate(inv_rows, 2):
    bg = FAIL_BG if "gap" in str(row[4]).lower() else PASS_BG
    wr(ws_inv, ri, row, bg)

# Sheet 3 — Security Findings
ws_sec = wb1.create_sheet("Security_Findings")
write_hdr(ws_sec,
    ["Finding_ID","Severity","Title","Description","Affected_Input","Reproduction_Steps","Recommendation"],
    [12, 10, 35, 55, 18, 60, 55])

sec_rows = [
    ("SF-001","Medium","Case-insensitive username login",
     "App accepts BUYERUSER (uppercase) — usernames are not case-sensitive",
     "Username field",
     "1. Open login. 2. Enter BUYERUSER. 3. Enter valid pwd. 4. Click Login. 5. Observe: login succeeds.",
     "Enforce case-sensitive username comparison in auth layer"),
    ("SF-002","High","Whitespace-only password accepted",
     "App accepts 5 spaces as a valid password when combined with valid username",
     "Password field",
     "1. Open login. 2. Enter buyeruser. 3. Enter 5 spaces. 4. Click Login. 5. Observe: login succeeds.",
     "Trim and validate password server-side; reject blank/whitespace-only passwords"),
    ("SF-003","Medium","Null-byte in username stripped — auth bypassed",
     "Server strips null byte (\\x00) from username before lookup, allowing buyeruser\\x00 to auth",
     "Username field",
     "1. Open login. 2. Enter buyeruser\\x00. 3. Enter valid pwd. 4. Click Login. 5. Observe: auth succeeds.",
     "Sanitize/reject inputs containing null bytes at the preprocessing layer"),
]
for ri, row in enumerate(sec_rows, 2):
    bg = FAIL_BG if row[1]=="High" else WARN_BG
    wr(ws_sec, ri, row, bg)

# Sheet 4 — Environment
ws_env = wb1.create_sheet("Environment")
write_hdr(ws_env,
    ["Parameter","Value"],
    [28, 60])
env_data = [
    ("Application Name",     "Ultriva Lean Suite"),
    ("Environment",          "QALS1011 (QA)"),
    ("Base URL",             "https://engineering.ultrivalms.com/qals1011"),
    ("Admin URL",            "https://engineering.ultrivalms.com/QALS1011/Admin/Home"),
    ("Test Date",            time.strftime("%Y-%m-%d")),
    ("Test Framework",       "Selenium 4.18 + pytest 8.1"),
    ("Browser",              "Google Chrome (Selenium Manager)"),
    ("Python Version",       "3.11.9"),
    ("Report Output",        "reports/test_results.xlsx + reports/report.html"),
]
for ri, (k, v) in enumerate(env_data, 2):
    ws_env.cell(row=ri, column=1, value=k).font = bf()
    ws_env.cell(row=ri, column=2, value=v).font = cf()
    for ci in range(1,3):
        ws_env.cell(row=ri, column=ci).border = border
        ws_env.cell(row=ri, column=ci).fill = fill(ALT_BG if ri%2==0 else "FFFFFF")
        ws_env.cell(row=ri, column=ci).alignment = left()

wb1.save("test_data/credentials.xlsx")
print("Saved: test_data/credentials.xlsx")

# ══════════════════════════════════════════════════════════════════════════════
# FILE 2 : admin_test_data.xlsx
# ══════════════════════════════════════════════════════════════════════════════
wb2 = openpyxl.Workbook()

# ── Sheet 1: Admin_Access ──────────────────────────────────────────────────
ws_acc = wb2.active
ws_acc.title = "Admin_Access"
write_hdr(ws_acc,
    ["TC_ID","Description","URL","Login_Required","Expected_Status","Expected_Behavior"],
    [10, 50, 55, 14, 16, 40])

access_rows = [
    ("ADTC-001","Admin Home page loads",                          "/QALS1011/Admin/Home",              "Yes","PASS","Page loaded, heading visible"),
    ("ADTC-002","Admin Home redirects to login if unauthenticated","/QALS1011/Admin/Home",             "No", "PASS","Redirect to login page"),
    ("ADTC-003","Business Unit page loads",                       "/QALS1011/Admin/BusUnit",           "Yes","PASS","Plant edit form visible"),
    ("ADTC-004","Business Unit Config page loads",                "/QALS1011/Admin/BusUnitConfiguration","Yes","PASS","Name/Value table visible"),
    ("ADTC-005","User Group list loads",                          "/QALS1011/Group/Index",             "Yes","PASS","Group list with rows"),
    ("ADTC-006","User List page loads",                           "/QALS1011/User/List",               "Yes","PASS","User list with rows"),
    ("ADTC-007","Supplier List page loads",                       "/QALS1011/Supplier/List",           "Yes","PASS","Supplier list with rows"),
    ("ADTC-008","Customer List page loads",                       "/QALS1011/Customer/List",           "Yes","PASS","Customer list with rows"),
    ("ADTC-009","Workcenter List page loads",                     "/QALS1011/Workcenter/List",         "Yes","PASS","Workcenter list with rows"),
    ("ADTC-010","Shift List page loads",                          "/QALS1011/Shift/List",              "Yes","PASS","Shift list with rows"),
    ("ADTC-011","Calendar/Holiday List page loads",               "/QALS1011/Calendar/List",           "Yes","PASS","Holiday list with rows"),
    ("ADTC-012","Kanban Item List page loads",                    "/QALS1011/KanbanItem/List",         "Yes","PASS","Kanban items list with rows"),
    ("ADTC-013","Warehouse/Dropoff Location List page loads",     "/QALS1011/Warehouse/List",          "Yes","PASS","Warehouse list with rows"),
    ("ADTC-014","Commodity Code list page loads",                 "/QALS1011/Entity/List",             "Yes","PASS","Entity list with rows"),
    ("ADTC-015","Admin Tools page loads",                         "/QALS1011/Admin/Tools",             "Yes","PASS","Kanban & Data Load sections"),
    ("ADTC-016","About Info page loads",                          "/QALS1011/Admin/AboutInfo",         "Yes","PASS","Version/About table visible"),
]
for ri, row in enumerate(access_rows, 2):
    wr(ws_acc, ri, row, PASS_BG if row[4]=="PASS" else FAIL_BG)

# ── Sheet 2: Admin_List_Validation ────────────────────────────────────────
ws_lst = wb2.create_sheet("Admin_List_Validation")
write_hdr(ws_lst,
    ["TC_ID","Description","URL","Column_To_Check","Filter_Input_Name","Min_Rows","Expected_Status"],
    [10, 50, 45, 30, 25, 10, 14])

list_rows = [
    ("ADTC-101","User List has Status column",           "/QALS1011/User/List",       "Status",                     "User_FindTextbox",    1,"PASS"),
    ("ADTC-102","User List has User ID column",          "/QALS1011/User/List",       "User ID",                    "User_FindTextbox",    1,"PASS"),
    ("ADTC-103","User List has Email ID column",         "/QALS1011/User/List",       "Email ID",                   "User_FindTextbox",    1,"PASS"),
    ("ADTC-104","User List search filter present",       "/QALS1011/User/List",       "",                           "User_FindTextbox",    1,"PASS"),
    ("ADTC-105","Supplier List has Name column",         "/QALS1011/Supplier/List",   "Name",                       "Supplier_FindTextbox",1,"PASS"),
    ("ADTC-106","Supplier List has Location column",     "/QALS1011/Supplier/List",   "Location",                   "Supplier_FindTextbox",1,"PASS"),
    ("ADTC-107","Customer List has Name column",         "/QALS1011/Customer/List",   "Name",                       "Customer_FindTextbox",1,"PASS"),
    ("ADTC-108","Customer List has Code column",         "/QALS1011/Customer/List",   "Code",                       "Customer_FindTextbox",1,"PASS"),
    ("ADTC-109","Workcenter List has Name column",       "/QALS1011/Workcenter/List", "Name",                       "Workcenter_FindTextbox",1,"PASS"),
    ("ADTC-110","Workcenter List has Time zone column",  "/QALS1011/Workcenter/List", "Time zone",                  "Workcenter_FindTextbox",1,"PASS"),
    ("ADTC-111","Shift List has Shift No column",        "/QALS1011/Shift/List",      "Shift No",                   "",                    1,"PASS"),
    ("ADTC-112","Shift List has Start Time column",      "/QALS1011/Shift/List",      "Start Time",                 "",                    1,"PASS"),
    ("ADTC-113","KanbanItem List has Item column",       "/QALS1011/KanbanItem/List", "Item",                       "MISC$SName",          1,"PASS"),
    ("ADTC-114","KanbanItem List has Status column",     "/QALS1011/KanbanItem/List", "Status",                     "MISC$RYG",            1,"PASS"),
    ("ADTC-115","Warehouse List has Name column",        "/QALS1011/Warehouse/List",  "Name",                       "Warehouse_FindTextbox",1,"PASS"),
    ("ADTC-116","Entity List has Code column",           "/QALS1011/Entity/List",     "Code",                       "Family_FindTextbox",  1,"PASS"),
    ("ADTC-117","Calendar List has Date column",         "/QALS1011/Calendar/List",   "Date",                       "",                    1,"PASS"),
    ("ADTC-118","BusUnit Config has Name col",           "/QALS1011/Admin/BusUnitConfiguration","Name",             "",                    1,"PASS"),
    ("ADTC-119","BusUnit Config has Value col",          "/QALS1011/Admin/BusUnitConfiguration","Value",            "",                    1,"PASS"),
    ("ADTC-120","User Group list has Title column",      "/QALS1011/Group/Index",     "Title",                      "",                    1,"PASS"),
]
for ri, row in enumerate(list_rows, 2):
    wr(ws_lst, ri, row, PASS_BG if row[6]=="PASS" else FAIL_BG)

# ── Sheet 3: Admin_Search_Filter ──────────────────────────────────────────
ws_srch = wb2.create_sheet("Admin_Search_Filter")
write_hdr(ws_srch,
    ["TC_ID","Description","URL","Filter_Name","Search_Term","Expected_Status","Notes"],
    [10, 48, 42, 25, 20, 14, 40])

search_rows = [
    ("ADTC-201","User filter input accepts text",     "/QALS1011/User/List",       "User_FindTextbox",    "buyer",     "PASS","Filter input field operable"),
    ("ADTC-202","Supplier filter accepts text",       "/QALS1011/Supplier/List",   "Supplier_FindTextbox","test",      "PASS","Filter input field operable"),
    ("ADTC-203","Customer filter accepts text",       "/QALS1011/Customer/List",   "Customer_FindTextbox","test",      "PASS","Filter input field operable"),
    ("ADTC-204","Workcenter filter accepts text",     "/QALS1011/Workcenter/List", "Workcenter_FindTextbox","test",    "PASS","Filter input field operable"),
    ("ADTC-205","Warehouse filter accepts text",      "/QALS1011/Warehouse/List",  "Warehouse_FindTextbox","test",     "PASS","Filter input field operable"),
    ("ADTC-206","KanbanItem supplier filter",         "/QALS1011/KanbanItem/List", "MISC$SName",           "test",     "PASS","Supplier name filter operable"),
    ("ADTC-207","Entity family filter accepts text",  "/QALS1011/Entity/List",     "Family_FindTextbox",   "test",     "PASS","Filter input field operable"),
    ("ADTC-208","Global search bar accepts text",     "/QALS1011/KanbanItem/List", "txtTbSearch",          "test",     "PASS","Top search bar operable"),
]
for ri, row in enumerate(search_rows, 2):
    wr(ws_srch, ri, row, PASS_BG if row[5]=="PASS" else FAIL_BG)

# ── Sheet 4: Admin_Navigation ─────────────────────────────────────────────
ws_nav = wb2.create_sheet("Admin_Navigation")
write_hdr(ws_nav,
    ["TC_ID","Description","Start_URL","Action","Expected_URL_Contains","Expected_Status"],
    [10, 48, 42, 30, 40, 14])

nav_rows = [
    ("ADTC-301","Admin Home nav link in sidebar",        "/QALS1011/KanbanItem/List",  "Click Admin nav link",    "Admin/Home",              "PASS"),
    ("ADTC-302","Breadcrumb link navigates to previous page","/QALS1011/Admin/AboutInfo", "Click Admin breadcrumb",  "QALS1011",                "PASS"),
    ("ADTC-303","About Info page accessible from Admin", "/QALS1011/Admin/Home",       "Navigate to AboutInfo",   "Admin/AboutInfo",         "PASS"),
    ("ADTC-304","Tools page accessible from Admin",      "/QALS1011/Admin/Home",       "Navigate to Tools",       "Admin/Tools",             "PASS"),
    ("ADTC-305","BusUnit page accessible from Admin",    "/QALS1011/Admin/Home",       "Navigate to BusUnit",     "Plant",                   "PASS"),
    ("ADTC-306","User List accessible from Admin",       "/QALS1011/Admin/Home",       "Navigate to User/List",   "User/List",               "PASS"),
    ("ADTC-307","Supplier List accessible from Admin",   "/QALS1011/Admin/Home",       "Navigate to Supplier",    "Supplier/List",           "PASS"),
    ("ADTC-308","Kanban Item page title correct",        "/QALS1011/KanbanItem/List",  "Check page heading",      "Items List",              "PASS"),
    ("ADTC-309","Shift list page title correct",         "/QALS1011/Shift/List",       "Check page heading",      "Shift List",              "PASS"),
    ("ADTC-310","Calendar page title correct",           "/QALS1011/Calendar/List",    "Check page heading",      "Holiday List",            "PASS"),
]
for ri, row in enumerate(nav_rows, 2):
    wr(ws_nav, ri, row, PASS_BG if row[5]=="PASS" else FAIL_BG)

# ── Sheet 5: Summary ──────────────────────────────────────────────────────
ws_sum = wb2.create_sheet("Summary")
ws_sum.column_dimensions["A"].width = 32
ws_sum.column_dimensions["B"].width = 14
ws_sum.column_dimensions["C"].width = 48

title = ws_sum.cell(row=1, column=1, value="Admin Module Test Data — Ultriva QALS1011")
title.font = Font(bold=True, size=14, color=HDR_BG, name="Calibri")
title.alignment = center()
ws_sum.merge_cells("A1:C1")
ws_sum.row_dimensions[1].height = 30

hrd_row = [("Test Sheet","# Cases","Coverage")]
for ci, h in enumerate(hrd_row[0], 1):
    c = ws_sum.cell(row=2, column=ci, value=h)
    c.font = hf(HDR_BG); c.fill = fill("BDD7EE")
    c.alignment = center(); c.border = border

sheets_info = [
    ("Admin_Access",         len(access_rows), "Page load & unauthenticated redirect checks"),
    ("Admin_List_Validation",len(list_rows),   "Column presence & row count on all list pages"),
    ("Admin_Search_Filter",  len(search_rows), "Filter input operability on each list page"),
    ("Admin_Navigation",     len(nav_rows),    "Breadcrumb, sidebar, and URL routing checks"),
]
for ri, (sh, cnt, cov) in enumerate(sheets_info, 3):
    ws_sum.cell(row=ri, column=1, value=sh).font = cf()
    ws_sum.cell(row=ri, column=2, value=cnt).font = cf()
    ws_sum.cell(row=ri, column=3, value=cov).font = cf()
    bg = ALT_BG if ri%2==0 else "FFFFFF"
    for ci in range(1,4):
        ws_sum.cell(row=ri, column=ci).fill = fill(bg)
        ws_sum.cell(row=ri, column=ci).alignment = center() if ci==2 else left()
        ws_sum.cell(row=ri, column=ci).border = border

total_ri = 3+len(sheets_info)
for ci, v in [(1,"TOTAL"),(2, sum(x[1] for x in sheets_info)),(3,"Full admin module backend coverage")]:
    c = ws_sum.cell(row=total_ri, column=ci, value=v)
    c.font = bf(); c.fill = fill(ADMIN_BG)
    c.alignment = center() if ci==2 else left()
    c.border = border

wb2.save("test_data/admin_test_data.xlsx")
print("Saved: test_data/admin_test_data.xlsx")
print()
print(f"credentials.xlsx    : {len(cre_rows)} valid creds, {len(inv_rows)} invalid scenarios, {len(sec_rows)} security findings")
print(f"admin_test_data.xlsx: {len(access_rows)+len(list_rows)+len(search_rows)+len(nav_rows)} test cases across 4 sheets")
