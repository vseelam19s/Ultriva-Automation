"""
Convert MCP_Servers_Documentation.html to a formatted DOCX with highlighted main points.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

OUTPUT_PATH = r"C:\Users\vseelam\Downloads\MCP_Servers_Documentation.docx"

# ── Colour palette ──────────────────────────────────────────────────────────
C_TITLE      = RGBColor(0x1F, 0x49, 0x7D)   # dark navy  – main title
C_H2         = RGBColor(0x2E, 0x74, 0xB5)   # medium blue – server names
C_H3         = RGBColor(0x20, 0x6A, 0x52)   # dark green  – sub‑headings
C_CODE_FG    = RGBColor(0x22, 0x22, 0x22)   # near‑black  – code text
C_CODE_BG    = RGBColor(0xF2, 0xF2, 0xF2)   # light grey  – code cell shading
C_HIGHLIGHT  = RGBColor(0xFF, 0xC0, 0x00)   # amber       – key points (bold)
C_WARN       = RGBColor(0xC0, 0x39, 0x2B)   # red         – security/warning notes

# ── Helpers ──────────────────────────────────────────────────────────────────

def set_cell_shading(paragraph, rgb: RGBColor):
    """Apply background colour to a paragraph via XML shading."""
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    # RGBColor is a tuple subclass (R, G, B)
    hex_color = "{:02X}{:02X}{:02X}".format(rgb[0], rgb[1], rgb[2])
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def add_title(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = C_TITLE
    doc.add_paragraph()   # spacer


def add_h2(doc: Document, text: str):
    """Server‑level heading (numbered, e.g. '1) selenium')."""
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = C_H2


def add_h3(doc: Document, text: str):
    """Sub‑section heading."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = C_H3


def add_body(doc: Document, text: str):
    doc.add_paragraph(text)


def add_bullet(doc: Document, text: str, level: int = 0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    p.add_run(text)


def add_highlight_bullet(doc: Document, text: str):
    """Amber-bold bullet for key/important points."""
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25)
    run = p.add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(0xC0, 0x6B, 0x00)   # darker amber – readable


def add_warning_bullet(doc: Document, text: str):
    """Red bullet for security / gotcha notes."""
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25)
    run = p.add_run(text)
    run.bold = True
    run.font.color.rgb = C_WARN


def add_code_block(doc: Document, code: str):
    """Monospaced grey‑background block."""
    for line in code.strip().splitlines():
        p = doc.add_paragraph()
        run = p.add_run(line if line else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        run.font.color.rgb = C_CODE_FG
        set_cell_shading(p, C_CODE_BG)
    doc.add_paragraph()   # spacer after code


def add_separator(doc: Document):
    p = doc.add_paragraph("─" * 80)
    p.runs[0].font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    p.runs[0].font.size = Pt(8)


# ── Document content ─────────────────────────────────────────────────────────

def build_document():
    doc = Document()

    # ── Page margins (narrower for code readability)
    for section in doc.sections:
        section.top_margin    = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin   = Inches(1.0)
        section.right_margin  = Inches(1.0)

    # ── Default body font
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    # ════════════════════════════════════════════════════════════════════════
    # TITLE & INTRO
    # ════════════════════════════════════════════════════════════════════════
    add_title(doc, "MCP Servers Documentation (Detailed)")

    p = doc.add_paragraph()
    p.add_run("Date: ").bold = True
    p.add_run("2026-04-13    ")
    r = p.add_run("User: ")
    r.bold = True
    p.add_run("vseelam19s")

    doc.add_paragraph(
        "This document explains each entry under the servers section of the "
        "MCP configuration.  In MCP (Model Context Protocol), a server is a "
        "tool provider that the AI client can connect to — either by launching "
        "a local process (stdio) or by connecting to a remote endpoint (http)."
    )

    # ── KEY CONCEPT BOX ──────────────────────────────────────────────────────
    doc.add_paragraph()
    add_h3(doc, "Key Concepts at a Glance")
    key_points = [
        "MCP = Model Context Protocol — lets AI clients call external tools.",
        'stdio type  → local process; client spawns it and talks via stdin/stdout.',
        'http type   → remote endpoint; client calls over the network.',
        "Every server exposes one or more callable tools to the AI model.",
        "Secrets should never be stored in plain text — use prompted inputs.",
        "Pin package versions (@x.y.z) instead of @latest for production stability.",
    ]
    for kp in key_points:
        add_highlight_bullet(doc, kp)

    add_separator(doc)

    # ════════════════════════════════════════════════════════════════════════
    # SERVER 1 — selenium
    # ════════════════════════════════════════════════════════════════════════
    add_h2(doc, "1)  selenium")
    add_h3(doc, "Config")
    add_code_block(doc, """\
{
  "selenium": {
    "command": "npx",
    "args": ["-y", "@angiejones/mcp-selenium@latest"],
    "type": "stdio"
  }
}""")
    add_h3(doc, "What it is")
    add_body(doc, "An MCP server that exposes Selenium browser automation capabilities as callable tools.")
    add_h3(doc, "How it runs")
    add_bullet(doc, 'type: "stdio" — the MCP client starts a local process and communicates over stdin/stdout.')
    add_bullet(doc, "Runs: npx -y @angiejones/mcp-selenium@latest")
    add_bullet(doc, "npx — executes a Node package without global install.", level=1)
    add_bullet(doc, "-y — auto-confirm prompts.", level=1)
    add_bullet(doc, "@latest — always uses the newest release.", level=1)
    add_h3(doc, "Typical Use Cases")
    add_bullet(doc, "UI testing and browser automation")
    add_bullet(doc, "Automated login flows")
    add_bullet(doc, "Form fills, clicks, text extraction, screenshots")
    add_h3(doc, "Notes / Gotchas")
    add_warning_bullet(doc, "Requires Node.js and a compatible browser/driver setup.")
    add_warning_bullet(doc, "@latest can introduce breaking changes — pin a version for stability.")
    add_separator(doc)

    # ════════════════════════════════════════════════════════════════════════
    # SERVER 2 — postman
    # ════════════════════════════════════════════════════════════════════════
    add_h2(doc, "2)  com.postman/postman-mcp-server")
    add_h3(doc, "Config")
    add_code_block(doc, """\
{
  "com.postman/postman-mcp-server": {
    "type": "stdio",
    "command": "npx",
    "args": ["-y", "@postman/postman-mcp-server@2.7.0"],
    "env": {
      "POSTMAN_API_KEY": "${input:POSTMAN_API_KEY}"
    }
  }
}""")
    add_h3(doc, "What it is")
    add_body(doc, "Postman's MCP server to interact with Postman assets (collections, environments, workspaces, APIs) via the Postman API.")
    add_h3(doc, "How it runs")
    add_bullet(doc, "Local process via stdio.")
    add_highlight_bullet(doc, "Uses pinned version @postman/postman-mcp-server@2.7.0  (good practice!).")
    add_h3(doc, "Authentication")
    add_bullet(doc, "Uses POSTMAN_API_KEY environment variable.")
    add_highlight_bullet(doc, "${input:POSTMAN_API_KEY} is securely prompted by the client — never hard-coded.")
    add_h3(doc, "Typical Use Cases")
    add_bullet(doc, "Listing / exporting collections")
    add_bullet(doc, "Reading / updating environments")
    add_bullet(doc, "Automating API testing workflows")
    add_h3(doc, "Notes / Gotchas")
    add_warning_bullet(doc, "API key must have permissions for the relevant Postman workspace/resources.")
    add_separator(doc)

    # ════════════════════════════════════════════════════════════════════════
    # SERVER 3 — postgres
    # ════════════════════════════════════════════════════════════════════════
    add_h2(doc, "3)  postgres")
    add_h3(doc, "Config")
    add_code_block(doc, """\
{
  "postgres": {
    "command": "npx",
    "args": [
      "-y",
      "@modelcontextprotocol/server-postgres",
      "postgresql://postgres:password@localhost:5432/ecommerce_db"
    ],
    "type": "stdio"
  }
}""")
    add_h3(doc, "What it is")
    add_body(doc, "An MCP server that exposes PostgreSQL operations (querying and schema exploration) as tools.")
    add_h3(doc, "How it runs")
    add_bullet(doc, "Local process via stdio.")
    add_bullet(doc, "Connection string is passed directly in args.")
    add_h3(doc, "Typical Use Cases")
    add_bullet(doc, "Running SQL queries")
    add_bullet(doc, "Exploring schemas, tables, and columns")
    add_bullet(doc, "Debugging application data issues")
    add_h3(doc, "Notes / Gotchas / Security")
    add_warning_bullet(doc, "SECURITY: Credentials are in plain text in args — avoid sharing this config as-is.")
    add_warning_bullet(doc, "PostgreSQL must be running and reachable at the connection string.")
    add_separator(doc)

    # ════════════════════════════════════════════════════════════════════════
    # SERVER 4 — atlassian
    # ════════════════════════════════════════════════════════════════════════
    add_h2(doc, "4)  com.atlassian/atlassian-mcp-server")
    add_h3(doc, "Config")
    add_code_block(doc, """\
{
  "com.atlassian/atlassian-mcp-server": {
    "type": "http",
    "url": "https://mcp.atlassian.com/v1/mcp",
    "headers": {
      "Authorization": "Bearer ${input:ATLASSIAN_API_TOKEN}"
    }
  }
}""")
    add_h3(doc, "What it is")
    add_body(doc, "An Atlassian-hosted MCP endpoint to interact with Atlassian services (Jira, Confluence).")
    add_h3(doc, "How it connects")
    add_highlight_bullet(doc, 'type: "http" — remote connection; your client calls the endpoint over the network (not a local process).')
    add_h3(doc, "Authentication")
    add_bullet(doc, "Sends Authorization: Bearer <token> in headers.")
    add_highlight_bullet(doc, "Token provided via ${input:ATLASSIAN_API_TOKEN} — prompted securely.")
    add_h3(doc, "Typical Use Cases")
    add_bullet(doc, "Jira issue search / create / update")
    add_bullet(doc, "Confluence content retrieval")
    add_h3(doc, "Notes / Gotchas")
    add_warning_bullet(doc, "Token must be valid and authorised.")
    add_warning_bullet(doc, "Network / firewall / proxy settings can affect access.")
    add_separator(doc)

    # ════════════════════════════════════════════════════════════════════════
    # SERVER 5 — excel
    # ════════════════════════════════════════════════════════════════════════
    add_h2(doc, "5)  excel")
    add_h3(doc, "Config")
    add_code_block(doc, """\
{
  "excel": {
    "type": "stdio",
    "command": "npx",
    "args": ["-y", "@negokaz/excel-mcp-server"]
  }
}""")
    add_h3(doc, "What it is")
    add_body(doc, "An MCP server for working with Excel files — read/write sheets and cells.")
    add_h3(doc, "Typical Use Cases")
    add_bullet(doc, "Generate Excel reports")
    add_bullet(doc, "Read existing workbooks and summarise / transform data")
    add_h3(doc, "Notes / Gotchas")
    add_warning_bullet(doc, "Ensure the server can access the Excel file paths on your machine.")
    add_separator(doc)

    # ════════════════════════════════════════════════════════════════════════
    # SERVER 6 — testrail
    # ════════════════════════════════════════════════════════════════════════
    add_h2(doc, "6)  testrail")
    add_h3(doc, "Config")
    add_code_block(doc, """\
{
  "testrail": {
    "command": "C:/Users/vseelam/AppData/Local/Python/pythoncore-3.14-64/python.exe",
    "args": ["C:/genai-training/Labs/Lab5/testrail.py"],
    "env": {
      "TESTRAIL_URL": "https://genai.testrail.io",
      "TESTRAIL_EMAIL": "user@example.com",
      "TESTRAIL_API_KEY": "your_api_key_here"
    },
    "type": "stdio"
  }
}""")
    add_h3(doc, "What it is")
    add_body(doc, "A custom Python MCP server script (testrail.py) that wraps the TestRail API.")
    add_h3(doc, "How it runs")
    add_bullet(doc, "Local process via stdio.")
    add_highlight_bullet(doc, "Uses a pinned Python interpreter path — script startup fails if path changes.")
    add_h3(doc, "Configuration")
    add_bullet(doc, "Uses environment variables: TESTRAIL_URL, TESTRAIL_EMAIL, TESTRAIL_API_KEY.")
    add_h3(doc, "Notes / Gotchas / Security")
    add_warning_bullet(doc, "SECURITY: API key is in plain text — prefer secure prompting.")
    add_warning_bullet(doc, "If Python or script path changes, server startup will fail.")
    add_separator(doc)

    # ════════════════════════════════════════════════════════════════════════
    # SERVER 7 — rag-retriever
    # ════════════════════════════════════════════════════════════════════════
    add_h2(doc, "7)  rag-retriever")
    add_h3(doc, "Config")
    add_code_block(doc, """\
{
  "rag-retriever": {
    "type": "stdio",
    "command": "python",
    "args": ["C:\\\\genai-training\\\\rag\\\\retriever_mcp_server.py"]
  }
}""")
    add_h3(doc, "What it is")
    add_body(doc, "A local Python MCP server for RAG (Retrieval-Augmented Generation) — retrieving relevant content from a knowledge base.")
    add_h3(doc, "Typical Use Cases")
    add_bullet(doc, "Query local documents and retrieve relevant passages")
    add_highlight_bullet(doc, "Provide grounded, factual context to LLM answers (reduces hallucinations).")
    add_h3(doc, "Notes / Gotchas")
    add_warning_bullet(doc, "Requires correct Python dependencies installed.")
    add_warning_bullet(doc, "The script path must exist on the machine running the client.")
    add_separator(doc)

    # ════════════════════════════════════════════════════════════════════════
    # SERVER 8 — playwright
    # ════════════════════════════════════════════════════════════════════════
    add_h2(doc, "8)  microsoft/playwright-mcp")
    add_h3(doc, "Config")
    add_code_block(doc, """\
{
  "microsoft/playwright-mcp": {
    "type": "stdio",
    "command": "npx",
    "args": ["@playwright/mcp@latest"]
  }
}""")
    add_h3(doc, "What it is")
    add_body(doc, "An MCP server powered by Playwright for browser automation (modern alternative to Selenium).")
    add_h3(doc, "Typical Use Cases")
    add_bullet(doc, "End-to-end browser automation")
    add_bullet(doc, "Navigation, clicking, screenshots, network inspection")
    add_h3(doc, "Notes / Gotchas")
    add_warning_bullet(doc, "Initial setup may download browser binaries (~hundreds of MB).")
    add_warning_bullet(doc, "Pin a version for long-term stability.")
    add_separator(doc)

    # ════════════════════════════════════════════════════════════════════════
    # SERVER 9 — github
    # ════════════════════════════════════════════════════════════════════════
    add_h2(doc, "9)  io.github.github/github-mcp-server")
    add_h3(doc, "Config")
    add_code_block(doc, """\
{
  "io.github.github/github-mcp-server": {
    "type": "http",
    "url": "https://api.githubcopilot.com/mcp/",
    "gallery": "https://api.mcp.github.com",
    "version": "0.33.0"
  }
}""")
    add_h3(doc, "What it is")
    add_body(doc, "A GitHub-hosted MCP endpoint providing GitHub / Copilot tool access.")
    add_h3(doc, "How it connects")
    add_highlight_bullet(doc, 'type: "http" — remote connection; no local process.')
    add_bullet(doc, "gallery: registry/discovery endpoint (client-dependent).")
    add_bullet(doc, "version: expected MCP server/protocol version (0.33.0).")
    add_h3(doc, "Notes / Gotchas")
    add_warning_bullet(doc, "Requires valid GitHub / Copilot authentication.")
    add_warning_bullet(doc, "Network access required — will not work offline.")
    add_separator(doc)

    # ════════════════════════════════════════════════════════════════════════
    # SERVER 10 — mysql
    # ════════════════════════════════════════════════════════════════════════
    add_h2(doc, "10)  mysql")
    add_h3(doc, "Config")
    add_code_block(doc, """\
{
  "mysql": {
    "command": "npx",
    "args": ["-y", "@sajithrw/mcp-mysql"],
    "env": {
      "MYSQL_HOST": "localhost",
      "MYSQL_USER": "your_user",
      "MYSQL_PASSWORD": "your_password",
      "MYSQL_DATABASE": "your_db"
    },
    "type": "stdio"
  }
}""")
    add_h3(doc, "What it is")
    add_body(doc, "An MCP server exposing MySQL database operations via tools.")
    add_h3(doc, "How it runs")
    add_bullet(doc, "Local process via stdio.")
    add_bullet(doc, "Uses environment variables for DB configuration.")
    add_h3(doc, "Notes / Gotchas / Security")
    add_warning_bullet(doc, "Replace all placeholder values with real credentials before use.")
    add_warning_bullet(doc, "SECURITY: Avoid storing real credentials in plain text.")
    add_separator(doc)

    # ════════════════════════════════════════════════════════════════════════
    # SERVER 11 — filesystem
    # ════════════════════════════════════════════════════════════════════════
    add_h2(doc, "11)  filesystem")
    add_h3(doc, "Config")
    add_code_block(doc, """\
{
  "filesystem": {
    "command": "npx",
    "args": [
      "-y",
      "@modelcontextprotocol/server-filesystem",
      "C:\\\\Prcatices-AI\\\\JSonfile"
    ],
    "type": "stdio"
  }
}""")
    add_h3(doc, "What it is")
    add_body(doc, "An MCP filesystem server providing controlled access to local files under an allowed root directory.")
    add_h3(doc, "How it runs")
    add_bullet(doc, "Local process via stdio.")
    add_highlight_bullet(doc, "Allowed root directory: C:\\Prcatices-AI\\JSonfile  — only files here are accessible.")
    add_h3(doc, "Notes / Gotchas / Security")
    add_warning_bullet(doc, "Keep the allowed root as restrictive as possible to avoid unintended access.")
    add_separator(doc)

    # ════════════════════════════════════════════════════════════════════════
    # SERVER 12 — rest-api
    # ════════════════════════════════════════════════════════════════════════
    add_h2(doc, "12)  rest-api")
    add_h3(doc, "Config")
    add_code_block(doc, """\
{
  "rest-api": {
    "command": "node",
    "args": ["C:/Users/.../dkmaker-mcp-rest-api/build/index.js"],
    "type": "stdio",
    "env": {
      "REST_BASE_URL": "https://engineering.ultrivalms.com/QALS1011/rest/api/2",
      "AUTH_BASIC_USERNAME": "your-username",
      "AUTH_BASIC_PASSWORD": "your-password",
      "AUTH_BEARER": "your-token",
      "AUTH_APIKEY_HEADER_NAME": "X-API-Key",
      "AUTH_APIKEY_VALUE": "your-api-key",
      "REST_ENABLE_SSL_VERIFY": "false",
      "REST_RESPONSE_SIZE_LIMIT": "10000"
    }
  }
}""")
    add_h3(doc, "What it is")
    add_body(doc, "A generic REST API MCP connector (Node-based) that can call internal/external REST endpoints. The base URL resembles a Jira REST API v2 style endpoint.")
    add_h3(doc, "Configuration Highlights")
    add_highlight_bullet(doc, "REST_BASE_URL — root URL for ALL REST requests made by this server.")
    add_bullet(doc, "Auth options: Basic (username/password), Bearer token, API Key header.")
    add_bullet(doc, "Custom headers: X-API-Version, Custom-Client, Accept.")
    add_warning_bullet(doc, 'SECURITY: REST_ENABLE_SSL_VERIFY: "false" disables TLS verification — risky in production.')
    add_bullet(doc, "REST_RESPONSE_SIZE_LIMIT — caps response body size (10 000 chars here).")
    add_separator(doc)

    # ════════════════════════════════════════════════════════════════════════
    # INPUTS SECTION
    # ════════════════════════════════════════════════════════════════════════
    add_h2(doc, "Inputs (Secrets Prompting)")
    add_body(doc,
        'The inputs section declares secret values your client will prompt for and then substitute '
        'using ${input:<ID>}.  Two inputs are declared:'
    )
    add_highlight_bullet(doc, "POSTMAN_API_KEY — prompted securely; used by postman-mcp-server.")
    add_highlight_bullet(doc, "ATLASSIAN_API_TOKEN — prompted securely; used by atlassian-mcp-server.")
    add_separator(doc)

    # ════════════════════════════════════════════════════════════════════════
    # BEST PRACTICES
    # ════════════════════════════════════════════════════════════════════════
    add_h2(doc, "Recommendations (Best Practices)")
    add_highlight_bullet(doc, "Pin versions — use @x.y.z instead of @latest to prevent unexpected breakage.")
    add_highlight_bullet(doc, "No plain-text credentials — use prompted inputs (${input:...}) or a secrets manager.")
    add_highlight_bullet(doc, "Filesystem server roots — restrict to the minimum required directory only.")
    add_highlight_bullet(doc, "SSL verification — never disable in production (REST_ENABLE_SSL_VERIFY).")
    add_highlight_bullet(doc, "Keep MCP config out of source control if it contains any real credentials.")

    # ════════════════════════════════════════════════════════════════════════
    # SUMMARY TABLE
    # ════════════════════════════════════════════════════════════════════════
    doc.add_paragraph()
    add_h2(doc, "Quick Reference Summary")
    headers = ["#", "Server", "Type", "Transport", "Auth Method"]
    rows = [
        ("1",  "selenium",                    "Browser automation",          "stdio", "None"),
        ("2",  "postman-mcp-server",          "Postman API",                 "stdio", "API Key (prompted)"),
        ("3",  "postgres",                    "PostgreSQL DB",               "stdio", "Conn string (plain)"),
        ("4",  "atlassian-mcp-server",        "Jira / Confluence",           "http",  "Bearer (prompted)"),
        ("5",  "excel",                       "Excel files",                 "stdio", "None"),
        ("6",  "testrail",                    "TestRail API",                "stdio", "API Key (env var)"),
        ("7",  "rag-retriever",               "RAG knowledge base",          "stdio", "None"),
        ("8",  "playwright-mcp",              "Browser automation",          "stdio", "None"),
        ("9",  "github-mcp-server",           "GitHub / Copilot",            "http",  "GitHub auth"),
        ("10", "mysql",                       "MySQL DB",                    "stdio", "Env vars (plain)"),
        ("11", "filesystem",                  "Local file access",           "stdio", "None"),
        ("12", "rest-api",                    "Generic REST calls",          "stdio", "Basic/Bearer/API Key"),
    ]

    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light List Accent 1"

    # Header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for para in hdr_cells[i].paragraphs:
            for run in para.runs:
                run.bold = True

    # Data rows
    for r_idx, row_data in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            cells[c_idx].text = val

    doc.save(OUTPUT_PATH)
    print(f"DOCX saved to: {OUTPUT_PATH}")


if __name__ == "__main__":
    build_document()
