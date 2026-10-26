"""Update test_data.xlsx: add Actual_Result + Notes columns; flag security findings."""
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

thin   = Side(style="thin")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
PASS_BG = "C6EFCE"
FAIL_BG = "FFC7CE"
INFO_BG = "FFE699"

def fill(h): return PatternFill(start_color=h, end_color=h, fill_type="solid")
def hdr():   return Font(bold=True, color="FFFFFF", size=11, name="Calibri")
def cell():  return Font(size=10, name="Calibri")
def left():  return Alignment(horizontal="left", vertical="center", wrap_text=True)
def center():return Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = load_workbook("test_data/test_data.xlsx")

security_findings = {
    "TC-010": ("Login ACCEPTED (case-insensitive username)",
               "SECURITY FINDING: App accepts uppercase/mixed-case usernames — "
               "username should be case-sensitive."),
    "TC-306": ("Login ACCEPTED (whitespace-only password)",
               "SECURITY FINDING: App accepts whitespace-only passwords without "
               "server-side trim/validation."),
    "TC-308": ("Login ACCEPTED (null-byte stripped)",
               "SECURITY FINDING: Server strips null byte and authenticates — "
               "input should be rejected before auth check."),
}


def annotate_sheet(ws, status_col_idx):
    """Add Actual_Result and Notes columns, then fill per-row."""
    # Ensure header columns exist
    for col_idx, header in [(status_col_idx + 1, "Actual_Result"),
                             (status_col_idx + 2, "Notes")]:
        c = ws.cell(row=1, column=col_idx, value=header)
        c.font = hdr(); c.fill = fill("1F4E79")
        c.alignment = center(); c.border = border
        ws.column_dimensions[c.column_letter].width = 40 if header == "Actual_Result" else 60

    for row in ws.iter_rows(min_row=2):
        tc_id = row[0].value
        if not tc_id:
            continue
        ar_cell   = ws.cell(row=row[0].row, column=status_col_idx + 1)
        note_cell = ws.cell(row=row[0].row, column=status_col_idx + 2)
        ar_cell.font = cell();   ar_cell.alignment = left();   ar_cell.border = border
        note_cell.font = cell(); note_cell.alignment = left(); note_cell.border = border

        if tc_id in security_findings:
            actual, note = security_findings[tc_id]
            row[status_col_idx - 1].value = "FAIL"
            row[status_col_idx - 1].fill  = fill(FAIL_BG)
            row[status_col_idx - 1].font  = Font(bold=True, size=10, name="Calibri")
            ar_cell.value = actual;  ar_cell.fill = fill(FAIL_BG)
            note_cell.value = note;  note_cell.fill = fill(INFO_BG)
        else:
            ar_cell.value = "As expected";  ar_cell.fill = fill(PASS_BG)
            note_cell.value = ""


# Login_TestData: status column = 6
annotate_sheet(wb["Login_TestData"], status_col_idx=6)

# EdgeCase_TestData: status column = 6
annotate_sheet(wb["EdgeCase_TestData"], status_col_idx=6)

wb.save("test_data/test_data.xlsx")
print("Updated test_data/test_data.xlsx with security findings.")
for tc, (actual, _) in security_findings.items():
    print(f"  {tc}: {actual}")
