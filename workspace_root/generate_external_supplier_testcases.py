"""
Generate External Supplier Test Cases DOCX (Short Form)
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import docx.opc.constants

# ── Helpers ───────────────────────────────────────────────────────────────────

def set_col_widths(table, widths):
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                cell.width = Inches(widths[i])

def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

def header_row(table, texts, bg="1F4E79", fg="FFFFFF", font_size=10):
    row = table.rows[0]
    for i, text in enumerate(texts):
        cell = row.cells[i]
        shade_cell(cell, bg)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor.from_string(fg)

def add_data_row(table, values, bg=None, bold_first=False, center_all=False):
    row = table.add_row()
    for i, val in enumerate(values):
        cell = row.cells[i]
        if bg:
            shade_cell(cell, bg)
        p = cell.paragraphs[0]
        if center_all:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(val))
        run.font.size = Pt(9)
        if bold_first and i == 0:
            run.bold = True
    return row

def section_title(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13 if level == 1 else 11)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)

def sub_note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    p.paragraph_format.space_after = Pt(4)

# ── Document ──────────────────────────────────────────────────────────────────

doc = Document()

# Page margins
section = doc.sections[0]
section.left_margin   = Inches(0.8)
section.right_margin  = Inches(0.8)
section.top_margin    = Inches(0.7)
section.bottom_margin = Inches(0.7)

# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("External Supplier — Test Cases (Short Form)")
run.bold = True
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
mr = meta.add_run("App: https://engineering.ultrivalms.com/QALS1011  |  Login: buyeruser / Ultriva!23")
mr.font.size = Pt(9)
mr.font.color.rgb = RGBColor(0x50, 0x50, 0x50)
doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# 1. TestExternalSupplier  (7 tests)
# ══════════════════════════════════════════════════════════════════════════════
section_title(doc, "1. TestExternalSupplier  (7 Tests)")
sub_note(doc, "File: qa_automation/admin/ExternalSupplier/ExternalSupplier.py")
sub_note(doc, "Scope: End-to-end — Add Supplier → Edit → View → Add Supplier User")

tbl = doc.add_table(rows=1, cols=4)
tbl.style = "Table Grid"
header_row(tbl, ["#", "Test Name", "Phase", "Description"])
set_col_widths(tbl, [0.3, 2.0, 1.2, 3.0])

tests_ext = [
    ("1", "test_01_login_and_setup",          "Login & Setup",      "Log in as buyeruser; open Admin module"),
    ("2", "test_02_add_supplier",             "Add Supplier",       "Open Supplier/Create; fill Name, Code, Address, Label, Timezone; save & verify in list"),
    ("3", "test_03_edit_supplier",            "Edit Supplier",      "Open Edit for created supplier; update Address, City, State, Zip, Country, DUNS; save"),
    ("4", "test_04_view_supplier",            "View Supplier",      "Open View for supplier; verify Address Info, Configuration, and Other Details sections"),
    ("5", "test_05_add_supplier_user_form",   "Add User — Nav",     "Capture supplier GID from list; navigate to User/CreateS?BusGID={gid}"),
    ("6", "test_06_fill_and_save_supplier_user", "Fill & Save User", "Fill First/Last Name, Phone, Email, Account Type, User ID, Password; save user"),
    ("7", "test_07_summary_report",           "Summary Report",     "Save screenshot; write full text report with supplier & user details"),
]

row_colors = ["FFFFFF", "EBF3FF"]
for i, (num, name, phase, desc) in enumerate(tests_ext):
    bg = row_colors[i % 2]
    add_data_row(tbl, [num, name, phase, desc], bg=bg)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# 2. TestSupplierUserCreation  (10 tests) — lives under Supplier/supplieruser
# ══════════════════════════════════════════════════════════════════════════════
section_title(doc, "2. TestSupplierUserCreation  (10 Tests)")
sub_note(doc, "File: qa_automation/tests/QALS11/Admin/Supplier/supplieruser/SupplierUserCreation.py")
sub_note(doc, "Scope: Create Supplier Administrator user for External Supplier \"TestSup197\"")

tbl2 = doc.add_table(rows=1, cols=4)
tbl2.style = "Table Grid"
header_row(tbl2, ["#", "Test Name", "Phase", "Description"])
set_col_widths(tbl2, [0.3, 2.4, 1.2, 2.6])

tests_usr = [
    ("1",  "test_01_login",                "Login",         "Log in to QA site with buyeruser credentials"),
    ("2",  "test_02_open_admin",           "Admin Nav",     "Click Admin module link"),
    ("3",  "test_03_click_supplier_menu",  "Supplier Menu", "Click Supplier menu → navigate to /Supplier/List"),
    ("4",  "test_04_select_supplier",      "Select Supp.",  "Select External Supplier row \"TestSup197\" from list"),
    ("5",  "test_05_click_add_user",       "Add User",      "Click 'Add User' button on the supplier page"),
    ("6",  "test_06_fill_user_form",       "Fill Form",     "Enter First Name, Last Name, Work Phone, Email in Create User form"),
    ("7",  "test_07_set_account_and_userid","Account & ID", "Set Account Type = Supplier Administrators; enter unique User ID"),
    ("8",  "test_08_set_password",         "Password",      "Enter password meeting all complexity rules (>8 chars, upper, lower, digit, symbol)"),
    ("9",  "test_09_save_user",            "Save",          "Submit the Create User form and assert success URL"),
    ("10", "test_10_show_results",         "Verify",        "Verify new user appears in the list and print results"),
]

for i, (num, name, phase, desc) in enumerate(tests_usr):
    bg = row_colors[i % 2]
    add_data_row(tbl2, [num, name, phase, desc], bg=bg)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# 3. Summary Table
# ══════════════════════════════════════════════════════════════════════════════
section_title(doc, "Summary")

tbl3 = doc.add_table(rows=1, cols=4)
tbl3.style = "Table Grid"
header_row(tbl3, ["Test Class", "File", "Tests", "Scope"])
set_col_widths(tbl3, [2.0, 2.8, 0.5, 2.2])

summary_rows = [
    ("TestExternalSupplier",
     "admin/ExternalSupplier/ExternalSupplier.py",
     "7",
     "Add / Edit / View Supplier + Add User"),
    ("TestSupplierUserCreation",
     "tests/.../supplieruser/SupplierUserCreation.py",
     "10",
     "Create Supplier Admin user for TestSup197"),
]

for i, row_data in enumerate(summary_rows):
    bg = row_colors[i % 2]
    add_data_row(tbl3, row_data, bg=bg)

# Total row
total_row = tbl3.add_row()
shade_cell(total_row.cells[0], "D6E4F0")
shade_cell(total_row.cells[1], "D6E4F0")
shade_cell(total_row.cells[2], "D6E4F0")
shade_cell(total_row.cells[3], "D6E4F0")
for i, val in enumerate(["Total", "", "17", ""]):
    p = total_row.cells[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(val)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

# ── Save ──────────────────────────────────────────────────────────────────────
out = r"c:\Prcatices-AI\qa_automation\admin\ExternalSupplier\ExternalSupplier_TestCases.docx"
doc.save(out)
print(f"Saved: {out}")
