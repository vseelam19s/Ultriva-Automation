from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
title = doc.add_heading("QA Automation – Test Cases Progress", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# --- Module 1: Workcenter ---
doc.add_heading("1. Workcenter", level=1)
cases = [
    "Login & Navigate to Workcenter",
    "Create New Workcenter",
    "Edit Workcenter",
    "Navigate to Workcenter Users",
    "Add Workcenter User",
    "Verify User in Workcenter List",
    "Edit Workcenter User",
    "View Workcenter User",
    "Disable and Enable Workcenter User",
    "Generate Workcenter Users Report",
]
for i, c in enumerate(cases, 1):
    doc.add_paragraph(f"{i}. {c}", style="List Number")

# --- Module 2: User Group (Local) ---
doc.add_heading("2. User Group (Local)", level=1)
cases = [
    "Create Local User Group",
    "Edit Local User Group",
    "View Local User Group",
    "Assign Users to User Group",
    "My Routines",
    "Clone User Group",
    "Delete User Group",
]
for i, c in enumerate(cases, 1):
    doc.add_paragraph(f"{i}. {c}", style="List Number")

# --- Module 3: Users ---
doc.add_heading("3. Users", level=1)
cases = [
    "Add User",
    "Edit User",
    "View User",
    "Disable User",
    "Enable User",
]
for i, c in enumerate(cases, 1):
    doc.add_paragraph(f"{i}. {c}", style="List Number")

# --- Module 4: Supplier ---
doc.add_heading("4. Supplier", level=1)
cases = [
    "Add Supplier",
    "Edit Supplier",
    "View Supplier",
    "Add Supplier User",
    "Summary Report",
]
for i, c in enumerate(cases, 1):
    doc.add_paragraph(f"{i}. {c}", style="List Number")

output = "qa_automation/reports/TestCases_Progress.docx"
doc.save(output)
print(f"Saved: {output}")
