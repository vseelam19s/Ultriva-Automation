"""Generate ULT-10612 (Steel Origin | End-to-end QA and Copeland barcode scan validation) test cases DOCX."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = Path(r"c:\Prcatices-AI\docs\ULT-10612_test_cases.docx")

COLS = ["ID", "Title", "Preconditions", "Steps", "Expected Result", "Pri", "Type", "Story"]
WIDTHS = [0.55, 1.35, 1.55, 2.6, 2.3, 0.4, 0.7, 0.55]

ENV = "10.10.6 QA site configured like Copeland (Emerson print plug-in, Copeland .lbl)."
FLAGGED = "Item X flagged for Steel Origin for Supplier A; open kanban cards available."
SHIP_SCREENS = ["Ship Cards", "Ship by Boxes", "Consolidated Ship"]


n = 0


def nxt():
    global n
    n += 1
    return f"TC-{n:02d}"


def mk(rows):
    return [(nxt(), *r) for r in rows]


SECTIONS = []

SECTIONS.append(("1. Environment and Pre-Release Baseline (capture BEFORE deploying)", mk([
    ("QA site configured like Copeland",
     "10.10.6 QA build available.",
     "1. Set SysConfig CUSTOM_PRINT_ASSEMBLY_NAME = Leansuite.Custom.Emerson:CustomPrintScanSvc.PrintData.\n2. Install Copeland's updated .lbl.\n3. Configure the desktop print client against the site.\n4. Confirm STEEL_ORIGIN_DI default.",
     "Site prints through the Emerson plug-in with the Copeland .lbl; desktop print client connects; build number is 10.10.6.", "High", "Setup", "All"),
    ("Capture pre-release baseline - label image",
     "Site still on the pre-release build. Unflagged supplier / item with a shipped card.",
     "1. Print the label for the card through the desktop print client.\n2. Save the print output (raw print stream / file) and a PDF/image.\n3. Record item, supplier, card, cycle, weight and print options used.",
     "Baseline label output saved and archived with a checksum (SHA-256) before deployment.", "High", "Setup", "Feature 3"),
    ("Capture pre-release baseline - PX2dBarCode",
     "As above.",
     "1. Print the label with the PX2dBarCode variant.\n2. Scan the 2D barcode into a hex viewer.\n3. Save the hex dump and the raw print stream.",
     "PX2dBarCode baseline bytes saved and checksummed.", "High", "Setup", "Feature 3"),
    ("Capture pre-release baseline - PXIR2dBarCode",
     "As above.",
     "1. Print the label with the PXIR2dBarCode variant.\n2. Scan the 2D barcode into a hex viewer.\n3. Save the hex dump and the raw print stream.",
     "PXIR2dBarCode baseline bytes saved and checksummed.", "High", "Setup", "Feature 3"),
    ("Deploy 10.10.6 and run DB scripts",
     "Baseline artifacts (TC-02 to TC-04) saved.",
     "1. Deploy the 10.10.6 build.\n2. Run the DB scripts.\n3. Verify app starts and login works.",
     "Deployment succeeds without errors; baseline remains untouched.", "High", "Setup", "All"),
    ("Test data prepared",
     ENV,
     "1. Create Supplier A and Supplier B.\n2. Create items X (flagged for A), Y (not flagged), plus X for Supplier B.\n3. Generate kanban cards (single, multi, boxed, consolidated).",
     "Cards available for every scenario in this plan.", "Medium", "Setup", "All"),
])))

SECTIONS.append(("2. Configuration (ULT-10605)", mk([
    ("Flag defaults to off",
     ENV + " New item / supplier relationship.",
     "1. Open the supplier item configuration for a new item.\n2. Check the Steel Origin flag.",
     "Flag is off (unchecked) by default.", "High", "Functional", "ULT-10605"),
    ("Enable flag for item X + Supplier A",
     ENV,
     "1. Open item X / Supplier A configuration.\n2. Enable the Steel Origin flag and save.\n3. Reopen the record.",
     "Flag saved and shown as enabled.", "High", "Functional", "ULT-10605"),
    ("Enabling does not affect Supplier B",
     "Item X flagged for Supplier A.",
     "1. Open item X / Supplier B.\n2. Check the flag.\n3. Ship a card of X for Supplier B.",
     "Flag remains off for Supplier B; no Steel Origin requirement on ship.", "High", "Functional", "ULT-10605"),
    ("Enabling does not affect other items",
     "Item X flagged for Supplier A.",
     "1. Open item Y / Supplier A.\n2. Check the flag.\n3. Ship a card of Y.",
     "Flag remains off for item Y; no requirement on ship.", "High", "Functional", "ULT-10605"),
    ("Disable flag removes the requirement",
     "Item X flagged for Supplier A.",
     "1. Disable the flag for X / A and save.\n2. Open the ship screens for an X card.",
     "No Steel Origin dropdown; ship is not blocked.", "High", "Functional", "ULT-10605"),
    ("Flagged-supplier-items query returns exactly flagged rows",
     "A mix of flagged and unflagged items across suppliers (including a re-disabled one).",
     "1. Run the flagged-supplier-items query / API against the DB.\n2. Compare the result with the configuration data.",
     "Result contains exactly the flagged supplier-item rows; no unflagged, disabled or other-supplier rows; no duplicates.", "High", "DB", "ULT-10605"),
    ("Flag persistence after re-login / cache refresh",
     "Item X flagged.",
     "1. Log out and in.\n2. Recycle app pool / refresh cache.\n3. Check the flag.",
     "Flag value persists.", "Medium", "Functional", "ULT-10605"),
    ("Flag configuration access by role",
     "Users with different roles.",
     "1. Open the configuration as buyer, supplier and admin.",
     "Behaviour matches the role design; unauthorised users cannot change the flag.", "Medium", "Security", "ULT-10605"),
])))

ship = []
ship += [("Dropdown appears only on flagged rows - " + s,
          FLAGGED + " Mixed shipment with flagged and unflagged cards.",
          f"1. Open {s} with flagged and unflagged cards selected.\n2. Inspect the Steel Origin column.",
          "Dropdown shown only on flagged rows; unflagged rows have no control.", "High", "Functional", "ULT-10606") for s in SHIP_SCREENS]
ship += [("Dropdown options US / AA, none preselected - " + s,
          FLAGGED,
          f"1. Open {s} with a flagged card.\n2. Inspect the dropdown before selecting.\n3. Open the list.",
          "Options are exactly US and AA; nothing is preselected (blank / placeholder).", "High", "Functional", "ULT-10606") for s in SHIP_SCREENS]
ship += [("Column and header hidden when no row is flagged - " + s,
          "Only unflagged cards selected.",
          f"1. Open {s} with unflagged cards only.\n2. Inspect the grid header and rows.",
          "Steel Origin column and header are hidden.", "High", "Functional", "ULT-10606") for s in SHIP_SCREENS]
ship += [("Ship blocked without value (client) - " + s,
          FLAGGED,
          f"1. Open {s} with a flagged card and leave Steel Origin empty.\n2. Click Ship.",
          "Client validation message; shipment is not submitted.", "High", "Negative", "ULT-10606") for s in SHIP_SCREENS]
ship += [("Ship blocked without value (server, client check bypassed) - " + s,
          FLAGGED + " Browser dev tools.",
          f"1. Open {s}, leave Steel Origin empty.\n2. Using dev tools, disable/remove the client validation (or replay the request without the value).\n3. Submit.",
          "Server rejects the request with an error; card is not shipped and no cycle data is written.", "High", "Security", "ULT-10606/10607") for s in SHIP_SCREENS]
ship += [("Ship blocked with invalid value (tampered) - " + s,
          FLAGGED + " Browser dev tools.",
          f"1. Open {s}.\n2. Using dev tools, post steelorigin = 'XX' (and lowercase 'us').\n3. Submit.",
          "Server rejects values other than US / AA; nothing is shipped.", "High", "Security", "ULT-10607") for s in SHIP_SCREENS]
ship += [("Ship with US - " + s,
          FLAGGED,
          f"1. Select US for the flagged card in {s}.\n2. Ship.\n3. Check the cycle data (IGParams).",
          "Ship succeeds; steelorigin=US stored for the cycle.", "High", "Functional", "ULT-10607") for s in SHIP_SCREENS]
ship += [("Ship with AA - " + s,
          FLAGGED,
          f"1. Select AA for the flagged card in {s}.\n2. Ship.\n3. Check the cycle data (IGParams).",
          "Ship succeeds; steelorigin=AA stored for the cycle.", "High", "Functional", "ULT-10607") for s in SHIP_SCREENS]
ship += [("Apply-to-all affects only flagged checked rows - " + s,
          FLAGGED + " Several flagged and unflagged cards, some unchecked.",
          f"1. In {s} check some flagged and unflagged cards, leave one flagged card unchecked.\n2. Use apply-to-all with AA.",
          "Only checked flagged rows get AA; unchecked flagged rows and unflagged rows unchanged.", "High", "Functional", "ULT-10606") for s in SHIP_SCREENS]
ship += [("Mixed shipment - all cards ship, unflagged has no value - " + s,
          FLAGGED + " One flagged and one unflagged card.",
          f"1. Select both cards in {s}; choose US for the flagged card.\n2. Ship.\n3. Check cycle data of both cards.",
          "Both cards ship; flagged card has US; unflagged card has no steelorigin.", "High", "Functional", "ULT-10607") for s in SHIP_SCREENS]
ship += [("Partial ship works - " + s,
          FLAGGED + " Card quantity > 1.",
          f"1. Ship a partial quantity in {s} with Steel Origin selected.\n2. Verify card state and cycle data.\n3. Ship the remainder.",
          "Partial ship succeeds; value stored; remainder can be shipped with its own selection.", "High", "Functional", "ULT-10607") for s in SHIP_SCREENS]
ship += [("Value differs per cycle - " + s,
          FLAGGED,
          f"1. Ship cycle 1 with US in {s} and receive.\n2. Ship cycle 2 with AA.\n3. Check cycle data for each cycle.",
          "Each cycle keeps its own value (US, AA); cycle 1 is not overwritten.", "High", "Functional", "ULT-10607") for s in SHIP_SCREENS]
ship += [("Charge No is preserved - " + s,
          FLAGGED + " Charge No entered.",
          f"1. Enter Charge No and select Steel Origin in {s}.\n2. Ship.\n3. Check Charge No in cycle data, card history and label.",
          "Charge No stored exactly as entered; unaffected by the Steel Origin value in IGParams.", "High", "Regression", "ULT-10607") for s in SHIP_SCREENS]
ship += [
    ("Selection retained after validation failure",
     FLAGGED + " Two flagged cards.",
     "1. Select a value for card 1 only.\n2. Click Ship (blocked).\n3. Review the grid.",
     "Selected value for card 1 retained; card 2 highlighted as missing.", "Medium", "Functional", "ULT-10606"),
    ("Scan / EDI ship of a flagged item without value",
     FLAGGED,
     "1. Ship a flagged card through scan or EDI.\n2. Check cycle data.",
     "Behaviour per design (shipped without value; can be set later through the correction page - see section 6).", "Medium", "Functional", "ULT-10607/10610"),
]
SECTIONS.append(("3. Ship (ULT-10606 / ULT-10607) - Ship Cards, Ship by Boxes, Consolidated Ship", mk(ship)))

SECTIONS.append(("4. Unaffected Supplier (Feature 3)", mk(
    [("No new UI on ship screen - " + s,
      "Supplier / item not flagged.",
      f"1. Open {s} for unflagged cards.\n2. Compare with the pre-release build screenshots.",
      "No Steel Origin column, dropdown, header or message; layout identical.", "High", "Regression", "Feature 3") for s in SHIP_SCREENS] + [
    ("Label byte-identical to baseline",
     "Baseline from TC-02. Same card data, weight, options.",
     "1. Ship the same unflagged item and print via the desktop print client.\n2. Compute SHA-256 of the raw print stream.\n3. Compare with baseline (allowing only documented dynamic fields such as date/serial, normalised).",
     "Print stream is byte-identical to the baseline.", "High", "Regression", "Feature 3"),
    ("PX2dBarCode byte-identical to baseline",
     "Baseline from TC-03.",
     "1. Print the PX2dBarCode variant for the unflagged card.\n2. Scan into a hex viewer; compare with the baseline hex.",
     "2D barcode bytes are identical to the baseline; no 5Z segment.", "High", "Regression", "Feature 3"),
    ("PXIR2dBarCode byte-identical to baseline",
     "Baseline from TC-04.",
     "1. Print the PXIR2dBarCode variant for the unflagged card.\n2. Scan into a hex viewer; compare with the baseline hex.",
     "2D barcode bytes are identical to the baseline; no 5Z segment.", "High", "Regression", "Feature 3"),
    ("Flagged item for another supplier not affected",
     "Item X flagged for A only.",
     "1. Ship and print X for Supplier B.",
     "No UI change, label and 2D barcodes same as baseline.", "High", "Regression", "Feature 3"),
    ("Unflagged card history / reports unchanged",
     "Unflagged shipments exist.",
     "1. Open card history and the correction page for an unflagged card.",
     "No Steel Origin values displayed / required.", "Medium", "Regression", "Feature 3"),
])))

SECTIONS.append(("5. Label (ULT-10608 / ULT-10609) - desktop print client", mk([
    ("Label text 'Steel Origin: US'",
     "Flagged card shipped with US. Printed via the desktop print client.",
     "1. Print the label.\n2. Inspect the printed label.",
     "Text 'Steel Origin: US' is printed legibly in the designed position.", "High", "Functional", "ULT-10608"),
    ("Label text 'Steel Origin: AA'",
     "Flagged card shipped with AA.",
     "1. Print the label.",
     "Text 'Steel Origin: AA' printed.", "High", "Functional", "ULT-10608"),
    ("Bottom-left Code128 linear barcode 5ZUS",
     "Card with US.",
     "1. Print the label.\n2. Locate the bottom-left linear barcode.\n3. Scan it.",
     "Code128 barcode present at the bottom-left; scans as 5ZUS.", "High", "Functional", "ULT-10608"),
    ("Bottom-left Code128 linear barcode 5ZAA",
     "Card with AA.",
     "1. Print the label.\n2. Scan the linear barcode.",
     "Scans as 5ZAA.", "High", "Functional", "ULT-10608"),
    ("PX2dBarCode ends with GS 5Z<value> RS EOT",
     "Card with US (repeat with AA); label with PX2dBarCode.",
     "1. Print the label.\n2. Scan the 2D barcode into a hex viewer.\n3. Inspect the tail of the payload.",
     "Payload ends with 0x1D '5ZUS' 0x1E 0x04 (GS 5Z<value> RS EOT).", "High", "Functional", "ULT-10609"),
    ("PXIR2dBarCode ends with GS 5Z<value> RS EOT",
     "Card with US (repeat with AA); label with PXIR2dBarCode.",
     "1. Print the label.\n2. Scan into a hex viewer.\n3. Inspect the tail.",
     "Payload ends with 0x1D '5Z<value>' 0x1E 0x04.", "High", "Functional", "ULT-10609"),
    ("Existing segments P, Q, V, S, K, 5K unchanged and in order",
     "Flagged card; baseline 2D of the same card without Steel Origin (or the pre-release layout).",
     "1. Decode both 2D variants in the hex viewer.\n2. Compare the segment order and content of P, Q, V, S, K and 5K with the baseline.",
     "P, Q, V, S, K, 5K values and order identical; only the 5Z segment is added before the trailer.", "High", "Regression", "ULT-10609"),
    ("Real control bytes 0x1E / 0x1D / 0x04 (not literal text)",
     "Both 2D variants printed.",
     "1. Scan into a hex viewer.\n2. Check that separators are single bytes 0x1D, 0x1E and 0x04.\n3. Make sure there are no literal strings like '<GS>', '\\x1D', '^^' or '[)>'-style substitutions.",
     "Real 0x1E / 0x1D / 0x04 bytes present in both variants.", "High", "Functional", "ULT-10609"),
    ("No value: no text",
     "Flagged card without Steel Origin (scan/EDI ship).",
     "1. Print the label.",
     "No 'Steel Origin' text, not even an empty caption.", "High", "Negative", "ULT-10608"),
    ("No value: no linear barcode glyph",
     "As above.",
     "1. Inspect the bottom-left area.",
     "No Code128 barcode or empty bars/quiet-zone artefact.", "High", "Negative", "ULT-10608"),
    ("No value: 2D equals baseline (both variants)",
     "As above; baseline for the same data.",
     "1. Print both PX2dBarCode and PXIR2dBarCode.\n2. Compare hex with the baseline.",
     "2D barcodes equal the baseline; no GS 5Z segment.", "High", "Regression", "ULT-10609"),
    ("SysConfig STEEL_ORIGIN_DI change - no rebuild",
     "Card with US. Admin access to SysConfig.",
     "1. Note the current DI in the 2D payload (default 5Z).\n2. Change STEEL_ORIGIN_DI to another value (e.g. 6Z) and save.\n3. Reprint without rebuilding / redeploying (refresh cache if required).\n4. Check text, linear barcode and both 2D variants.",
     "The DI changes in the linear barcode and both 2D barcodes (e.g. 6ZUS); no rebuild needed. Restore the original after the test.", "High", "Functional", "ULT-10609"),
    ("STEEL_ORIGIN_DI invalid / empty value",
     "Admin access.",
     "1. Set STEEL_ORIGIN_DI empty or invalid.\n2. Print.",
     "Behaviour per design (falls back to default or is rejected); the label does not break and no exception is logged.", "Medium", "Negative", "ULT-10609"),
    ("Reprint after ship - same value",
     "Card shipped with AA and printed.",
     "1. Reprint the label from the ship / print screen.",
     "Same value, text and barcodes as the original print.", "High", "Functional", "ULT-10608"),
    ("Reprint from card history - same value",
     "As above.",
     "1. Open card history.\n2. Reprint the label for the cycle.",
     "Label shows the cycle's value (AA) in text, linear and both 2D barcodes.", "High", "Functional", "ULT-10608/10611"),
    ("Reprint of earlier cycle shows that cycle's value",
     "Cycle 1 shipped US, cycle 2 shipped AA.",
     "1. Reprint cycle 1 and cycle 2 from card history.",
     "Cycle 1 prints US; cycle 2 prints AA.", "High", "Functional", "ULT-10608"),
    ("Multi-card print: correct item weight per card (row-0 fix)",
     "Several cards in one print job with different items and different weights, mixed flagged/unflagged.",
     "1. Select multiple cards and print them in one job.\n2. For each label verify the weight in the text and in the 2D payload (Q/K/etc. per design).\n3. Compare card 1 vs rows 2..n.",
     "Every label uses its own item's weight, not row 0's; Steel Origin value is also per card.", "High", "Regression", "ULT-10608"),
    ("Multi-card print: mixed values",
     "Cards with US, AA and no value in one job.",
     "1. Print all in one job.",
     "Each label has the correct text/barcodes; the unflagged card has none.", "High", "Functional", "ULT-10608"),
    ("Label layout / print quality",
     "Printer used by Copeland (or equivalent).",
     "1. Print on the real label stock.\n2. Check overlap, alignment, quiet zone and scannability of the linear barcode with a handheld scanner.",
     "No overlap with existing fields; barcode scans reliably on first attempt.", "Medium", "Usability", "ULT-10608"),
    ("Emerson plug-in path is used",
     "CUSTOM_PRINT_ASSEMBLY_NAME configured.",
     "1. Print the label.\n2. Check the print client / plug-in log for CustomPrintScanSvc.PrintData.",
     "Print goes through the Emerson plug-in; Steel Origin data is passed to it.", "High", "Integration", "ULT-10608/10609"),
    ("Labels in other print paths (non desktop client)",
     "Web/PDF print if available.",
     "1. Print a flagged card using the other print path.",
     "Behaviour documented; no exception. (Primary acceptance is via the desktop print client.)", "Low", "Functional", "ULT-10608"),
])))

SECTIONS.append(("6. Correction (ULT-10610)", mk([
    ("Menu item visible to the same users as Set Tracking/Packing Slip Numbers",
     "Users with and without the Set Tracking/Packing Slip Numbers permission.",
     "1. Log in as each user type.\n2. Check the menu for the new Steel Origin correction item.",
     "New menu item is visible to exactly the same users as Set Tracking/Packing Slip Numbers; hidden for others.", "High", "Security", "ULT-10610"),
    ("Direct URL access without permission",
     "User without permission.",
     "1. Open the new page URL directly.",
     "Access denied / redirect; no data shown.", "High", "Security", "ULT-10610"),
    ("Old page unchanged",
     "Set Tracking/Packing Slip Numbers page.",
     "1. Open the old page and compare with the pre-release behaviour.\n2. Update a tracking and a packing slip number.",
     "Page layout and behaviour are unchanged; no Steel Origin control.", "High", "Regression", "ULT-10610"),
    ("Correct US to AA, then reprint",
     "Card shipped with US.",
     "1. Open the new page and find the card.\n2. Change US to AA and save.\n3. Reprint the label.",
     "Text 'Steel Origin: AA'; linear barcode 5ZAA; both 2D variants end with GS 5ZAA RS EOT.", "High", "Functional", "ULT-10610"),
    ("Correct AA to US, then reprint",
     "Card shipped with AA.",
     "1. Correct AA to US.\n2. Reprint.",
     "Text and all barcodes show US.", "High", "Functional", "ULT-10610"),
    ("EboKanbanHistory keeps previous version with US, ActionBy and time",
     "Correction US to AA done.",
     "1. Query EboKanbanHistory for the card/cycle.\n2. Review the rows.",
     "Previous version retains US; new version has AA; ActionBy = correcting user; action time recorded.", "High", "DB", "ULT-10610"),
    ("Unflagged rows behave like the old page",
     "Unflagged card.",
     "1. Search for an unflagged card on the new page.\n2. Edit the tracking / packing slip fields.",
     "Row behaves as on the old page; no Steel Origin control.", "High", "Regression", "ULT-10610"),
    ("Set missing value from a scan / EDI ship",
     "Flagged card shipped via scan or EDI without value.",
     "1. Open the card on the new page.\n2. Select AA and save.\n3. Reprint and check card history.",
     "Value saved; label and card history show AA; history entry created.", "High", "Functional", "ULT-10610"),
    ("Invalid value rejected on the server",
     "Browser dev tools.",
     "1. Post an invalid Steel Origin value from the correction page.",
     "Server rejects values other than US / AA.", "High", "Security", "ULT-10610"),
    ("Cannot clear the value on a flagged row",
     "Flagged card with value.",
     "1. Try to save the row with the value blank.",
     "Blocked per design with a message; existing value stays.", "Medium", "Negative", "ULT-10610"),
    ("Correct one cycle does not affect other cycles",
     "Card with several cycles.",
     "1. Correct the value on cycle 2.\n2. Check cycle 1.",
     "Only the chosen cycle changes.", "High", "Functional", "ULT-10610"),
    ("Charge No and other fields preserved on correction",
     "Card with Charge No, tracking number.",
     "1. Correct the Steel Origin.\n2. Compare the other fields.",
     "Charge No, tracking and packing slip numbers are unchanged.", "Medium", "Regression", "ULT-10610"),
    ("Localized labels in a non-English locale",
     "User with a non-English language (e.g. Spanish / French).",
     "1. Switch language.\n2. Open the new menu item and page.\n3. Check column headers, dropdown caption, buttons and messages.",
     "All strings are localized (no missing resource keys / English fallback unexpected); layout not broken.", "Medium", "Localization", "ULT-10610"),
    ("Concurrent correction",
     "Two users with the same card open.",
     "1. User 1 saves AA.\n2. User 2 saves US from the stale page.",
     "Last save wins or conflict is reported per design; history shows both versions.", "Low", "Negative", "ULT-10610"),
])))

SECTIONS.append(("7. Card History (ULT-10611)", mk([
    ("On-demand export shows per-cycle value",
     "Cycles shipped with US and AA.",
     "1. Reports > Card History; export on demand.\n2. Review the Steel Origin column.",
     "Each row shows the value of its cycle.", "High", "Functional", "ULT-10611"),
    ("Scheduled export shows per-cycle value",
     "Scheduled Card History report configured.",
     "1. Run the schedule.\n2. Open the output.",
     "Same column and values as the on-demand export.", "High", "Functional", "ULT-10611"),
    ("Completed cycles keep their own value",
     "Cycle 1 US received, cycle 2 AA received.",
     "1. Export Card History.",
     "Cycle 1 = US; cycle 2 = AA.", "High", "Functional", "ULT-10611"),
    ("Corrected values show",
     "Value corrected US to AA (section 6).",
     "1. Export Card History (on-demand and scheduled).",
     "The corrected value AA is shown.", "High", "Functional", "ULT-10611/10610"),
    ("Unflagged rows blank",
     "Unflagged card.",
     "1. Export Card History.",
     "Steel Origin cell is empty.", "High", "Functional", "ULT-10611"),
    ("Customer layouts show no regression",
     "Customer-specific UI config overrides (Custom*UIConfig).",
     "1. Export Card History for each customer layout.\n2. Compare with the pre-release export.",
     "Existing columns/order unchanged; row count unchanged.", "High", "Regression", "ULT-10611"),
    ("Buyer and supplier exports consistent",
     "Same data.",
     "1. Export on the buyer and supplier sides.",
     "Both exports show identical Steel Origin values.", "Medium", "Functional", "ULT-10611"),
])))

SECTIONS.append(("8. Copeland Validation (sign-off)", mk([
    ("Print sample label - US",
     "Production-like print setup; flagged card shipped with US.",
     "1. Print the label through the desktop print client.\n2. Label the sample (case, card, date, tester).",
     "Physical label with 'Steel Origin: US', 5ZUS linear barcode.", "High", "Validation", "Copeland"),
    ("Print sample label - AA",
     "Flagged card shipped with AA.",
     "1. Print and label the sample.",
     "Physical label with 'Steel Origin: AA', 5ZAA linear barcode.", "High", "Validation", "Copeland"),
    ("Print sample label - unflagged card",
     "Unflagged card.",
     "1. Print and label the sample.",
     "Physical label with no Steel Origin text or barcode; same as the baseline.", "High", "Validation", "Copeland"),
    ("Print sample - PX2dBarCode variant",
     "Samples for US, AA and unflagged.",
     "1. Print each case using the PX2dBarCode label.",
     "Samples for each case with 2D payload per design.", "High", "Validation", "Copeland"),
    ("Print sample - PXIR2dBarCode variant",
     "As above.",
     "1. Print each case using the PXIR2dBarCode label.",
     "Samples for each case.", "High", "Validation", "Copeland"),
    ("Pre-send hex verification of samples",
     "Printed samples.",
     "1. Scan every sample into a hex viewer.\n2. Confirm the expected bytes (0x1D 5Z<value> 0x1E 0x04) and the baseline for unflagged.\n3. Attach the hex dumps to the ticket.",
     "All samples match expectation; evidence attached.", "High", "Validation", "Copeland"),
    ("Send samples / scans to David / Copeland IT",
     "Verified samples.",
     "1. Ship the physical labels or send high-resolution scans with the hex dumps to David / Copeland IT.\n2. Include the agreed Data Identifier, the segment order and the format description.\n3. Log the date and contact in the ticket.",
     "Samples delivered and receipt acknowledged.", "High", "Validation", "Copeland"),
    ("Copeland parser decodes the Steel Origin field with the agreed DI",
     "Samples received by Copeland.",
     "1. Copeland scans the samples with their scanner and parser.\n2. Check the decoded Steel Origin field for US and AA.\n3. Check that the unflagged sample parses as before.\n4. Check that P, Q, V, S, K, 5K are parsed unchanged.",
     "Parser decodes Steel Origin = US / AA with the agreed DI; unflagged and the existing fields unaffected; no parse errors.", "High", "Validation", "Copeland"),
    ("Parser mismatch / rejection handling",
     "Copeland reports a problem.",
     "1. Record the symptoms and samples.\n2. Adjust STEEL_ORIGIN_DI or layout as agreed; reprint.\n3. Resend.",
     "Defect or config change logged; samples re-validated until confirmed.", "Medium", "Validation", "Copeland"),
    ("Record Copeland confirmation in the ticket",
     "Confirmation received.",
     "1. Add a comment on ULT-10612 with the confirmer name, date, sample references and DI.\n2. Attach the confirmation email / screenshot.",
     "Confirmation recorded in the ticket; open item closed; ticket ready for sign-off.", "High", "Validation", "Copeland"),
])))

total = sum(len(c) for _, c in SECTIONS)

STORY_MATRIX = [
    ("ULT-10605 - Configuration", "Section 2"),
    ("ULT-10606 / ULT-10607 - Ship (Ship Cards, Ship by Boxes, Consolidated Ship)", "Section 3"),
    ("Feature 3 - Unaffected supplier", "Section 1 (baseline), Section 4"),
    ("ULT-10608 / ULT-10609 - Label and 2D barcodes", "Section 5"),
    ("ULT-10610 - Correction page", "Section 6"),
    ("ULT-10611 - Card History", "Section 7"),
    ("Copeland validation / sign-off", "Section 8"),
]


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def write(cell, text, size=8, bold=False, color=None, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lines = text.split("\n")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if i < len(lines) - 1:
            run.add_break()


def heading(doc, text, size=13):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def set_widths(table, widths):
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Inches(0.5))

t = doc.add_paragraph()
r = t.add_run("ULT-10612 - Test Cases")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
doc.add_paragraph("Steel Origin | End-to-end QA and Copeland barcode scan validation").runs[0].font.size = Pt(11)

info = doc.add_table(rows=0, cols=2)
info.style = "Table Grid"
for k, v in [
    ("Ticket", "ULT-10612 (Epic ULT-10604) - Task - Status: In Progress - Fix version 10.10.6"),
    ("Labels", "Copeland, SteelOrigin"),
    ("Runs after", "ULT-10605 ... ULT-10611"),
    ("Sign-off needs", "Copeland parser confirmation (open item)"),
    ("Goal", "Full regression of the Steel Origin feature on a 10.10.6 QA site configured like Copeland: Emerson print plug-in "
             "(CUSTOM_PRINT_ASSEMBLY_NAME = Leansuite.Custom.Emerson:CustomPrintScanSvc.PrintData) and Copeland's updated .lbl; "
             "then hand physical label samples to Copeland."),
    ("Scope", "Configuration (ULT-10605), Ship (ULT-10606/10607) on Ship Cards, Ship by Boxes and Consolidated Ship, unaffected supplier (Feature 3), "
              "Label and 2D barcodes PX2dBarCode / PXIR2dBarCode (ULT-10608/10609) via the desktop print client, Correction page (ULT-10610), "
              "Card History (ULT-10611), Copeland validation."),
    ("Important", "Capture the pre-release baseline (label and both 2D variants) BEFORE deploying 10.10.6 (TC-02 to TC-04)."),
    ("Total test cases", str(total)),
]:
    row = info.add_row()
    write(row.cells[0], k, size=9, bold=True)
    shade(row.cells[0], "DCE6F1")
    write(row.cells[1], v, size=9)
set_widths(info, [1.5, 8.5])

for title, cases in SECTIONS:
    heading(doc, f"{title} ({len(cases)})")
    tbl = doc.add_table(rows=1, cols=len(COLS))
    tbl.style = "Table Grid"
    for i, h in enumerate(COLS):
        write(tbl.rows[0].cells[i], h, size=9, bold=True, color="FFFFFF", center=True)
        shade(tbl.rows[0].cells[i], "1F4E79")
    for k, case in enumerate(cases):
        row = tbl.add_row()
        for i, val in enumerate(case):
            write(row.cells[i], val, bold=(i == 0), center=(i in (0, 5, 6, 7)))
            if k % 2:
                shade(row.cells[i], "F2F2F2")
    set_widths(tbl, WIDTHS)

heading(doc, "Traceability")
m = doc.add_table(rows=1, cols=2)
m.style = "Table Grid"
for i, h in enumerate(["Story / Area", "Where covered"]):
    write(m.rows[0].cells[i], h, size=9, bold=True, color="FFFFFF")
    shade(m.rows[0].cells[i], "1F4E79")
for a, b in STORY_MATRIX:
    row = m.add_row()
    write(row.cells[0], a, size=9)
    write(row.cells[1], b, size=9)
set_widths(m, [5.0, 5.0])

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"Created: {OUT} ({total} test cases)")
