from pathlib import Path
from docx import Document

src = Path(r'c:\Prcatices-AI\docs\ULT-10595_test_plan.md')
out = Path(r'c:\Prcatices-AI\docs\ULT-10595_test_plan.docx')

text = src.read_text(encoding='utf-8')
doc = Document()

for line in text.splitlines():
    s = line.strip()
    if not s:
        doc.add_paragraph('')
    elif s.startswith('# '):
        doc.add_heading(s[2:], level=1)
    elif s.startswith('## '):
        doc.add_heading(s[3:], level=2)
    elif s.startswith('### '):
        doc.add_heading(s[4:], level=3)
    elif s.startswith('---'):
        continue
    elif s.startswith('|'):
        doc.add_paragraph(s)
    else:
        doc.add_paragraph(s)

doc.save(out)
print(f'Created: {out}')
